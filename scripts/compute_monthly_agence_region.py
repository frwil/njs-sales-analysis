"""
Calcule les volumes mensuels par agence et par région (objectif vs réalisé) pour S1 2026.
- Par agence: 14 agences × 6 mois × (obj, réel)
- Par région: 3 régions × 6 mois × (obj, réel)
"""
import openpyxl
import json
import re
from collections import defaultdict

with open('/home/z/my-project/scripts/product_category_map.json') as f:
    cat_map = json.load(f)

MANUAL_WEIGHTS = {"M1051": 50.0, "CF101": 1.0, "S101": 1.0, "P109": 25.0}
SERVICES_VALIDEE = {"PONT_BASCULE", "CONTRIBUTION_CARBURANT"}
WEIGHT_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(KG|Kg|kg|KG|GRAMMES?|G|L)\b", re.IGNORECASE)

def parse_weight_kg(ref, desc):
    ref_str = str(ref).strip() if ref else ""
    if ref_str in MANUAL_WEIGHTS: return MANUAL_WEIGHTS[ref_str]
    if not desc: return 0.0
    matches = WEIGHT_RE.findall(str(desc))
    if not matches: return 0.0
    val_str, unit = matches[-1]
    val = float(val_str)
    u = unit.upper()
    if u == "KG": return val
    if u in ("G", "GRAMME", "GRAMMES"): return val/1000.0
    if u == "L": return val
    return 0.0

def normalize_agence(agence_str):
    if not agence_str: return ""
    s = str(agence_str).strip()
    if s.upper().startswith("SPC") or s.upper().startswith("PDC"): return ""
    if "AGENCE " in s.upper(): s = s[7:]
    if s.upper().startswith("DE "): s = s[3:]
    if "BAMENDA" in s.upper() and "MBOUDA" in s.upper(): return "Mbouda"
    if "BERTOUA" in s.upper(): return "Bertoua"
    if s: s = s[0].upper() + s[1:].lower()
    return s

AGENCE_REGION = {
    "Ahala": "Centre", "Bertoua": "Centre", "Buea": "Littoral",
    "Djeleng": "Ouest", "Famla": "Ouest", "Mbouda": "Ouest",
    "Messassi": "Centre", "Ndobo": "Littoral", "Ngaoundere": "Centre",
    "Nkoabang": "Centre", "Nkolbisson": "Centre", "Nkongsamba": "Littoral",
    "Pk11": "Littoral", "Village": "Littoral",
}

# ===== LIRE VENTES 2026 (mensuel par agence) =====
wb = openpyxl.load_workbook('/home/z/my-project/upload/ventes janv a juin 2026.xlsx', read_only=True, data_only=True)
sheet_to_month = {"Sheet 1": 1, "Feuil1": 2, "Feuil2": 3, "Feuil3": 4, "Feuil4": 5, "Feuil5": 6}

agence_month_vol = defaultdict(lambda: defaultdict(float))  # agence → month → vol
region_month_vol = defaultdict(lambda: defaultdict(float))  # region → month → vol

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    if ws.cell(2, 1).value != 'Réf. produit': continue
    month = sheet_to_month.get(sheet_name)
    if month is None: continue
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or len(row) < 16: continue
        ref = row[0]
        desc = row[1]
        qte = row[2]
        ca = row[8]
        etat = row[15]
        agence = row[17] if len(row) > 17 else None
        if not ref: continue
        ref_str = str(ref).strip()
        etat_str = str(etat) if etat else ""
        if etat_str != "Livrée":
            if not (ref_str in SERVICES_VALIDEE and etat_str == "Validée"):
                continue
        try:
            q = float(qte) if qte else 0
            c = float(ca) if ca else 0
        except: continue
        if ref_str == "M1051" and c == 0: continue
        weight = parse_weight_kg(ref_str, desc)
        vol_t = q * weight / 1000.0
        
        agence_norm = normalize_agence(agence)
        if agence_norm:
            agence_month_vol[agence_norm][month] += vol_t
            region = AGENCE_REGION.get(agence_norm, "Autre")
            region_month_vol[region][month] += vol_t

wb.close()

# ===== LIRE OBJECTIFS MENSUELS PAR AGENCE =====
with open('/home/z/my-project/scripts/objectives_comparison.json') as f:
    obj_data = json.load(f)

# agence_objectives[agence][category][month_str] = tons
agence_month_obj = defaultdict(lambda: defaultdict(float))
for agence, cats in obj_data.get('agence_objectives', {}).items():
    agence_norm = normalize_agence(agence)
    if not agence_norm: continue
    for cat, months in cats.items():
        for m_str, vol in months.items():
            m = int(m_str)
            if m <= 6:  # S1 uniquement
                agence_month_obj[agence_norm][m] += float(vol)

# Région mensuel objectif
region_month_obj = defaultdict(lambda: defaultdict(float))
for agence_norm, months in agence_month_obj.items():
    region = AGENCE_REGION.get(agence_norm, "Autre")
    for m, vol in months.items():
        region_month_obj[region][m] += vol

# ===== AFFICHAGE PAR AGENCE =====
print("=" * 120)
print("TENDANCE MENSUELLE PAR AGENCE — S1 2026 (volume objectif vs réalisé, tonnes)")
print("=" * 120)
print(f"\n{'Agence':<15}", end="")
for m in range(1, 7):
    print(f"{'M'+str(m)+' obj':>10}{'M'+str(m)+' réel':>10}{'%':>7}", end="")
print(f"{'S1 obj':>10}{'S1 réel':>10}{'%':>7}")
print("-" * 130)

agences_order = ["Famla", "Ndobo", "Messassi", "Djeleng", "Mbouda", "Village", "Bertoua",
                 "Nkongsamba", "Nkoabang", "Ngaoundere", "Buea", "Ahala", "Nkolbisson", "Pk11"]

total_s1_obj = 0
total_s1_real = 0
for ag in agences_order:
    print(f"{ag:<15}", end="")
    s1_obj = 0
    s1_real = 0
    for m in range(1, 7):
        obj = agence_month_obj.get(ag, {}).get(m, 0)
        real = agence_month_vol.get(ag, {}).get(m, 0)
        pct = real/obj*100 if obj > 0 else 0
        print(f"{obj:>10.0f}{real:>10.0f}{pct:>6.0f}%", end="")
        s1_obj += obj
        s1_real += real
    pct_s1 = s1_real/s1_obj*100 if s1_obj > 0 else 0
    print(f"{s1_obj:>10.0f}{s1_real:>10.0f}{pct_s1:>6.0f}%")
    total_s1_obj += s1_obj
    total_s1_real += s1_real

print("-" * 130)
pct_total = total_s1_real/total_s1_obj*100 if total_s1_obj > 0 else 0
print(f"{'TOTAL':<15}", end="")
for m in range(1, 7):
    obj = sum(agence_month_obj.get(ag, {}).get(m, 0) for ag in agences_order)
    real = sum(agence_month_vol.get(ag, {}).get(m, 0) for ag in agences_order)
    pct = real/obj*100 if obj > 0 else 0
    print(f"{obj:>10.0f}{real:>10.0f}{pct:>6.0f}%", end="")
print(f"{total_s1_obj:>10.0f}{total_s1_real:>10.0f}{pct_total:>6.0f}%")

# ===== AFFICHAGE PAR RÉGION =====
print("\n" + "=" * 120)
print("TENDANCE MENSUELLE PAR RÉGION — S1 2026 (volume objectif vs réalisé, tonnes)")
print("=" * 120)
print(f"\n{'Région':<12}", end="")
for m in range(1, 7):
    print(f"{'M'+str(m)+' obj':>10}{'M'+str(m)+' réel':>10}{'%':>7}", end="")
print(f"{'S1 obj':>10}{'S1 réel':>10}{'%':>7}")
print("-" * 130)

regions_order = ["Ouest", "Centre", "Littoral"]
total_r_obj = 0
total_r_real = 0
for r in regions_order:
    print(f"{r:<12}", end="")
    s1_obj = 0
    s1_real = 0
    for m in range(1, 7):
        obj = region_month_obj.get(r, {}).get(m, 0)
        real = region_month_vol.get(r, {}).get(m, 0)
        pct = real/obj*100 if obj > 0 else 0
        print(f"{obj:>10.0f}{real:>10.0f}{pct:>6.0f}%", end="")
        s1_obj += obj
        s1_real += real
    pct_s1 = s1_real/s1_obj*100 if s1_obj > 0 else 0
    print(f"{s1_obj:>10.0f}{s1_real:>10.0f}{pct_s1:>6.0f}%")
    total_r_obj += s1_obj
    total_r_real += s1_real

print("-" * 130)
pct_total_r = total_r_real/total_r_obj*100 if total_r_obj > 0 else 0
print(f"{'TOTAL':<12}", end="")
for m in range(1, 7):
    obj = sum(region_month_obj.get(r, {}).get(m, 0) for r in regions_order)
    real = sum(region_month_vol.get(r, {}).get(m, 0) for r in regions_order)
    pct = real/obj*100 if obj > 0 else 0
    print(f"{obj:>10.0f}{real:>10.0f}{pct:>6.0f}%", end="")
print(f"{total_r_obj:>10.0f}{total_r_real:>10.0f}{pct_total_r:>6.0f}%")

# Sauvegarder
output = {
    'agence_month': {
        ag: {
            'obj': {str(m): agence_month_obj.get(ag, {}).get(m, 0) for m in range(1, 7)},
            'real': {str(m): agence_month_vol.get(ag, {}).get(m, 0) for m in range(1, 7)},
        } for ag in agences_order
    },
    'region_month': {
        r: {
            'obj': {str(m): region_month_obj.get(r, {}).get(m, 0) for m in range(1, 7)},
            'real': {str(m): region_month_vol.get(r, {}).get(m, 0) for m in range(1, 7)},
        } for r in regions_order
    }
}
with open('/home/z/my-project/scripts/monthly_agence_region.json', 'w', encoding='utf-8') as f:
    json.dump(output, f, ensure_ascii=False, indent=2)
print("\n✓ Sauvegardé dans monthly_agence_region.json")
