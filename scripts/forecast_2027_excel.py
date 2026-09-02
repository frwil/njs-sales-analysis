from openpyxl.utils import get_column_letter
"""
Excel - Forecast 2027 complet (12 mois, S3)
"""
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import os

fcst = pd.read_csv("/home/z/my-project/scripts/forecast_2027_S3.csv", parse_dates=['date'])
print(f"Loaded {len(fcst)} forecast records")

desc_df = pd.read_csv("/home/z/my-project/scripts/dataset_consolide_v2.csv", low_memory=False)
desc_map = desc_df[['ref', 'description']].drop_duplicates().set_index('ref')['description'].to_dict()

HEAD_FILL = PatternFill('solid', fgColor='1F4E78')
HEAD_FONT = Font(bold=True, color='FFFFFF', size=11)
SUBHEAD_FILL = PatternFill('solid', fgColor='D9E1F2')
SUBHEAD_FONT = Font(bold=True, color='1F4E78', size=11)
TOTAL_FILL = PatternFill('solid', fgColor='FFF2CC')
TOTAL_FONT = Font(bold=True, size=11)
S3_COLOR = 'C6EFCE'
THIN = Side(border_style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FAMILIES = ['TOURTEAUX', 'CONCENTRES', 'INGREDIENTS', 'ALIMENT_COMPLET', 'MATERIEL_ELEVAGE', 'PREMIX']
MONTHS = list(range(1, 13))
MONTH_NAMES = ['Jan', 'Fév', 'Mar', 'Avr', 'Mai', 'Juin', 'Juil', 'Août', 'Sep', 'Oct', 'Nov', 'Déc']
QUARTERS = {'Q1 (Jan-Mar)': [1,2,3], 'Q2 (Avr-Juin)': [4,5,6], 'Q3 (Juil-Sept)': [7,8,9], 'Q4 (Oct-Déc)': [10,11,12]}

def style_header_row(ws, row, n_cols):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL; cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER

def style_total_row(ws, row, n_cols):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = TOTAL_FILL; cell.font = TOTAL_FONT; cell.border = BORDER

def style_data_row(ws, row, n_cols, color=None):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.border = BORDER
        if color: cell.fill = PatternFill('solid', fgColor=color)

wb = openpyxl.Workbook()
wb.remove(wb.active)

# === Sheet 1: Synthèse ===
ws = wb.create_sheet("1. Synthèse")
ws['A1'] = 'BELGOCAM SA - Forecast 2027 (Janvier - Décembre)'
ws['A1'].font = Font(bold=True, size=16, color='1F4E78')
ws['A2'] = 'Scénario S3 | Désaisonnalisation effet soja | En cours+Validées inclus | Prix soja 25 000 FCFA/sac'
ws['A2'].font = Font(italic=True, size=10, color='595959')

total_t = fcst['tonnes'].sum()
total_ca = fcst['ca_m_fcfa'].sum()

row = 4
ws.cell(row=row, column=1, value='SYNTHÈSE GLOBALE 2027')
ws.cell(row=row, column=1).font = SUBHEAD_FONT
row += 1

headers = ['Indicateur', 'Valeur', 'Unité']
for i, h in enumerate(headers, 1): ws.cell(row=row, column=i, value=h)
style_header_row(ws, row, 3)
row += 1

global_data = [
    ('Volume total 2027', f"{total_t:,.0f}", 'tonnes'),
    ('CA total 2027', f"{total_ca:,.1f}", 'M FCFA'),
    ('Période forecast', 'Janvier - Décembre 2027', '12 mois'),
    ('Scénario', 'S3 - Réappro soja 100%', '—'),
    ('Nb familles', '6', 'TOURTEAUX, CONCENTRÉS, INGREDIENTS, ALIMENT_COMPLET, MATERIEL_ELEVAGE, PREMIX'),
    ('Nb produits', f"{fcst['ref'].nunique()}", 'références'),
    ('Nb agences', f"{fcst['agence'].nunique()}", 'agences'),
    ('Méthode', 'Prophet (4 familles) + Extrapolation (MAT+PREMIX)', '—'),
    ('Désaisonnalisation', 'Effet soja Jul-Août 2026 neutralisé (cap S1 avg)', '—'),
    ('En cours + Validées', 'Incluses comme potentielles Livrées', '—'),
]
for label, val, unit in global_data:
    ws.cell(row=row, column=1, value=label).font = Font(bold=True)
    ws.cell(row=row, column=2, value=val)
    ws.cell(row=row, column=3, value=unit)
    style_data_row(ws, row, 3, color=S3_COLOR)
    row += 1

row += 1
ws.cell(row=row, column=1, value='SYNTHÈSE PAR FAMILLE')
ws.cell(row=row, column=1).font = SUBHEAD_FONT
row += 1

headers = ['Famille', 'Volume 2027 (t)', 'CA 2027 (M FCFA)', 'Part CA (%)']
for i, h in enumerate(headers, 1): ws.cell(row=row, column=i, value=h)
style_header_row(ws, row, 4)
row += 1

fam_synth = fcst.groupby('family').agg(tonnes=('tonnes', 'sum'), ca_m_fcfa=('ca_m_fcfa', 'sum')).reindex(FAMILIES)
for fam in FAMILIES:
    if fam in fam_synth.index:
        t = fam_synth.loc[fam, 'tonnes']
        ca = fam_synth.loc[fam, 'ca_m_fcfa']
        pct = ca / total_ca * 100 if total_ca > 0 else 0
        ws.cell(row=row, column=1, value=fam)
        ws.cell(row=row, column=2, value=round(t, 0))
        ws.cell(row=row, column=3, value=round(ca, 1))
        ws.cell(row=row, column=4, value=f"{pct:.1f}%")
        style_data_row(ws, row, 4, color=S3_COLOR)
        row += 1
ws.cell(row=row, column=1, value='TOTAL')
ws.cell(row=row, column=2, value=round(total_t, 0))
ws.cell(row=row, column=3, value=round(total_ca, 1))
ws.cell(row=row, column=4, value='100.0%')
style_total_row(ws, row, 4)

row += 2
ws.cell(row=row, column=1, value='SYNTHÈSE PAR TRIMESTRE')
ws.cell(row=row, column=1).font = SUBHEAD_FONT
row += 1
headers = ['Trimestre', 'Volume (t)', 'CA (M FCFA)', 'Part CA (%)']
for i, h in enumerate(headers, 1): ws.cell(row=row, column=i, value=h)
style_header_row(ws, row, 4)
row += 1
for qname, qmonths in QUARTERS.items():
    sub = fcst[fcst['month'].isin(qmonths)]
    t = sub['tonnes'].sum()
    ca = sub['ca_m_fcfa'].sum()
    pct = ca / total_ca * 100 if total_ca > 0 else 0
    ws.cell(row=row, column=1, value=qname)
    ws.cell(row=row, column=2, value=round(t, 0))
    ws.cell(row=row, column=3, value=round(ca, 1))
    ws.cell(row=row, column=4, value=f"{pct:.1f}%")
    style_data_row(ws, row, 4, color=S3_COLOR)
    row += 1
ws.cell(row=row, column=1, value='TOTAL')
ws.cell(row=row, column=2, value=round(total_t, 0))
ws.cell(row=row, column=3, value=round(total_ca, 1))
ws.cell(row=row, column=4, value='100.0%')
style_total_row(ws, row, 4)

ws.column_dimensions['A'].width = 30
ws.column_dimensions['B'].width = 20
ws.column_dimensions['C'].width = 20
ws.column_dimensions['D'].width = 12

# === Sheet 2: Par Famille × Mois ===
ws = wb.create_sheet("2. Par Famille × Mois")
ws['A1'] = 'FORECAST 2027 - DÉTAIL PAR FAMILLE × MOIS'
ws['A1'].font = Font(bold=True, size=14, color='1F4E78')

row = 3
headers = ['Famille'] + MONTH_NAMES + ['Total (t)', 'CA (M FCFA)']
for i, h in enumerate(headers, 1): ws.cell(row=row, column=i, value=h)
style_header_row(ws, row, len(headers))
row += 1

for fam in FAMILIES:
    sub = fcst[fcst['family'] == fam]
    if len(sub) == 0: continue
    monthly_t = sub.groupby('month')['tonnes'].sum().to_dict()
    monthly_ca = sub.groupby('month')['ca_m_fcfa'].sum().to_dict()
    total_t_fam = sub['tonnes'].sum()
    total_ca_fam = sub['ca_m_fcfa'].sum()
    ws.cell(row=row, column=1, value=fam)
    for i, m in enumerate(MONTHS, 2):
        ws.cell(row=row, column=i, value=round(monthly_t.get(m, 0), 0))
    ws.cell(row=row, column=14, value=round(total_t_fam, 0))
    ws.cell(row=row, column=15, value=round(total_ca_fam, 1))
    style_data_row(ws, row, 15, color=S3_COLOR)
    row += 1
ws.cell(row=row, column=1, value='TOTAL')
for i, m in enumerate(MONTHS, 2):
    ws.cell(row=row, column=i, value=round(fcst.groupby('month')['tonnes'].sum().get(m, 0), 0))
ws.cell(row=row, column=14, value=round(total_t, 0))
ws.cell(row=row, column=15, value=round(total_ca, 1))
style_total_row(ws, row, 15)

ws.column_dimensions['A'].width = 22
for col in [get_column_letter(c) for c in range(2, 16)]:
    ws.column_dimensions[col].width = 10

# === Sheet 3: Par Région × Mois ===
ws = wb.create_sheet("3. Par Région × Mois")
ws['A1'] = 'FORECAST 2027 - DÉTAIL PAR RÉGION × MOIS'
ws['A1'].font = Font(bold=True, size=14, color='1F4E78')

row = 3
headers = ['Région'] + MONTH_NAMES + ['Total (t)', 'CA (M FCFA)']
for i, h in enumerate(headers, 1): ws.cell(row=row, column=i, value=h)
style_header_row(ws, row, len(headers))
row += 1
for reg in ['Ouest', 'Centre', 'Littoral']:
    sub = fcst[fcst['region'] == reg]
    monthly_t = sub.groupby('month')['tonnes'].sum().to_dict()
    ws.cell(row=row, column=1, value=reg)
    for i, m in enumerate(MONTHS, 2):
        ws.cell(row=row, column=i, value=round(monthly_t.get(m, 0), 0))
    ws.cell(row=row, column=14, value=round(sub['tonnes'].sum(), 0))
    ws.cell(row=row, column=15, value=round(sub['ca_m_fcfa'].sum(), 1))
    style_data_row(ws, row, 15, color=S3_COLOR)
    row += 1
ws.cell(row=row, column=1, value='TOTAL')
for i, m in enumerate(MONTHS, 2):
    ws.cell(row=row, column=i, value=round(fcst.groupby('month')['tonnes'].sum().get(m, 0), 0))
ws.cell(row=row, column=14, value=round(total_t, 0))
ws.cell(row=row, column=15, value=round(total_ca, 1))
style_total_row(ws, row, 15)
ws.column_dimensions['A'].width = 15
for col in [get_column_letter(c) for c in range(2, 16)]:
    ws.column_dimensions[col].width = 10

# === Sheet 4: Par Agence ===
ws = wb.create_sheet("4. Par Agence")
ws['A1'] = 'FORECAST 2027 - DÉTAIL PAR AGENCE'
ws['A1'].font = Font(bold=True, size=14, color='1F4E78')
row = 3
headers = ['Agence', 'Région', 'Volume 2027 (t)', 'CA 2027 (M FCFA)', 'Part CA (%)']
for i, h in enumerate(headers, 1): ws.cell(row=row, column=i, value=h)
style_header_row(ws, row, 5)
row += 1
by_ag = fcst.groupby(['agence', 'region']).agg(tonnes=('tonnes', 'sum'), ca_m_fcfa=('ca_m_fcfa', 'sum')).reset_index().sort_values('ca_m_fcfa', ascending=False)
for _, r in by_ag.iterrows():
    pct = r['ca_m_fcfa'] / total_ca * 100 if total_ca > 0 else 0
    ws.cell(row=row, column=1, value=r['agence'])
    ws.cell(row=row, column=2, value=r['region'])
    ws.cell(row=row, column=3, value=round(r['tonnes'], 0))
    ws.cell(row=row, column=4, value=round(r['ca_m_fcfa'], 1))
    ws.cell(row=row, column=5, value=f"{pct:.1f}%")
    style_data_row(ws, row, 5, color=S3_COLOR)
    row += 1
ws.cell(row=row, column=1, value='TOTAL')
ws.cell(row=row, column=3, value=round(total_t, 0))
ws.cell(row=row, column=4, value=round(total_ca, 1))
ws.cell(row=row, column=5, value='100.0%')
style_total_row(ws, row, 5)
ws.column_dimensions['A'].width = 18; ws.column_dimensions['B'].width = 12
for col in 'CDE': ws.column_dimensions[col].width = 18

# === Sheet 5: Par Produit ===
ws = wb.create_sheet("5. Par Produit")
ws['A1'] = 'FORECAST 2027 - DÉTAIL PAR PRODUIT'
ws['A1'].font = Font(bold=True, size=14, color='1F4E78')
row = 3
headers = ['Réf', 'Description', 'Famille', 'Volume 2027 (t)', 'CA 2027 (M FCFA)', 'Part CA (%)']
for i, h in enumerate(headers, 1): ws.cell(row=row, column=i, value=h)
style_header_row(ws, row, 6)
row += 1
by_prod = fcst.groupby(['ref', 'family']).agg(tonnes=('tonnes', 'sum'), ca_m_fcfa=('ca_m_fcfa', 'sum')).reset_index().sort_values('ca_m_fcfa', ascending=False)
for _, r in by_prod.iterrows():
    pct = r['ca_m_fcfa'] / total_ca * 100 if total_ca > 0 else 0
    ws.cell(row=row, column=1, value=r['ref'])
    ws.cell(row=row, column=2, value=desc_map.get(r['ref'], ''))
    ws.cell(row=row, column=3, value=r['family'])
    ws.cell(row=row, column=4, value=round(r['tonnes'], 1))
    ws.cell(row=row, column=5, value=round(r['ca_m_fcfa'], 1))
    ws.cell(row=row, column=6, value=f"{pct:.1f}%")
    style_data_row(ws, row, 6, color=S3_COLOR)
    row += 1
ws.cell(row=row, column=1, value='TOTAL')
ws.cell(row=row, column=4, value=round(total_t, 0))
ws.cell(row=row, column=5, value=round(total_ca, 1))
ws.cell(row=row, column=6, value='100.0%')
style_total_row(ws, row, 6)
ws.column_dimensions['A'].width = 12; ws.column_dimensions['B'].width = 35; ws.column_dimensions['C'].width = 22
for col in 'DEF': ws.column_dimensions[col].width = 18

# === Sheet 6: Détail complet ===
ws = wb.create_sheet("6. Détail complet")
ws['A1'] = 'FORECAST 2027 - DÉTAIL COMPLET (Produit × Agence × Mois)'
ws['A1'].font = Font(bold=True, size=14, color='1F4E78')
row = 3
headers = ['Réf', 'Description', 'Famille', 'Agence', 'Région', 'Mois', 'Année', 'Tonnes', 'Sacs 50kg', 'Prix TTC/sac', 'CA (M FCFA)']
for i, h in enumerate(headers, 1): ws.cell(row=row, column=i, value=h)
style_header_row(ws, row, 11)
row += 1
for _, r in fcst.sort_values(['family', 'ref', 'agence', 'month']).iterrows():
    ws.cell(row=row, column=1, value=r['ref'])
    ws.cell(row=row, column=2, value=desc_map.get(r['ref'], ''))
    ws.cell(row=row, column=3, value=r['family'])
    ws.cell(row=row, column=4, value=r['agence'])
    ws.cell(row=row, column=5, value=r['region'])
    ws.cell(row=row, column=6, value=f"{int(r['month']):02d}/2027")
    ws.cell(row=row, column=7, value=int(r['year']))
    ws.cell(row=row, column=8, value=round(r['tonnes'], 2))
    ws.cell(row=row, column=9, value=round(r['sacs_50'], 1))
    ws.cell(row=row, column=10, value=int(r['prix_ttc_sac']) if r['prix_ttc_sac'] else 0)
    ws.cell(row=row, column=11, value=round(r['ca_m_fcfa'], 2))
    style_data_row(ws, row, 11)
    row += 1
ws.column_dimensions['A'].width = 12; ws.column_dimensions['B'].width = 32; ws.column_dimensions['C'].width = 22
ws.column_dimensions['D'].width = 15; ws.column_dimensions['E'].width = 12
for col in 'FGHIJK': ws.column_dimensions[col].width = 13
ws.freeze_panes = 'A4'

# === Sheet 7: Hypothèses ===
ws = wb.create_sheet("7. Hypothèses")
ws['A1'] = 'HYPOTHÈSES DU FORECAST 2027'
ws['A1'].font = Font(bold=True, size=14, color='1F4E78')
hyp = [
    ('PÉRIODE', 'Janvier - Décembre 2027 (12 mois)'),
    ('SCÉNARIO', 'S3 - Réappro soja 100% (situation normale)'),
    ('MÉTHODE', 'Prophet (4 familles alimentaires) + Extrapolation (MAT+PREMIX)'),
    ('', ''),
    ('INNOVATIONS VS FORECAST Q4 2026', ''),
    ('En cours + Validées', 'Commandes En cours et Validées d\'août 2026 incluses comme potentielles Livrées'),
    ('Désaisonnalisation soja', 'Volumes soja Jul-Août 2026 neutralisés (cap à moyenne S1) pour éviter le biais de la rupture concurrente exceptionnelle'),
    ('Prix soja actualisé', '25 000 FCFA/sac (prix au 24/08/2026) au lieu de 20 600'),
    ('Forecast 12 mois', 'Période complète Jan-Déc 2027 (vs 4 mois pour Q4 2026)'),
    ('', ''),
    ('FAMILLES INCLUSES', '6 familles (TOURTEAUX, CONCENTRÉS, INGREDIENTS, ALIMENT_COMPLET, MATERIEL_ELEVAGE, PREMIX)'),
    ('FAMILLES EXCLUES', 'MAIS (produit opportuniste), DIVERS, COMPLEMENT ALIMENTAIRE'),
    ('PRODUITS OPPORTUNISTES', 'MAÏS, ELVOR TONIC, CARBONATE DE CALCIUM — exclus du périmètre'),
    ('', ''),
    ('DONNÉES HISTORIQUES', '115 086 enregistrements (Jan 2025 - Août 2026) + En cours/Validées août'),
    ('LIMITES', '20 mois d\'historique (Prophet recommande 2+ ans pour yearly seasonality)'),
    ('MISE À JOUR', 'Pipeline automatisé: python scripts/forecast_pipeline.py'),
]
row = 3
for label, value in hyp:
    if label and not value:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=1).font = SUBHEAD_FONT
        ws.cell(row=row, column=1).fill = SUBHEAD_FILL
    elif label:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        ws.cell(row=row, column=1).border = BORDER
        ws.cell(row=row, column=2).border = BORDER
    row += 1
ws.column_dimensions['A'].width = 35; ws.column_dimensions['B'].width = 85

output_path = "/home/z/my-project/download/forecast_2027_S3_volumes_valeurs.xlsx"
wb.save(output_path)
print(f"\n=== EXCEL 2027 SAVED ===")
print(f"Path: {output_path}")
print(f"Size: {os.path.getsize(output_path) / 1024:.0f} KB")
print(f"Sheets: {wb.sheetnames}")
