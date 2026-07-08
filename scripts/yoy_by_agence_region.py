"""
Calcule le YoY par agence et par région en croisant les ventes 2025 (S1) et 2026 (S1).
- Fichier 2025: col 13 (Agence), col 14 (Région), col 12 (Qté en tonnes), col 7 (CA HT), col 10 (État)
- Fichier 2026: col 18 (Agence), col 9 (CA HT), col 16 (État), col 2 (Description), col 3 (Qté)
- Filtres: État=Livrée (ou services Validée pour 2026), exclure M1051 à CA=0
"""
import openpyxl
import json
import re
from collections import defaultdict

# Charger mapping
with open('/home/z/my-project/scripts/product_category_map.json') as f:
    cat_map = json.load(f)

MANUAL_WEIGHTS = {"M1051": 50.0, "CF101": 1.0, "S101": 1.0, "P109": 25.0}
SERVICES_VALIDEE = {"PONT_BASCULE", "CONTRIBUTION_CARBURANT"}
WEIGHT_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(KG|Kg|kg|KG|GRAMMES?|G|L)\b", re.IGNORECASE)

def parse_weight_kg(ref, desc):
    ref_str = str(ref).strip() if ref else ""
    if ref_str in MANUAL_WEIGHTS:
        return MANUAL_WEIGHTS[ref_str]
    if not desc: return 0.0
    matches = WEIGHT_RE.findall(str(desc))
    if not matches: return 0.0
    val_str, unit = matches[-1]
    val = float(val_str)
    u = unit.upper()
    if u == "KG": return val
    if u in ("G", "GRAMME", "GRAMMES"): return val/1000.0
    if u == "L": return val
    return 0.0

# ===== 2025 SALES (avec agence et région natives) =====
print("Reading 2025 sales...")
wb25 = openpyxl.load_workbook('/home/z/my-project/upload/86d96135-9db7-45bc-bba6-a69efa2c5ee3.xlsx', read_only=True, data_only=True)
ws25 = wb25.active

# 2025: par agence et par région
agence_2025 = defaultdict(lambda: {"vol_t": 0, "ca": 0})
region_2025 = defaultdict(lambda: {"vol_t": 0, "ca": 0})
agence_cat_2025 = defaultdict(lambda: defaultdict(lambda: {"vol_t": 0, "ca": 0}))

for row in ws25.iter_rows(min_row=2, values_only=True):
    if not row or len(row) < 14: continue
    ref = row[0]
    desc = row[1]
    ca = row[6]
    etat = row[9]
    cat = row[10]
    vol_t = row[11]
    agence = row[12]
    region = row[13]
    
    if not ref: continue
    etat_str = str(etat) if etat else ""
    if etat_str != "Livrée": continue
    
    try:
        v = float(vol_t) if vol_t else 0
        c = float(ca) if ca else 0
    except: continue
    
    agence_str = str(agence).strip() if agence else ""
    region_str = str(region).strip() if region else ""
    
    if agence_str:
        agence_2025[agence_str]["vol_t"] += v
        agence_2025[agence_str]["ca"] += c
    if region_str:
        region_2025[region_str]["vol_t"] += v
        region_2025[region_str]["ca"] += c
    
    # par agence × catégorie (pour CONCENTRES)
    if agence_str and cat:
        cat_str = str(cat).strip()
        if cat_str == "DIVERS2": cat_str = "DIVERS"
        agence_cat_2025[agence_str][cat_str]["vol_t"] += v
        agence_cat_2025[agence_str][cat_str]["ca"] += c

wb25.close()

print(f"\n2025 S1 — {len(agence_2025)} agences, {len(region_2025)} régions")
print(f"  Volume total: {sum(d['vol_t'] for d in agence_2025.values()):.1f} t")
print(f"  CA total: {sum(d['ca'] for d in agence_2025.values())/1e6:.1f} M FCFA")

# ===== 2026 SALES (avec agence — col 18, mais pas de région native) =====
print("\nReading 2026 sales...")
wb26 = openpyxl.load_workbook('/home/z/my-project/upload/ventes janv a juin 2026.xlsx', read_only=True, data_only=True)

agence_2026 = defaultdict(lambda: {"vol_t": 0, "ca": 0})
agence_cat_2026 = defaultdict(lambda: defaultdict(lambda: {"vol_t": 0, "ca": 0}))

# Mapping agence → région (pour 2026 qui n'a pas la région native)
AGENCE_REGION = {
    "Ahala": "Centre", "Bertoua": "Centre", "Buea": "Littoral",
    "Djeleng": "Ouest", "Famla": "Ouest", "Mbouda": "Ouest",
    "Messassi": "Centre", "Ndobo": "Littoral", "Ngaoundere": "Centre",
    "Nkoabang": "Centre", "Nkolbisson": "Centre", "Nkongsamba": "Littoral",
    "Pk11": "Littoral", "Village": "Littoral",
}

def normalize_agence(agence_str):
    if not agence_str: return ""
    s = str(agence_str).strip()
    if s.upper().startswith("AGENCE "): s = s[7:]
    if s.upper().startswith("DE "): s = s[3:]
    if "BAMENDA" in s.upper() and "MBOUDA" in s.upper(): return "Mbouda"
    if "BERTOUA" in s.upper(): return "Bertoua"
    if s: s = s[0].upper() + s[1:].lower()
    return s

sheet_to_month = {"Sheet 1": 1, "Feuil1": 2, "Feuil2": 3, "Feuil3": 4, "Feuil4": 5, "Feuil5": 6}

for sheet_name in wb26.sheetnames:
    ws = wb26[sheet_name]
    if ws.cell(2, 1).value != 'Réf. produit': continue
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or len(row) < 16: continue
        ref = row[0]
        desc = row[1]
        qte = row[2]
        ca = row[8]
        etat = row[15]
        agence = row[17] if len(row) > 17 else None
        
        if not ref: continue
        ref_str = str(ref).strip()
        etat_str = str(etat) if etat else ""
        
        if etat_str != "Livrée":
            if not (ref_str in SERVICES_VALIDEE and etat_str == "Validée"):
                continue
        
        try:
            q = float(qte) if qte else 0
            c = float(ca) if ca else 0
        except: continue
        
        if ref_str == "M1051" and c == 0: continue
        
        weight = parse_weight_kg(ref_str, desc)
        vol_t = q * weight / 1000.0
        
        agence_norm = normalize_agence(agence)
        if agence_norm:
            agence_2026[agence_norm]["vol_t"] += vol_t
            agence_2026[agence_norm]["ca"] += c
            
            cat = cat_map.get(ref_str, "DIVERS")
            agence_cat_2026[agence_norm][cat]["vol_t"] += vol_t
            agence_cat_2026[agence_norm][cat]["ca"] += c

wb26.close()

# Calculer région 2026 à partir des agences
region_2026 = defaultdict(lambda: {"vol_t": 0, "ca": 0})
for agence, d in agence_2026.items():
    region = AGENCE_REGION.get(agence, "Autre")
    region_2026[region]["vol_t"] += d["vol_t"]
    region_2026[region]["ca"] += d["ca"]

print(f"\n2026 S1 — {len(agence_2026)} agences, {len(region_2026)} régions")
print(f"  Volume total: {sum(d['vol_t'] for d in agence_2026.values()):.1f} t")
print(f"  CA total: {sum(d['ca'] for d in agence_2026.values())/1e6:.1f} M FCFA")

# ===== YOY PAR AGENCE =====
print("\n" + "=" * 100)
print("YoY PAR AGENCE (S1 2025 vs S1 2026)")
print("=" * 100)
print(f"{'Agence':<20} {'Vol 2025':>10} {'Vol 2026':>10} {'YoY vol':>10} {'CA 25 (M)':>10} {'CA 26 (M)':>10} {'YoY CA':>10}")
print("-" * 90)

all_agences = sorted(set(list(agence_2025.keys()) + list(agence_2026.keys())))
yoy_agence = {}
for ag in all_agences:
    d25 = agence_2025.get(ag, {"vol_t": 0, "ca": 0})
    d26 = agence_2026.get(ag, {"vol_t": 0, "ca": 0})
    yoy_v = (d26["vol_t"] - d25["vol_t"]) / d25["vol_t"] * 100 if d25["vol_t"] > 0 else 0
    yoy_c = (d26["ca"] - d25["ca"]) / d25["ca"] * 100 if d25["ca"] > 0 else 0
    yoy_agence[ag] = {
        "vol_2025": d25["vol_t"],
        "vol_2026": d26["vol_t"],
        "yoy_vol_pct": yoy_v,
        "ca_2025": d25["ca"],
        "ca_2026": d26["ca"],
        "yoy_ca_pct": yoy_c
    }
    print(f"{ag:<20} {d25['vol_t']:>10.1f} {d26['vol_t']:>10.1f} {yoy_v:>+9.1f}% {d25['ca']/1e6:>10.1f} {d26['ca']/1e6:>10.1f} {yoy_c:>+9.1f}%")

# ===== YOY PAR RÉGION =====
print("\n" + "=" * 100)
print("YoY PAR RÉGION (S1 2025 vs S1 2026)")
print("=" * 100)
print(f"{'Région':<15} {'Vol 2025':>10} {'Vol 2026':>10} {'YoY vol':>10} {'CA 25 (M)':>10} {'CA 26 (M)':>10} {'YoY CA':>10}")
print("-" * 80)

all_regions = sorted(set(list(region_2025.keys()) + list(region_2026.keys())))
yoy_region = {}
for r in all_regions:
    d25 = region_2025.get(r, {"vol_t": 0, "ca": 0})
    d26 = region_2026.get(r, {"vol_t": 0, "ca": 0})
    yoy_v = (d26["vol_t"] - d25["vol_t"]) / d25["vol_t"] * 100 if d25["vol_t"] > 0 else 0
    yoy_c = (d26["ca"] - d25["ca"]) / d25["ca"] * 100 if d25["ca"] > 0 else 0
    yoy_region[r] = {
        "vol_2025": d25["vol_t"],
        "vol_2026": d26["vol_t"],
        "yoy_vol_pct": yoy_v,
        "ca_2025": d25["ca"],
        "ca_2026": d26["ca"],
        "yoy_ca_pct": yoy_c
    }
    print(f"{r:<15} {d25['vol_t']:>10.1f} {d26['vol_t']:>10.1f} {yoy_v:>+9.1f}% {d25['ca']/1e6:>10.1f} {d26['ca']/1e6:>10.1f} {yoy_c:>+9.1f}%")

# ===== YOY CONCENTRES PAR AGENCE =====
print("\n" + "=" * 100)
print("YoY CONCENTRÉS PAR AGENCE (S1 2025 vs S1 2026)")
print("=" * 100)
print(f"{'Agence':<20} {'Conc 2025':>10} {'Conc 2026':>10} {'YoY conc':>10}")
print("-" * 60)

yoy_conc_agence = {}
for ag in all_agences:
    d25 = agence_cat_2025.get(ag, {}).get("CONCENTRES", {"vol_t": 0, "ca": 0})
    d26 = agence_cat_2026.get(ag, {}).get("CONCENTRES", {"vol_t": 0, "ca": 0})
    yoy = (d26["vol_t"] - d25["vol_t"]) / d25["vol_t"] * 100 if d25["vol_t"] > 0 else 0
    yoy_conc_agence[ag] = {
        "conc_2025": d25["vol_t"],
        "conc_2026": d26["vol_t"],
        "yoy_pct": yoy
    }
    if d25["vol_t"] > 0 or d26["vol_t"] > 0:
        print(f"{ag:<20} {d25['vol_t']:>10.1f} {d26['vol_t']:>10.1f} {yoy:>+9.1f}%")

# Sauvegarder
output = {
    "yoy_agence": yoy_agence,
    "yoy_region": yoy_region,
    "yoy_conc_agence": yoy_conc_agence,
    "agence_2025": {ag: {"vol_t": d["vol_t"], "ca": d["ca"]} for ag, d in agence_2025.items()},
    "agence_2026": {ag: {"vol_t": d["vol_t"], "ca": d["ca"]} for ag, d in agence_2026.items()},
    "region_2025": {r: {"vol_t": d["vol_t"], "ca": d["ca"]} for r, d in region_2025.items()},
    "region_2026": {r: {"vol_t": d["vol_t"], "ca": d["ca"]} for r, d in region_2026.items()},
}
with open('/home/z/my-project/scripts/yoy_agence_region.json', 'w', encoding='utf-8') as f:
    json.dump(output, f, ensure_ascii=False, indent=2)
print(f"\n✓ Sauvegardé dans yoy_agence_region.json")
