"""
Identify internal/filiale clients to exclude from analysis.
Look for: SPC, PDC, PROVENDERIE, NJS, FILIALE, GROUPE NJS, SOLDE COMPTA, etc.
"""
from openpyxl import load_workbook
import re

SRC = "/home/z/my-project/download/ventes_livrees.xlsx"

wb = load_workbook(SRC, read_only=True, data_only=True)

# Pattern of internal/filiale keywords
INTERNAL_PATTERNS = [
    r"\bSPC\b",
    r"\bPDC\b",
    r"\bPROVENDERIE\b",
    r"\bNJS\b",
    r"\bFILIALE\b",
    r"\bGROUPE NJS\b",
    r"\bSOLDE\b",
    r"\bCOMPTA\b",
    r"\bBELGOCAM\b",
]
INTERNAL_RE = re.compile("|".join(INTERNAL_PATTERNS), re.IGNORECASE)

# Collect all unique client names
clients = {}
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for row in ws.iter_rows(min_row=3, values_only=True):
        tiers = row[5] if len(row) > 5 else None
        if tiers is None or str(tiers).strip() == "":
            continue
        if tiers not in clients:
            clients[tiers] = 0
        ca = row[8] if len(row) > 8 else 0
        try:
            clients[tiers] += float(ca) if ca is not None else 0
        except:
            pass

wb.close()

# Filter for internal-like clients
print(f"Total unique clients (raw tiers): {len(clients)}")
print(f"\nClients matching internal patterns:")
print(f"{'Tiers':<80} {'CA HT 6 mois':>15}")
print("-" * 100)
internal_clients = []
for tiers, ca in sorted(clients.items(), key=lambda x: -x[1]):
    if INTERNAL_RE.search(str(tiers)):
        internal_clients.append((tiers, ca))
        print(f"{str(tiers)[:78]:<80} {ca:>15,.0f}")

print(f"\nTotal internal-like clients: {len(internal_clients)}")
total_ca_internal = sum(c for _, c in internal_clients)
print(f"Total CA of these clients: {total_ca_internal:,.0f} FCFA")
