"""
Q1 -> Q2 transition analysis on CONCENTRÉS only (12 products: BELGO 10% + BELGO 5%, no soja).

Same methodology as transition_q1_q2.py but with CONCENTRE_REFS instead of TARGET_REFS.

Output: /home/z/my-project/download/transition_q1_q2_concentres.xlsx
"""
import re
import os
import time
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

# ---- Paths ----
SRC = "/home/z/my-project/download/ventes_livrees.xlsx"
OUT = "/home/z/my-project/download/transition_q1_q2_concentres.xlsx"

# ---- Product categories ----
SOJA_REFS = {"T102", "T1021", "T1023", "T1024"}
CONCENTRE_REFS = {
    "C102", "C1022", "C104", "C1042", "C1043", "C1044",
    "C105", "C1053", "C1054", "C1055",
    "C101", "C103",
}
# THIS time we use CONCENTRE_REFS as the target category
CATEGORY_REFS = CONCENTRE_REFS
CATEGORY_LABEL = "concentrés (12 produits : BELGO 10% + BELGO 5%)"

WEIGHT_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(KG|Kg|kg|KG|GRAMMES?|G|L)\b", re.IGNORECASE)

def parse_weight_kg(desc):
    if not desc: return 0.0
    matches = WEIGHT_RE.findall(str(desc))
    if not matches: return 0.0
    val_str, unit = matches[-1]
    val = float(val_str)
    u = unit.upper()
    if u == "KG": return val
    if u in ("G","GRAMME","GRAMMES"): return val/1000.0
    if u == "L": return val
    return 0.0

Q1_SHEETS = {"Sheet 1","Feuil1","Feuil2"}
Q2_SHEETS = {"Feuil3","Feuil4","Feuil5"}

TIERS_RE = re.compile(r"^\s*([A-Z]{2,}[\w-]*?)\s*-\s+(.+?)\s*$")

def split_tiers(tiers):
    if tiers is None: return ("","")
    s = str(tiers).strip()
    m = TIERS_RE.match(s)
    if m: return (m.group(1).strip(), m.group(2).strip())
    return ("", s)

print(f"Reading source: {SRC}")
t0 = time.time()
wb_src = load_workbook(SRC, read_only=True, data_only=True)
print(f"  Loaded in {time.time()-t0:.1f}s")

product_weight = {}
client_q1_kg = defaultdict(float)
client_q2_kg = defaultdict(float)
client_q1_ca = defaultdict(float)
client_q2_ca = defaultdict(float)
client_q1_months = defaultdict(set)
client_q2_months = defaultdict(set)
client_ca_total_all = defaultdict(float)

sheet_to_month = {
    "Sheet 1":1,"Feuil1":2,"Feuil2":3,
    "Feuil3":4,"Feuil4":5,"Feuil5":6,
}

for sheet_name in wb_src.sheetnames:
    ws = wb_src[sheet_name]
    is_q1 = sheet_name in Q1_SHEETS
    is_q2 = sheet_name in Q2_SHEETS
    if not (is_q1 or is_q2): continue
    month_num = sheet_to_month.get(sheet_name)
    if month_num is None: continue
    n = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        tiers = row[5] if len(row)>5 else None
        if tiers is None or str(tiers).strip()=="": continue
        n += 1
        ref_prod = row[0] if len(row)>0 else None
        desc = row[1] if len(row)>1 else None
        qte = row[2] if len(row)>2 else 0
        ca_ht = row[8] if len(row)>8 else 0
        ref_c, name_c = split_tiers(tiers)
        key = (ref_c, name_c)
        ref_prod_str = str(ref_prod).strip() if ref_prod is not None else ""
        if ref_prod_str and ref_prod_str not in product_weight:
            product_weight[ref_prod_str] = parse_weight_kg(desc)
        weight_kg = product_weight.get(ref_prod_str, 0.0)
        try:
            qte_f = float(qte) if qte is not None else 0.0
        except (ValueError, TypeError):
            qte_f = 0.0
        try:
            ca_f = float(ca_ht) if ca_ht is not None else 0.0
        except (ValueError, TypeError):
            ca_f = 0.0
        client_ca_total_all[key] += ca_f
        if ref_prod_str not in CATEGORY_REFS: continue
        vol_kg = qte_f * weight_kg
        if is_q1:
            client_q1_kg[key] += vol_kg
            client_q1_ca[key] += ca_f
            client_q1_months[key].add(month_num)
        else:
            client_q2_kg[key] += vol_kg
            client_q2_ca[key] += ca_f
            client_q2_months[key].add(month_num)
    print(f"  {sheet_name} (month {month_num}): {n} client occurrences")

wb_src.close()

clients_all = sorted(client_ca_total_all.keys(), key=lambda k:(k[1].upper(), k[0]))
print(f"\nTotal unique clients: {len(clients_all)}")

# Pareto
total_ca_all = sum(client_ca_total_all.values())
sorted_by_ca = sorted(clients_all, key=lambda k:-client_ca_total_all[k])
pareto_threshold = 0.80 * total_ca_all
pareto_clients = set()
cumulative = 0.0
for key in sorted_by_ca:
    cumulative += client_ca_total_all[key]
    pareto_clients.add(key)
    if cumulative >= pareto_threshold: break

def freq_cat(months_set):
    n = len(months_set)
    if n >= 3: return "Fidèle"
    elif n >= 1: return "Semi-fidèle"
    else: return "Zero"

SEGMENT_LABELS = {
    ("Zero Q1","Zero Q2"):   "Q1 Zero → Q2 Zero (persistant)",
    ("Zero Q1","Active Q2"): "Q1 Zero → Q2 Active (réactivé)",
    ("Active Q1","Zero Q2"): "Q1 Active → Q2 Zero (churned / perdu)",
    ("Active Q1","Active Q2"):"Q1 Active → Q2 Active (fidèle / retenu)",
}

client_segment = {}
for k in clients_all:
    q1_active = len(client_q1_months[k])>0
    q2_active = len(client_q2_months[k])>0
    client_segment[k] = ("Active Q1" if q1_active else "Zero Q1",
                          "Active Q2" if q2_active else "Zero Q2")

seg_counts = defaultdict(int)
for k, seg in client_segment.items():
    seg_counts[seg] += 1

print(f"\n{'='*70}")
print(f"TRANSITION MATRIX (CONCENTRÉS ONLY — 12 produits)")
print(f"{'='*70}")
print(f"{'Q1 \\ Q2':<25} {'Zero Q2':>12} {'Active Q2':>12} {'Total':>10}")
print("-"*60)

zz = seg_counts[("Zero Q1","Zero Q2")]
za = seg_counts[("Zero Q1","Active Q2")]
az = seg_counts[("Active Q1","Zero Q2")]
aa = seg_counts[("Active Q1","Active Q2")]

print(f"{'Zero Q1':<25} {zz:>12} {za:>12} {zz+za:>10}")
print(f"{'Active Q1':<25} {az:>12} {aa:>12} {az+aa:>10}")
print(f"{'Total':<25} {zz+az:>12} {za+aa:>12} {len(clients_all):>10}")

def seg_stats(seg_key):
    clients = [k for k in clients_all if client_segment[k]==seg_key]
    q1_vol_t = sum(client_q1_kg[k] for k in clients)/1000.0
    q2_vol_t = sum(client_q2_kg[k] for k in clients)/1000.0
    q1_ca = sum(client_q1_ca[k] for k in clients)
    q2_ca = sum(client_q2_ca[k] for k in clients)
    return {
        "n_clients": len(clients),
        "q1_vol_t": q1_vol_t,
        "q2_vol_t": q2_vol_t,
        "q1_ca": q1_ca,
        "q2_ca": q2_ca,
        "delta_vol_t": q2_vol_t - q1_vol_t,
        "delta_ca": q2_ca - q1_ca,
    }

stats_zz = seg_stats(("Zero Q1","Zero Q2"))
stats_za = seg_stats(("Zero Q1","Active Q2"))
stats_aa = seg_stats(("Active Q1","Active Q2"))
stats_az = seg_stats(("Active Q1","Zero Q2"))

print(f"\n{'='*70}")
print(f"STATS BY SEGMENT (volumes en tonnes, CA en FCFA — concentrés only)")
print(f"{'='*70}")
print(f"{'Segment':<40} {'N':>5} {'Vol Q1':>10} {'Vol Q2':>10} {'dVol':>10} {'CA Q1':>15} {'CA Q2':>15} {'dCA':>15}")
for label, st in [
    ("Q1 Zero → Q2 Zero (persistant)", stats_zz),
    ("Q1 Zero → Q2 Active (réactivé)", stats_za),
    ("Q1 Active → Q2 Active (retenu)", stats_aa),
    ("Q1 Active → Q2 Zero (churned)", stats_az),
]:
    print(f"{label:<40} {st['n_clients']:>5} {st['q1_vol_t']:>10.2f} {st['q2_vol_t']:>10.2f} "
          f"{st['delta_vol_t']:>+10.2f} {st['q1_ca']:>15,.0f} {st['q2_ca']:>15,.0f} {st['delta_ca']:>+15,.0f}")

gain_ca = stats_za["q2_ca"]
loss_ca = stats_az["q1_ca"]
gain_vol = stats_za["q2_vol_t"]
loss_vol = stats_az["q1_vol_t"]
net_ca = gain_ca - loss_ca
net_vol = gain_vol - loss_vol

print(f"\n{'='*70}")
print(f"BILAN NET Q1 -> Q2 (concentrés only)")
print(f"{'='*70}")
print(f"  GAIN (réactivés) : {gain_vol:>10.2f} t   {gain_ca:>15,.0f} FCFA")
print(f"  LOSS (churned)   : {loss_vol:>10.2f} t   {loss_ca:>15,.0f} FCFA")
print(f"  NET              : {net_vol:>+10.2f} t   {net_ca:>+15,.0f} FCFA")

# Fréquence
q1_active_total = sum(1 for k in clients_all if client_segment[k][0]=="Active Q1")
q2_active_total = sum(1 for k in clients_all if client_segment[k][1]=="Active Q2")
freq_dist_q1 = defaultdict(int)
for k in clients_all:
    if client_segment[k][0]=="Active Q1":
        freq_dist_q1[freq_cat(client_q1_months[k])] += 1
freq_dist_q2 = defaultdict(int)
for k in clients_all:
    if client_segment[k][1]=="Active Q2":
        freq_dist_q2[freq_cat(client_q2_months[k])] += 1

print(f"\nFRÉQUENCE Q1 (actifs Q1, n={q1_active_total}): Fidèle={freq_dist_q1['Fidèle']}, Semi-fidèle={freq_dist_q1['Semi-fidèle']}")
print(f"FRÉQUENCE Q2 (actifs Q2, n={q2_active_total}): Fidèle={freq_dist_q2['Fidèle']}, Semi-fidèle={freq_dist_q2['Semi-fidèle']}")

q1_fidele = [k for k in clients_all if freq_cat(client_q1_months[k])=="Fidèle"]
q1_fidele_to_q2_zero = sum(1 for k in q1_fidele if client_segment[k][1]=="Zero Q2")
q1_fidele_to_q2_active = sum(1 for k in q1_fidele if client_segment[k][1]=="Active Q2")
q1_fidele_to_q2_fidele = sum(1 for k in q1_fidele if freq_cat(client_q2_months[k])=="Fidèle")
q1_fidele_to_q2_semi = sum(1 for k in q1_fidele if freq_cat(client_q2_months[k])=="Semi-fidèle")

print(f"\nDestin des Q1 Fidèles (n={len(q1_fidele)}):")
print(f"  → Q2 Fidèle     : {q1_fidele_to_q2_fidele}")
print(f"  → Q2 Semi-fidèle: {q1_fidele_to_q2_semi}")
print(f"  → Q2 Zero       : {q1_fidele_to_q2_zero}")

# ---- Build output workbook ----
print(f"\n{'='*70}")
print(f"Building output: {OUT}")
print(f"{'='*70}")

# Styles
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F4E78")
TITLE_ALIGN = Alignment(horizontal="left", vertical="center")
SUB_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
BAND_FILL = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
BODY_FONT = Font(name="Calibri", size=10)
BOLD_FONT = Font(name="Calibri", size=10, bold=True)
BODY_ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
BODY_ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
BODY_ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
PARETO_FILL = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
PARETO_FONT = Font(name="Calibri", size=14, bold=True, color="7F6000")
PARETO_SYMBOL = "★"

SEG_COLORS = {
    ("Zero Q1","Zero Q2"):   "D9D9D9",
    ("Zero Q1","Active Q2"): "C6EFCE",
    ("Active Q1","Zero Q2"): "FCE4E4",
    ("Active Q1","Active Q2"): "DDEBF7",
}
SEG_FONTS = {
    ("Zero Q1","Zero Q2"):   Font(name="Calibri", size=10, color="595959"),
    ("Zero Q1","Active Q2"): Font(name="Calibri", size=10, color="006100", bold=True),
    ("Active Q1","Zero Q2"): Font(name="Calibri", size=10, color="C00000", bold=True),
    ("Active Q1","Active Q2"): Font(name="Calibri", size=10, color="1F4E78", bold=True),
}

CA_NUM_FMT = '#,##0" FCFA"'
VOL_NUM_FMT = '#,##0.00" t"'
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()
wb.properties.creator = "Z.ai"

# ===== Sheet 1: Synthèse =====
ws1 = wb.active
ws1.title = "Synthèse"

ws1.cell(row=1, column=1, value=f"BELGOCAM SA - Transition Q1 → Q2 2026 (CONCENTRÉS only : 12 produits BELGO 10% + BELGO 5%)").font = TITLE_FONT
ws1.cell(row=1, column=1).alignment = TITLE_ALIGN
ws1.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
ws1.row_dimensions[1].height = 28

ws1.cell(row=2, column=1, value=("Analyse du comportement des 1 394 clients entre Q1 (Jan-Mar) et Q2 (Avr-Jun) 2026 sur les CONCENTRÉS uniquement "
                                  "(10 BELGO 10% + 2 BELGO 5%, sans tourteaux de soja).")
).font = SUB_FONT
ws1.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)
ws1.row_dimensions[2].height = 30

# Section 1: Matrix
ws1.cell(row=4, column=1, value="1. Matrice de transition Q1 → Q2 (nombre de clients) — Concentrés").font = Font(name="Calibri", size=12, bold=True, color="1F4E78")
ws1.merge_cells(start_row=4, start_column=1, end_row=4, end_column=8)
ws1.row_dimensions[4].height = 22

c = ws1.cell(row=5, column=1, value=""); c.fill = HEADER_FILL; c.border = BORDER
c = ws1.cell(row=5, column=2, value="Q2 : Zero (pas d'achat concentré)")
c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN; c.border = BORDER
c = ws1.cell(row=5, column=3, value="Q2 : Active (≥1 achat concentré)")
c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN; c.border = BORDER
c = ws1.cell(row=5, column=4, value="Total Q1")
c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN; c.border = BORDER

c = ws1.cell(row=6, column=1, value="Q1 : Zero (pas d'achat concentré)")
c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = Alignment(horizontal="left", vertical="center"); c.border = BORDER
c = ws1.cell(row=6, column=2, value=zz); c.alignment = BODY_ALIGN_CENTER; c.border = BORDER
c.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"); c.font = BOLD_FONT
c = ws1.cell(row=6, column=3, value=za); c.alignment = BODY_ALIGN_CENTER; c.border = BORDER
c.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
c.font = Font(name="Calibri", size=11, bold=True, color="006100")
c = ws1.cell(row=6, column=4, value=zz+za); c.alignment = BODY_ALIGN_CENTER; c.border = BORDER; c.font = BOLD_FONT

c = ws1.cell(row=7, column=1, value="Q1 : Active (≥1 achat concentré)")
c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = Alignment(horizontal="left", vertical="center"); c.border = BORDER
c = ws1.cell(row=7, column=2, value=az); c.alignment = BODY_ALIGN_CENTER; c.border = BORDER
c.fill = PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid")
c.font = Font(name="Calibri", size=11, bold=True, color="C00000")
c = ws1.cell(row=7, column=3, value=aa); c.alignment = BODY_ALIGN_CENTER; c.border = BORDER
c.fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
c.font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
c = ws1.cell(row=7, column=4, value=az+aa); c.alignment = BODY_ALIGN_CENTER; c.border = BORDER; c.font = BOLD_FONT

c = ws1.cell(row=8, column=1, value="Total Q2")
c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = Alignment(horizontal="left", vertical="center"); c.border = BORDER
c = ws1.cell(row=8, column=2, value=zz+az); c.alignment = BODY_ALIGN_CENTER; c.border = BORDER; c.font = BOLD_FONT
c = ws1.cell(row=8, column=3, value=za+aa); c.alignment = BODY_ALIGN_CENTER; c.border = BORDER; c.font = BOLD_FONT
c = ws1.cell(row=8, column=4, value=len(clients_all)); c.alignment = BODY_ALIGN_CENTER; c.border = BORDER; c.font = BOLD_FONT

ws1.column_dimensions['A'].width = 35
ws1.column_dimensions['B'].width = 24
ws1.column_dimensions['C'].width = 24
ws1.column_dimensions['D'].width = 14

# Section 2: Interprétation
ws1.cell(row=10, column=1, value="2. Lecture des 4 segments").font = Font(name="Calibri", size=12, bold=True, color="1F4E78")
ws1.merge_cells(start_row=10, start_column=1, end_row=10, end_column=8)
ws1.row_dimensions[10].height = 22

interpretations = [
    ("Q1 Zero → Q2 Zero (persistant)", zz,
     "Clients qui n'ont JAMAIS acheté de concentrés sur les 6 mois. Peuvent avoir acheté du soja mais pas de concentrés. Action : prospection sur les concentrés."),
    ("Q1 Zero → Q2 Active (réactivé)", za,
     "Clients réactivés sur les concentrés en Q2. C'est un GAIN. Action : comprendre ce qui a déclenché l'achat Q2 et le reproduire."),
    ("Q1 Active → Q2 Active (retenu)", aa,
     "Clients qui achètent des concentrés en Q1 et Q2. C'est le cœur de clientèle concentrés. Action : fidélisation."),
    ("Q1 Active → Q2 Zero (churned)", az,
     "Clients qui ont arrêté d'acheter des concentrés en Q2. PERTE sur les concentrés. Action : recontact urgent."),
]
r = 11
for label, n, comment in interpretations:
    color = SEG_COLORS[
        ("Zero Q1","Zero Q2") if "persistant" in label else
        ("Zero Q1","Active Q2") if "réactivé" in label else
        ("Active Q1","Active Q2") if "retenu" in label else
        ("Active Q1","Zero Q2")
    ]
    c = ws1.cell(row=r, column=1, value=label); c.font = BOLD_FONT
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True); c.border = BORDER
    c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    c = ws1.cell(row=r, column=2, value=n); c.font = BOLD_FONT; c.alignment = BODY_ALIGN_CENTER; c.border = BORDER
    c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    c = ws1.cell(row=r, column=3, value=comment); c.font = BODY_FONT
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True); c.border = BORDER
    ws1.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    ws1.row_dimensions[r].height = 36
    r += 1

# Section 3: Volumes & CA par segment
r_title = r + 1
ws1.cell(row=r_title, column=1, value="3. Volumes (tonnes) et CA (FCFA) par segment — Concentrés uniquement").font = Font(name="Calibri", size=12, bold=True, color="1F4E78")
ws1.merge_cells(start_row=r_title, start_column=1, end_row=r_title, end_column=8)
ws1.row_dimensions[r_title].height = 22

r_h = r_title + 1
headers = ["Segment", "Nb clients", "Vol Q1 (t)", "Vol Q2 (t)", "Δ Vol (t)", "CA Q1 (FCFA)", "CA Q2 (FCFA)", "Δ CA (FCFA)"]
for col_idx, h in enumerate(headers, start=1):
    c = ws1.cell(row=r_h, column=col_idx, value=h)
    c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN; c.border = BORDER
ws1.row_dimensions[r_h].height = 32

seg_rows = [
    ("Q1 Zero → Q2 Zero (persistant)", stats_zz, ("Zero Q1","Zero Q2")),
    ("Q1 Zero → Q2 Active (réactivé)", stats_za, ("Zero Q1","Active Q2")),
    ("Q1 Active → Q2 Active (retenu)", stats_aa, ("Active Q1","Active Q2")),
    ("Q1 Active → Q2 Zero (churned)", stats_az, ("Active Q1","Zero Q2")),
]
r_data = r_h + 1
for label, st, seg_key in seg_rows:
    color = SEG_COLORS[seg_key]
    c = ws1.cell(row=r_data, column=1, value=label); c.font = BOLD_FONT
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True); c.border = BORDER
    c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    c = ws1.cell(row=r_data, column=2, value=st["n_clients"]); c.font = BODY_FONT
    c.alignment = BODY_ALIGN_CENTER; c.border = BORDER
    c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    c = ws1.cell(row=r_data, column=3, value=st["q1_vol_t"]); c.font = BODY_FONT
    c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER; c.number_format = VOL_NUM_FMT
    c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    c = ws1.cell(row=r_data, column=4, value=st["q2_vol_t"]); c.font = BODY_FONT
    c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER; c.number_format = VOL_NUM_FMT
    c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    c = ws1.cell(row=r_data, column=5, value=st["delta_vol_t"]); c.font = BOLD_FONT
    c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER
    c.number_format = '+#,##0.00" t";-#,##0.00" t";"-"'
    c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    c = ws1.cell(row=r_data, column=6, value=st["q1_ca"]); c.font = BODY_FONT
    c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER; c.number_format = CA_NUM_FMT
    c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    c = ws1.cell(row=r_data, column=7, value=st["q2_ca"]); c.font = BODY_FONT
    c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER; c.number_format = CA_NUM_FMT
    c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    c = ws1.cell(row=r_data, column=8, value=st["delta_ca"]); c.font = BOLD_FONT
    c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER
    c.number_format = '+#,##0" FCFA";-#,##0" FCFA";"-"'
    c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    r_data += 1

# Total row
total_q1_vol = sum(s["q1_vol_t"] for _, s, _ in seg_rows)
total_q2_vol = sum(s["q2_vol_t"] for _, s, _ in seg_rows)
total_q1_ca = sum(s["q1_ca"] for _, s, _ in seg_rows)
total_q2_ca = sum(s["q2_ca"] for _, s, _ in seg_rows)
c = ws1.cell(row=r_data, column=1, value="TOTAL"); c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
c.fill = HEADER_FILL; c.alignment = BODY_ALIGN_LEFT; c.border = BORDER
c = ws1.cell(row=r_data, column=2, value=len(clients_all)); c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
c.fill = HEADER_FILL; c.alignment = BODY_ALIGN_CENTER; c.border = BORDER
c = ws1.cell(row=r_data, column=3, value=total_q1_vol); c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
c.fill = HEADER_FILL; c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER; c.number_format = VOL_NUM_FMT
c = ws1.cell(row=r_data, column=4, value=total_q2_vol); c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
c.fill = HEADER_FILL; c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER; c.number_format = VOL_NUM_FMT
c = ws1.cell(row=r_data, column=5, value=total_q2_vol-total_q1_vol); c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
c.fill = HEADER_FILL; c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER
c.number_format = '+#,##0.00" t";-#,##0.00" t";"-"'
c = ws1.cell(row=r_data, column=6, value=total_q1_ca); c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
c.fill = HEADER_FILL; c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER; c.number_format = CA_NUM_FMT
c = ws1.cell(row=r_data, column=7, value=total_q2_ca); c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
c.fill = HEADER_FILL; c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER; c.number_format = CA_NUM_FMT
c = ws1.cell(row=r_data, column=8, value=total_q2_ca-total_q1_ca); c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
c.fill = HEADER_FILL; c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER
c.number_format = '+#,##0" FCFA";-#,##0" FCFA";"-"'

for col_letter, w in zip(['A','B','C','D','E','F','G','H'], [40, 11, 14, 14, 14, 18, 18, 18]):
    ws1.column_dimensions[col_letter].width = w

# Section 4: Bilan net
r_bilan = r_data + 2
ws1.cell(row=r_bilan, column=1, value="4. Bilan net Q1 → Q2 (Concentrés)").font = Font(name="Calibri", size=12, bold=True, color="1F4E78")
ws1.merge_cells(start_row=r_bilan, start_column=1, end_row=r_bilan, end_column=8)
ws1.row_dimensions[r_bilan].height = 22

r_bilan_h = r_bilan + 1
bilan_headers = ["", "Clients", "Volume (t)", "CA (FCFA)", "Commentaire"]
for col_idx, h in enumerate(bilan_headers, start=1):
    c = ws1.cell(row=r_bilan_h, column=col_idx, value=h)
    c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN; c.border = BORDER
ws1.merge_cells(start_row=r_bilan_h, start_column=5, end_row=r_bilan_h, end_column=8)
ws1.row_dimensions[r_bilan_h].height = 28

r_gain = r_bilan_h + 1
c = ws1.cell(row=r_gain, column=1, value="GAIN — Clients réactivés (Q1 Zero → Q2 Active)")
c.font = Font(name="Calibri", size=10, bold=True, color="006100")
c.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True); c.border = BORDER
c = ws1.cell(row=r_gain, column=2, value=za); c.font = BOLD_FONT; c.alignment = BODY_ALIGN_CENTER; c.border = BORDER
c.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
c = ws1.cell(row=r_gain, column=3, value=gain_vol); c.font = BOLD_FONT; c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER
c.number_format = VOL_NUM_FMT
c.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
c = ws1.cell(row=r_gain, column=4, value=gain_ca); c.font = BOLD_FONT; c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER
c.number_format = CA_NUM_FMT
c.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
c = ws1.cell(row=r_gain, column=5, value="Nouveaux achats de concentrés en Q2 par des clients qui n'en faisaient pas en Q1.")
c.font = BODY_FONT; c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True); c.border = BORDER
ws1.merge_cells(start_row=r_gain, start_column=5, end_row=r_gain, end_column=8)
ws1.row_dimensions[r_gain].height = 30

r_loss = r_gain + 1
c = ws1.cell(row=r_loss, column=1, value="PERTE — Clients churned (Q1 Active → Q2 Zero)")
c.font = Font(name="Calibri", size=10, bold=True, color="C00000")
c.fill = PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid")
c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True); c.border = BORDER
c = ws1.cell(row=r_loss, column=2, value=az); c.font = BOLD_FONT; c.alignment = BODY_ALIGN_CENTER; c.border = BORDER
c.fill = PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid")
c = ws1.cell(row=r_loss, column=3, value=loss_vol); c.font = BOLD_FONT; c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER
c.number_format = VOL_NUM_FMT
c.fill = PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid")
c = ws1.cell(row=r_loss, column=4, value=loss_ca); c.font = BOLD_FONT; c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER
c.number_format = CA_NUM_FMT
c.fill = PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid")
c = ws1.cell(row=r_loss, column=5, value="CA concentrés Q1 perdu car ces clients n'ont plus rien acheté de concentrés en Q2.")
c.font = BODY_FONT; c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True); c.border = BORDER
ws1.merge_cells(start_row=r_loss, start_column=5, end_row=r_loss, end_column=8)
ws1.row_dimensions[r_loss].height = 30

r_net = r_loss + 1
c = ws1.cell(row=r_net, column=1, value="BILAN NET")
c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
c.fill = HEADER_FILL; c.alignment = BODY_ALIGN_LEFT; c.border = BORDER
c = ws1.cell(row=r_net, column=2, value=za-az); c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
c.fill = HEADER_FILL; c.alignment = BODY_ALIGN_CENTER; c.border = BORDER
c = ws1.cell(row=r_net, column=3, value=net_vol); c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
c.fill = HEADER_FILL; c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER
c.number_format = '+#,##0.00" t";-#,##0.00" t";"-"'
c = ws1.cell(row=r_net, column=4, value=net_ca); c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
c.fill = HEADER_FILL; c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER
c.number_format = '+#,##0" FCFA";-#,##0" FCFA";"-"'
c = ws1.cell(row=r_net, column=5, value="Si positif : plus de clients réactivés que perdus sur les concentrés. Si négatif : dégradation nette.")
c.font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
c.fill = HEADER_FILL; c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True); c.border = BORDER
ws1.merge_cells(start_row=r_net, start_column=5, end_row=r_net, end_column=8)
ws1.row_dimensions[r_net].height = 30

# Section 5: Destin des Q1 Fidèles
r_freq = r_net + 2
ws1.cell(row=r_freq, column=1, value="5. Destin des clients fidèles Q1 (3 mois d'achat de concentrés en Q1)").font = Font(name="Calibri", size=12, bold=True, color="1F4E78")
ws1.merge_cells(start_row=r_freq, start_column=1, end_row=r_freq, end_column=8)
ws1.row_dimensions[r_freq].height = 22

r_freq_h = r_freq + 1
freq_headers = ["Statut Q2", "Nb clients", "% des Q1 Fidèles", "Commentaire"]
for col_idx, h in enumerate(freq_headers, start=1):
    c = ws1.cell(row=r_freq_h, column=col_idx, value=h)
    c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN; c.border = BORDER
ws1.merge_cells(start_row=r_freq_h, start_column=4, end_row=r_freq_h, end_column=8)
ws1.row_dimensions[r_freq_h].height = 28

r_freq_d = r_freq_h + 1
n_q1_fidele = len(q1_fidele)
freq_rows = [
    ("Q2 Fidèle (3 mois)", q1_fidele_to_q2_fidele, "Maintenu à un rythme mensuel — client stable sur les concentrés."),
    ("Q2 Semi-fidèle (1-2 mois)", q1_fidele_to_q2_semi, "Déclin de fréquence — risque d'attrition sur les concentrés."),
    ("Q2 Zero (churned)", q1_fidele_to_q2_zero, "PERTE SÈCHE — gros client fidèle devenu inactif sur les concentrés."),
]
for label, n, comment in freq_rows:
    color = "DDEBF7" if "Fidèle" in label and "Zero" not in label else "FFF2CC" if "Semi" in label else "FCE4E4"
    c = ws1.cell(row=r_freq_d, column=1, value=label); c.font = BOLD_FONT
    c.alignment = BODY_ALIGN_LEFT; c.border = BORDER
    c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    c = ws1.cell(row=r_freq_d, column=2, value=n); c.font = BOLD_FONT
    c.alignment = BODY_ALIGN_CENTER; c.border = BORDER
    c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    pct = (n/n_q1_fidele*100) if n_q1_fidele>0 else 0
    c = ws1.cell(row=r_freq_d, column=3, value=pct/100); c.font = BODY_FONT
    c.alignment = BODY_ALIGN_CENTER; c.border = BORDER; c.number_format = '0.0%'
    c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    c = ws1.cell(row=r_freq_d, column=4, value=comment); c.font = BODY_FONT
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True); c.border = BORDER
    ws1.merge_cells(start_row=r_freq_d, start_column=4, end_row=r_freq_d, end_column=8)
    c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    ws1.row_dimensions[r_freq_d].height = 24
    r_freq_d += 1

c = ws1.cell(row=r_freq_d, column=1, value="TOTAL Q1 Fidèles")
c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
c.fill = HEADER_FILL; c.alignment = BODY_ALIGN_LEFT; c.border = BORDER
c = ws1.cell(row=r_freq_d, column=2, value=n_q1_fidele)
c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
c.fill = HEADER_FILL; c.alignment = BODY_ALIGN_CENTER; c.border = BORDER
c = ws1.cell(row=r_freq_d, column=3, value=1.0)
c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
c.fill = HEADER_FILL; c.alignment = BODY_ALIGN_CENTER; c.border = BORDER; c.number_format = '0.0%'
c = ws1.cell(row=r_freq_d, column=4, value=""); c.fill = HEADER_FILL; c.border = BORDER
ws1.merge_cells(start_row=r_freq_d, start_column=4, end_row=r_freq_d, end_column=8)

# ===== Sheet 2: Détail clients =====
ws2 = wb.create_sheet("Détail clients")
ws2.cell(row=1, column=1, value="BELGOCAM SA - Détail par client : transition Q1 → Q2 sur CONCENTRÉS uniquement").font = TITLE_FONT
ws2.cell(row=1, column=1).alignment = TITLE_ALIGN
ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=13)
ws2.row_dimensions[1].height = 26

HEADERS = ["N°", "Réf. client", "Nom du client", "Segment Q1→Q2",
           "Mois actifs Q1", "Fréquence Q1", "Mois actifs Q2", "Fréquence Q2",
           "Vol Q1 (t)", "CA Q1 (FCFA)", "Vol Q2 (t)", "CA Q2 (FCFA)", "20/80"]
for col_idx, h in enumerate(HEADERS, start=1):
    c = ws2.cell(row=3, column=col_idx, value=h)
    c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN; c.border = BORDER
ws2.row_dimensions[3].height = 36

seg_order = {
    ("Active Q1","Zero Q2"): 0,
    ("Zero Q1","Active Q2"): 1,
    ("Active Q1","Active Q2"): 2,
    ("Zero Q1","Zero Q2"): 3,
}

def sort_key(k):
    seg = client_segment[k]
    return (seg_order[seg], -(client_q1_ca[k] + client_q2_ca[k]))

sorted_clients = sorted(clients_all, key=sort_key)

START_ROW = 4
for i, key in enumerate(sorted_clients, start=1):
    r = START_ROW + i - 1
    ref, name = key
    seg = client_segment[key]
    seg_label = SEGMENT_LABELS[seg]
    color = SEG_COLORS[seg]
    font = SEG_FONTS[seg]
    q1_months = sorted(client_q1_months[key])
    q2_months = sorted(client_q2_months[key])
    q1_freq = freq_cat(client_q1_months[key])
    q2_freq = freq_cat(client_q2_months[key])
    q1_vol_t = client_q1_kg[key]/1000.0
    q2_vol_t = client_q2_kg[key]/1000.0
    q1_ca = client_q1_ca[key]
    q2_ca = client_q2_ca[key]
    fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

    cells = [
        (1, i, BODY_ALIGN_CENTER, None),
        (2, ref, BODY_ALIGN_CENTER, None),
        (3, name, BODY_ALIGN_LEFT, None),
        (4, seg_label, Alignment(horizontal="left", vertical="center", wrap_text=True), None),
        (5, ", ".join(str(m) for m in q1_months) if q1_months else "-", BODY_ALIGN_CENTER, None),
        (6, q1_freq, BODY_ALIGN_CENTER, None),
        (7, ", ".join(str(m) for m in q2_months) if q2_months else "-", BODY_ALIGN_CENTER, None),
        (8, q2_freq, BODY_ALIGN_CENTER, None),
        (9, q1_vol_t, BODY_ALIGN_RIGHT, VOL_NUM_FMT),
        (10, q1_ca, BODY_ALIGN_RIGHT, CA_NUM_FMT),
        (11, q2_vol_t, BODY_ALIGN_RIGHT, VOL_NUM_FMT),
        (12, q2_ca, BODY_ALIGN_RIGHT, CA_NUM_FMT),
    ]
    for col_idx, val, align, fmt in cells:
        c = ws2.cell(row=r, column=col_idx, value=val)
        c.font = font; c.alignment = align; c.border = BORDER
        if fmt: c.number_format = fmt
        c.fill = fill
    c = ws2.cell(row=r, column=13)
    c.border = BORDER; c.alignment = BODY_ALIGN_CENTER
    if key in pareto_clients:
        c.value = PARETO_SYMBOL
        c.fill = PARETO_FILL
        c.font = PARETO_FONT
    else:
        c.fill = fill
        c.font = font

widths = [6, 16, 38, 38, 14, 14, 14, 14, 12, 18, 12, 18, 9]
for i, w in enumerate(widths, start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w

ws2.freeze_panes = "E4"
ws2.auto_filter.ref = f"A3:M{START_ROW + len(sorted_clients) - 1}"

wb.save(OUT)
print(f"\n  Saved: {OUT}")
print(f"  File size: {os.path.getsize(OUT):,} bytes")
