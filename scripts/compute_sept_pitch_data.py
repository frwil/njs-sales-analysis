"""Compute September pitch data (nationwide + per-region) for build_pitch_dg_pdf.py.
Based on sept_pitch_data_08.json structure — regenerated from the full-month extraction (51).xlsx.

Output: /home/z/my-project/scripts/sept_pitch_data_10.json
"""
import os
import json
from collections import defaultdict
from datetime import date, timedelta

import pandas as pd

SEPT_SRC = "/home/z/my-project/upload/NJS GROUP ERP - Lignes de commandes + multicompany (51).xlsx"

SOJA_REFS = {'T102': 50, 'T1021': 1, 'T1023': 5, 'T1024': 25}
CONC_REFS = {
    'C101': 50, 'C102': 50, 'C103': 50, 'C104': 50, 'C1042': 1, 'C1043': 5, 'C1044': 25,
    'C105': 50, 'C1053': 1, 'C1054': 5, 'C1055': 25, 'C108': 50
}

AGENCE_MAP = {
    'AGENCE FAMLA': ('Famla', 'Ouest'),
    'AGENCE MESSASSI': ('Messassi', 'Centre'),
    'AGENCE BERTOUA': ('Bertoua', 'Centre'),
    'AGENCE NDOBO': ('Ndobo', 'Littoral'),
    'AGENCE DJELENG': ('Djeleng', 'Ouest'),
    'AGENCE VILLAGE': ('Village', 'Littoral'),
    'AGENCE PK11': ('Pk11', 'Littoral'),
    'AGENCE NGAOUNDERE': ('Ngaoundere', 'Centre'),
    'AGENCE AHALA': ('Ahala', 'Centre'),
    'AGENCE NKONGSAMBA': ('Nkongsamba', 'Littoral'),
    'AGENCE NKOLBISSON': ('Nkolbisson', 'Centre'),
    'AGENCE NKOABANG': ('Nkoabang', 'Centre'),
    'AGENCE DE BAMENDA - DEPOT MBOUDA': ('Mbouda', 'Ouest'),
    'AGENCE BUEA': ('Buea', 'Littoral'),
    'AGENCE MAROUA': ('Maroua', 'Centre'),
    'SPC BAF-CHEFFERIE': ('Baf-Chefferie', 'Ouest'),
    'PDC Emana': ('Emana', 'Centre'),
    'SPC-NDERE': ('Ndere', 'Centre'),
    'SPC-DSCHANG': ('Dschang', 'Ouest'),
    'SPC BUEA': ('Buea-SPC', 'Littoral'),
    'SPC-YASSA': ('Yassa', 'Littoral'),
}

# Objectives September — S2 recalibrés (s2_recaled_objectives.json, mois 9)
OBJ = {'TOURTEAUX': 3781.3057056936536, 'CONCENTRES': 1601.3205305244583}

# Per-region objectives: S2 recalibrés mois 9 (somme des objectifs agences = objectif global)
REGION_OBJ = {
    'Ouest': {'TOURTEAUX': 1527.2507708563594, 'CONCENTRES': 631.6206108717748},
    'Centre': {'TOURTEAUX': 868.1861973111836, 'CONCENTRES': 512.7503113808139},
    'Littoral': {'TOURTEAUX': 1385.8687375261105, 'CONCENTRES': 456.9496082718696},
}


def load_livree(path):
    """Load 'Livrée' rows from the ERP extraction (.csv or .xlsx, same 16-col layout:
    0=ref produit, 2=qte, 3=ref commande, 6=date, 13=état, 15=agence)."""
    if path.endswith('.csv'):
        df = pd.read_csv(path, encoding='latin-1', sep=';', header=1,
                         on_bad_lines='skip', engine='python')
        ncols = len(df.columns)
        col_etat = 13 if ncols <= 16 else 15
        col_agence = 15 if ncols <= 16 else 17
        rows = []
        for t in df.itertuples(index=False, name=None):
            r = list(t)
            if not r or len(r) < max(col_etat, col_agence) + 1:
                continue
            if str(r[0] or '').strip().upper().startswith('TOTAL'):
                continue
            if str(r[col_etat] or '').strip() != 'Livrée':
                continue
            try:
                r[2] = float(str(r[2]).replace(' ', '').replace('\xa0', '').replace(',', '.'))
            except (ValueError, TypeError):
                r[2] = 0.0
            rows.append((r, col_agence))
        return rows

    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['Sheet 1']
    header = None
    for row in ws.iter_rows(min_row=2, max_row=2, values_only=True):
        header = row
        break
    col_etat = 13 if len(header) <= 16 else 15
    col_agence = 15 if len(header) <= 16 else 17
    for i, h in enumerate(header):
        if h == 'État':
            col_etat = i
        elif h == 'agence':
            col_agence = i
    rows = []
    for t in ws.iter_rows(min_row=3, values_only=True):
        r = list(t)
        if not r or len(r) < max(col_etat, col_agence) + 1:
            continue
        if str(r[0] or '').strip().upper().startswith('TOTAL') or r[0] == 'Total':
            continue
        if str(r[col_etat] or '').strip() != 'Livrée':
            continue
        try:
            r[2] = float(str(r[2]).replace(' ', '').replace('\xa0', '').replace(',', '.'))
        except (ValueError, TypeError):
            r[2] = 0.0
        rows.append((r, col_agence))
    return rows


def parse_date_str(d):
    s = str(d)[:10]
    return s


print(f"Loading September extraction: {os.path.basename(SEPT_SRC)}...")
rows = load_livree(SEPT_SRC)
print(f"  {len(rows)} Livrée rows total")

# Latest date
latest_date = None
for r, _ in rows:
    ds = parse_date_str(r[6])
    if '/09/2026' in ds:
        try:
            d = date(int(ds[6:10]), int(ds[3:5]), int(ds[:2]))
            if latest_date is None or d > latest_date:
                latest_date = d
        except Exception:
            pass
print(f"  Latest Livrée date: {latest_date}")

days_elapsed = 0
for off in range((latest_date - date(2026, 9, 1)).days + 1):
    dt = date(2026, 9, 1) + timedelta(days=off)
    if dt.weekday() < 6:
        days_elapsed += 1
total_days_sep = 0
for off in range(30):
    dt = date(2026, 9, 1) + timedelta(days=off)
    if dt.month != 9:
        break
    if dt.weekday() < 6:
        total_days_sep += 1
print(f"  Days elapsed: {days_elapsed}/{total_days_sep} ({days_elapsed/total_days_sep*100:.0f}%)")

# Aggregates: (region, cat) -> kg ; commands per region
vol = defaultdict(lambda: defaultdict(float))          # region -> cat -> kg
sacs = defaultdict(lambda: defaultdict(float))         # region -> cat -> sacs
cmds_soja_r = defaultdict(set)
cmds_conc_r = defaultdict(set)
cmds_soja_only_r = defaultdict(set)

for r, col_agence in rows:
    ref = str(r[0] or '').strip().upper()
    qte = r[2]
    ds = parse_date_str(r[6])
    if '/09/2026' not in ds:
        continue
    agence_raw = str(r[col_agence] or '') if len(r) > col_agence else ''
    _, region = AGENCE_MAP.get(agence_raw, (agence_raw, '?'))
    cmd_ref = str(r[3] or '').strip() if len(r) > 3 else ''

    if ref in SOJA_REFS:
        kg = qte * SOJA_REFS[ref]
        vol[region]['TOURTEAUX'] += kg
        sacs[region]['TOURTEAUX'] += kg / 50
        if cmd_ref:
            cmds_soja_r[region].add(cmd_ref)
    elif ref in CONC_REFS:
        kg = qte * CONC_REFS[ref]
        vol[region]['CONCENTRES'] += kg
        sacs[region]['CONCENTRES'] += kg / 50
        if cmd_ref:
            cmds_conc_r[region].add(cmd_ref)

# Nationwide
nat = {}
nat['soja_t_mtd'] = sum(vol[rg]['TOURTEAUX'] for rg in vol) / 1000
nat['conc_t_mtd'] = sum(vol[rg]['CONCENTRES'] for rg in vol) / 1000
nat['soja_sacs_mtd'] = sum(sacs[rg]['TOURTEAUX'] for rg in sacs)
nat['conc_sacs_mtd'] = sum(sacs[rg]['CONCENTRES'] for rg in sacs)
nat['soja_proj_t'] = nat['soja_t_mtd'] / days_elapsed * total_days_sep
nat['conc_proj_t'] = nat['conc_t_mtd'] / days_elapsed * total_days_sep
nat['soja_obj_t'] = OBJ['TOURTEAUX']
nat['conc_obj_t'] = OBJ['CONCENTRES']
nat['soja_pct_obj'] = nat['soja_proj_t'] / OBJ['TOURTEAUX'] * 100
nat['conc_pct_obj'] = nat['conc_proj_t'] / OBJ['CONCENTRES'] * 100
nat['soja_moy_t_j'] = nat['soja_t_mtd'] / days_elapsed
nat['conc_moy_t_j'] = nat['conc_t_mtd'] / days_elapsed
nat['ratio_global'] = nat['soja_sacs_mtd'] / nat['conc_sacs_mtd'] if nat['conc_sacs_mtd'] else 0
nat['pct_soja_volume'] = nat['soja_t_mtd'] / (nat['soja_t_mtd'] + nat['conc_t_mtd']) * 100
nat['pct_conc_volume'] = nat['conc_t_mtd'] / (nat['soja_t_mtd'] + nat['conc_t_mtd']) * 100

# Per region
regions = {}
for region in ['Ouest', 'Centre', 'Littoral']:
    rg = {}
    rg['soja_t_mtd'] = vol[region]['TOURTEAUX'] / 1000
    rg['conc_t_mtd'] = vol[region]['CONCENTRES'] / 1000
    rg['soja_sacs_mtd'] = sacs[region]['TOURTEAUX']
    rg['conc_sacs_mtd'] = sacs[region]['CONCENTRES']
    rg['soja_proj_t'] = rg['soja_t_mtd'] / days_elapsed * total_days_sep
    rg['conc_proj_t'] = rg['conc_t_mtd'] / days_elapsed * total_days_sep
    rg['soja_obj_t'] = REGION_OBJ[region]['TOURTEAUX']
    rg['conc_obj_t'] = REGION_OBJ[region]['CONCENTRES']
    rg['soja_pct_obj'] = rg['soja_proj_t'] / rg['soja_obj_t'] * 100
    rg['conc_pct_obj'] = rg['conc_proj_t'] / rg['conc_obj_t'] * 100
    rg['soja_moy_t_j'] = rg['soja_t_mtd'] / days_elapsed
    rg['conc_moy_t_j'] = rg['conc_t_mtd'] / days_elapsed
    rg['ratio'] = rg['soja_sacs_mtd'] / rg['conc_sacs_mtd'] if rg['conc_sacs_mtd'] else 0
    bundle = cmds_soja_r[region] & cmds_conc_r[region]
    soja_only = cmds_soja_r[region] - cmds_conc_r[region]
    rg['cmds_bundle'] = len(bundle)
    rg['cmds_soja_only'] = len(soja_only)
    rg['pct_bundle'] = len(bundle) / (len(bundle) + len(soja_only)) * 100 if (len(bundle) + len(soja_only)) else 0
    regions[region] = rg

out = {
    'nationwide': nat,
    'regions': regions,
    'days_elapsed': days_elapsed,
    'total_days_sep': total_days_sep,
    'pct_elapsed': round(days_elapsed / total_days_sep * 100, 1),
}

OUT = '/home/z/my-project/scripts/sept_pitch_data_10.json'
with open(OUT, 'w', encoding='utf-8') as f:
    json.dump(out, f, indent=2, ensure_ascii=False, default=str)
print(f"\n=== SAVED: {OUT} ===")
print(f"  Nat: soja {nat['soja_t_mtd']:.0f} t (proj {nat['soja_proj_t']:.0f} t = {nat['soja_pct_obj']:.0f}%), "
      f"conc {nat['conc_t_mtd']:.0f} t (proj {nat['conc_proj_t']:.0f} t = {nat['conc_pct_obj']:.0f}%), ratio {nat['ratio_global']:.2f}:1")
for region in ['Ouest', 'Centre', 'Littoral']:
    rg = regions[region]
    print(f"  {region:10s}: soja {rg['soja_t_mtd']:.0f} t ({rg['soja_pct_obj']:.0f}% obj), "
          f"conc {rg['conc_t_mtd']:.0f} t ({rg['conc_pct_obj']:.0f}% obj), ratio {rg['ratio']:.2f}:1, "
          f"bundle {rg['cmds_bundle']}/{rg['cmds_bundle']+rg['cmds_soja_only']}")
