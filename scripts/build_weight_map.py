"""
Build a product -> unit weight (kg) mapping by parsing product descriptions.

Most products have a weight indicator at the end like "50 KG", "5Kg", "1KG",
"25 Kg", "200GRAMMES", "1L" (liquid ≈ 1 kg/L).
For non-feed products (equipment, manuals, fees, "MAIS" without unit, etc.),
weight = 0.

Output: prints the mapping for review.
"""
import re
from openpyxl import load_workbook
from collections import defaultdict

SRC = "/home/z/my-project/download/ventes_livrees.xlsx"

# Regex: capture weight + unit anywhere in the description
# Units: KG, Kg, kg, KG, G, GRAMMES, L
WEIGHT_RE = re.compile(
    r"(\d+(?:\.\d+)?)\s*(KG|Kg|kg|KG|GRAMMES?|G|L)\b",
    re.IGNORECASE
)

def parse_weight_kg(desc):
    """Return weight in kg, or 0 if no weight found."""
    if not desc:
        return 0.0
    s = str(desc)
    matches = WEIGHT_RE.findall(s)
    if not matches:
        return 0.0
    # Use the LAST match (the weight is typically at the end of the description)
    val_str, unit = matches[-1]
    val = float(val_str)
    unit_upper = unit.upper()
    if unit_upper in ("KG",):
        return val
    if unit_upper in ("G", "GRAMME", "GRAMMES"):
        return val / 1000.0
    if unit_upper == "L":
        return val  # 1L ≈ 1 kg
    return 0.0

# Read source file and collect unique products
wb = load_workbook(SRC, read_only=True, data_only=True)
products = {}  # ref -> (desc, weight_kg)
ref_counts = defaultdict(int)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for row in ws.iter_rows(min_row=3, values_only=True):
        ref = row[0] if len(row) > 0 else None
        desc = row[1] if len(row) > 1 else None
        if ref is None:
            continue
        ref_str = str(ref).strip()
        if ref_str not in products:
            w = parse_weight_kg(desc)
            products[ref_str] = (desc, w)
        ref_counts[ref_str] += 1

wb.close()

# Print sorted by reference
print(f"{'Réf':<12} {'Poids (kg)':>10}  {'Lignes':>7}  Description")
print("-" * 100)
for ref in sorted(products.keys()):
    desc, w = products[ref]
    print(f"{ref:<12} {w:>10.2f}  {ref_counts[ref]:>7}  {str(desc)[:60]}")

# Stats
total_products = len(products)
with_weight = sum(1 for _, w in products.values() if w > 0)
without_weight = sum(1 for _, w in products.values() if w == 0)
print(f"\nTotal products: {total_products}")
print(f"  With parsed weight: {with_weight}")
print(f"  Without weight (0 kg): {without_weight}")
