"""Inspect the 'agence' column (column R, index 17) in source file."""
from openpyxl import load_workbook
from collections import Counter, defaultdict
import json

SRC = "/home/z/my-project/download/ventes_livrees.xlsx"

# Load excluded clients
with open("/home/z/my-project/scripts/excluded_clients.json", "r", encoding="utf-8") as f:
    excluded_tiers = set(json.load(f))

wb = load_workbook(SRC, read_only=True, data_only=True)
agency_counter = Counter()
agency_ca = defaultdict(float)
total_rows = 0
empty_agency_rows = 0

# Map client -> set of agencies (and CA per agency)
client_agencies = defaultdict(lambda: defaultdict(float))  # key -> {agency: ca}

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for row in ws.iter_rows(min_row=3, values_only=True):
        tiers = row[5] if len(row) > 5 else None
        if tiers is None or str(tiers).strip() == "":
            continue
        if tiers in excluded_tiers:
            continue
        ca = row[8] if len(row) > 8 else 0
        agence = row[17] if len(row) > 17 else None  # column R = index 17
        total_rows += 1
        if agence is None or str(agence).strip() == "":
            empty_agency_rows += 1
            agence_str = "(vide)"
        else:
            agence_str = str(agence).strip()
        try:
            ca_f = float(ca) if ca is not None else 0.0
        except:
            ca_f = 0.0
        agency_counter[agence_str] += 1
        agency_ca[agence_str] += ca_f
        client_agencies[tiers][agence_str] += ca_f

wb.close()

print(f"Total rows (after exclusion): {total_rows}")
print(f"Rows with empty agency: {empty_agency_rows}")
print(f"\nDistinct agencies: {len(agency_counter)}")
print(f"\n{'Agence':<40} {'Lignes':>8} {'CA HT':>15}")
print("-" * 70)
for ag, n in agency_counter.most_common():
    print(f"{ag[:38]:<40} {n:>8} {agency_ca[ag]:>15,.0f}")

# Distribution of number of agencies per client
from collections import Counter as C
n_agencies_per_client = C(len(v) for v in client_agencies.values())
print(f"\nDistribution of agencies per client:")
for n, c in sorted(n_agencies_per_client.items()):
    print(f"  {n} agence(s): {c} clients")
