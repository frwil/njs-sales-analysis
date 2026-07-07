"""Inspect the objectives file in detail."""
from openpyxl import load_workbook

SRC = "/home/z/my-project/upload/objectifs 2026 - Takou.xlsx"

wb = load_workbook(SRC, data_only=False)
ws = wb.active
print(f"Sheet: {ws.title}")
print(f"Dimensions: {ws.dimensions}")
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")

print("\n" + "="*100)
print("FIRST 15 ROWS (with formulas where applicable):")
print("="*100)
for r in range(1, 16):
    print(f"\n--- Row {r} ---")
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row=r, column=c)
        col_letter = chr(ord('A') + c - 1) if c <= 26 else f"Col{c}"
        if cell.value is not None:
            val_repr = repr(cell.value)[:80]
            print(f"  {col_letter}{r}: {val_repr}")

# Now load with data_only=True to see computed values
print("\n" + "="*100)
print("FIRST 15 ROWS (with computed values):")
print("="*100)
wb2 = load_workbook(SRC, data_only=True)
ws2 = wb2.active
for r in range(1, 16):
    print(f"\n--- Row {r} ---")
    for c in range(1, ws2.max_column + 1):
        cell = ws2.cell(row=r, column=c)
        col_letter = chr(ord('A') + c - 1) if c <= 26 else f"Col{c}"
        if cell.value is not None:
            val_repr = repr(cell.value)[:80]
            print(f"  {col_letter}{r}: {val_repr}")
