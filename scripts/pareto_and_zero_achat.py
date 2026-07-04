"""
20/80 Pareto marking + Zero-achat (Jan-Mar) detection on targeted products.

Targeted products (16 total):
- Tourteaux de soja: T102, T1021, T1023, T1024
- BELGO 10%: C102, C1022, C104, C1042, C1043, C1044, C105, C1053, C1054, C1055
- BELGO 5%:  C101, C103

Outputs:
1. Updates the existing "Clients uniques - check mensuel.xlsx" by adding:
   - A "Catégorie 20/80" column (★ mark for top clients contributing ~80% of CA HT)
   - A "CA Total HT (6 mois)" column for transparency
2. Creates "Clients zero achat - Janv-Mars.xlsx" listing all clients who
   did NOT buy any targeted product in Jan/Feb/Mar.
"""
import re
import os
import time
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

# ---- Paths ----
SRC = "/home/z/my-project/download/ventes janv a juin 2026 - Livrées uniquement.xlsx"
CHECK_FILE = "/home/z/my-project/download/Clients uniques - check mensuel.xlsx"
ZERO_ACHAT_FILE = "/home/z/my-project/download/Clients zero achat - Janv-Mars.xlsx"

# ---- Constants ----
TARGET_REFS = {
    # Tourteaux de soja
    "T102", "T1021", "T1023", "T1024",
    # BELGO 10%
    "C102", "C1022", "C104", "C1042", "C1043", "C1044",
    "C105", "C1053", "C1054", "C1055",
    # BELGO 5%
    "C101", "C103",
}
assert len(TARGET_REFS) == 16

MONTHS = [
    ("Sheet 1", "Janvier", 0),
    ("Feuil1",  "Février", 1),
    ("Feuil2",  "Mars",    2),
    ("Feuil3",  "Avril",   3),
    ("Feuil4",  "Mai",     4),
    ("Feuil5",  "Juin",    5),
]
FIRST_QUARTER_SHEETS = {"Sheet 1", "Feuil1", "Feuil2"}

TIERS_RE = re.compile(r"^\s*([A-Z]{2,}[\w-]*?)\s*-\s+(.+?)\s*$")

def split_tiers(tiers):
    if tiers is None:
        return ("", "")
    s = str(tiers).strip()
    m = TIERS_RE.match(s)
    if m:
        return (m.group(1).strip(), m.group(2).strip())
    return ("", s)

# ---- 1. Read source & aggregate per-client data ---------------------------
print(f"Reading source: {SRC}")
t0 = time.time()
wb_src = load_workbook(SRC, read_only=True, data_only=True)
print(f"  Loaded in {time.time()-t0:.1f}s")

# Per-client aggregates
client_ca_total = defaultdict(float)               # (ref, name) -> total CA HT (6 months)
client_ca_q1 = defaultdict(float)                  # (ref, name) -> CA HT Jan-Mar (all products)
client_targeted_q1 = defaultdict(bool)             # (ref, name) -> True if bought >=1 targeted product in Jan-Mar
client_targeted_any = defaultdict(bool)            # (ref, name) -> True if bought >=1 targeted product in any month
client_months = defaultdict(set)                   # for verification: month indices purchased (any product)

for sheet_name, month_label, m_idx in MONTHS:
    ws = wb_src[sheet_name]
    n_rows = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        n_rows += 1
        ref_prod = row[0] if len(row) > 0 else None      # A: Réf. produit
        tiers    = row[5] if len(row) > 5 else None      # F: Tiers
        ca_ht    = row[8] if len(row) > 8 else 0          # I: Montant HT

        if tiers is None:
            continue
        ref_c, name_c = split_tiers(tiers)
        key = (ref_c, name_c)

        # Skip non-numeric CA
        try:
            ca_ht_f = float(ca_ht) if ca_ht is not None else 0.0
        except (ValueError, TypeError):
            ca_ht_f = 0.0

        client_ca_total[key] += ca_ht_f
        client_months[key].add(m_idx)

        if sheet_name in FIRST_QUARTER_SHEETS:
            client_ca_q1[key] += ca_ht_f
            if ref_prod is not None and str(ref_prod).strip() in TARGET_REFS:
                client_targeted_q1[key] = True

        if ref_prod is not None and str(ref_prod).strip() in TARGET_REFS:
            client_targeted_any[key] = True

    print(f"  {month_label:<10} ({sheet_name}): {n_rows} data rows processed")

wb_src.close()

clients_all = sorted(client_ca_total.keys(), key=lambda k: (k[1].upper(), k[0]))
print(f"\nTotal unique clients: {len(clients_all)}")
print(f"Total CA HT (6 months, all clients): {sum(client_ca_total.values()):,.0f} FCFA")

# ---- 2. Pareto 20/80 -----------------------------------------------------
# Sort clients by CA HT descending
sorted_by_ca = sorted(clients_all, key=lambda k: -client_ca_total[k])
total_ca = sum(client_ca_total.values())
pareto_threshold = 0.80 * total_ca

# Walk down the list and stop when cumulative CA reaches 80% of total
pareto_clients = set()
cumulative = 0.0
for key in sorted_by_ca:
    cumulative += client_ca_total[key]
    pareto_clients.add(key)
    if cumulative >= pareto_threshold:
        break

pareto_count = len(pareto_clients)
pareto_pct_clients = pareto_count / len(clients_all) * 100
pareto_ca = sum(client_ca_total[k] for k in pareto_clients)
pareto_pct_ca = pareto_ca / total_ca * 100

print(f"\n{'='*70}")
print(f"PARETO 20/80")
print(f"{'='*70}")
print(f"Total clients              : {len(clients_all)}")
print(f"Total CA HT                : {total_ca:,.0f} FCFA")
print(f"Pareto clients (top)       : {pareto_count} ({pareto_pct_clients:.1f}% of clients)")
print(f"CA HT generated by them    : {pareto_ca:,.0f} FCFA ({pareto_pct_ca:.1f}% of total CA)")
print(f"Cumulative threshold (80%) : {pareto_threshold:,.0f} FCFA")

# ---- 3. Update the existing check file with 20/80 column ------------------
print(f"\n{'='*70}")
print(f"Updating check file: {CHECK_FILE}")
print(f"{'='*70}")

wb_check = load_workbook(CHECK_FILE)
ws = wb_check.active

# Current structure (before update):
#   Row 1: Title (merged A1:J1)
#   Row 2: spacer
#   Row 3: Headers A..J (N°, Réf. client, Nom du client, Jan..Juin, Nb mois d'achat)
#   Row 4+: data

# We will:
# - Insert a new column K "CA Total HT (6 mois)"
# - Insert a new column L "20/80" with star marker for Pareto clients
# - Re-merge title row to cover A1:L1
# - Update auto_filter range

# Unmerge existing title
for mr in list(ws.merged_cells.ranges):
    if mr.min_row == 1 and mr.max_row == 1:
        ws.unmerge_cells(str(mr))

# ---- Styles ----
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
BAND_FILL = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
BODY_FONT = Font(name="Calibri", size=10)
BODY_ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
BODY_ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
BODY_ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

# Pareto (20/80) styles - gold star
PARETO_FILL = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
PARETO_FONT = Font(name="Calibri", size=14, bold=True, color="7F6000")
PARETO_SYMBOL = "★"

# CA total style
CA_FONT = Font(name="Calibri", size=10)
CA_NUM_FMT = '#,##0" FCFA"'

THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Add new headers in K (CA HT) and L (20/80)
c = ws.cell(row=3, column=11, value="CA Total HT (6 mois)")
c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN; c.border = BORDER

c = ws.cell(row=3, column=12, value="Catégorie 20/80")
c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = HEADER_ALIGN; c.border = BORDER

# Fill data for new columns
START_ROW = 4
n_clients = len(clients_all)
for i in range(n_clients):
    r = START_ROW + i
    ref = ws.cell(row=r, column=2).value
    name = ws.cell(row=r, column=3).value
    key = (ref, name)
    banding = (i % 2 == 0)

    # Column K: CA Total HT
    ca_val = client_ca_total.get(key, 0.0)
    c = ws.cell(row=r, column=11, value=ca_val)
    c.font = CA_FONT
    c.alignment = BODY_ALIGN_RIGHT
    c.border = BORDER
    c.number_format = CA_NUM_FMT
    if banding: c.fill = BAND_FILL

    # Column L: 20/80 marker
    c = ws.cell(row=r, column=12)
    c.border = BORDER
    c.alignment = BODY_ALIGN_CENTER
    if key in pareto_clients:
        c.value = PARETO_SYMBOL
        c.fill = PARETO_FILL
        c.font = PARETO_FONT
    else:
        if banding:
            c.fill = BAND_FILL
        c.font = BODY_FONT

# Column widths for new columns
ws.column_dimensions['K'].width = 22
ws.column_dimensions['L'].width = 14

# Re-merge title across A1:L1
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=12)

# Update auto filter range
ws.auto_filter.ref = f"A3:L{START_ROW + n_clients - 1}"

# Save
wb_check.save(CHECK_FILE)
wb_check.close()
print(f"  Updated: {CHECK_FILE}")

# ---- 4. Build zero-achat (Jan-Mar) file ---------------------------------
print(f"\n{'='*70}")
print(f"Building zero-achat file: {ZERO_ACHAT_FILE}")
print(f"{'='*70}")

# Zero-achat = clients who did NOT buy any targeted product in Jan-Mar
# (whether or not they bought other products in those months)
zero_achat_clients = [k for k in clients_all if not client_targeted_q1[k]]

# Sort: by name (case-insensitive), then ref
zero_achat_clients.sort(key=lambda k: (k[1].upper(), k[0]))

print(f"Total clients              : {len(clients_all)}")
print(f"Clients with targeted Q1   : {sum(1 for k in clients_all if client_targeted_q1[k])}")
print(f"Clients 'zero achat' Q1    : {len(zero_achat_clients)}")

# Build output workbook
wb = Workbook()
wb.properties.creator = "Z.ai"
ws = wb.active
ws.title = "Zero achat Q1"

# Title
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F4E78")
TITLE_ALIGN = Alignment(horizontal="left", vertical="center")
ws.cell(row=1, column=1,
        value=("BELGOCAM SA - Clients 'zéro achat' Janvier-Mars 2026 "
               "(Tourteaux de soja + BELGO 10% + BELGO 5%)")
).font = TITLE_FONT
ws.cell(row=1, column=1).alignment = TITLE_ALIGN
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
ws.row_dimensions[1].height = 26

# Spacer
ws.row_dimensions[2].height = 6

# Subtitle with criteria explanation
SUB_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
ws.cell(row=2, column=1,
        value=("Critère : client n'ayant acheté AUCUN des 16 produits ciblés "
               "(4 tourteaux de soja + 10 BELGO 10% + 2 BELGO 5%) en janvier, février ou mars 2026."
        )).font = SUB_FONT
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)
ws.row_dimensions[2].height = 22

# Row 3: Headers
HEADERS = ["N°", "Réf. client", "Nom du client",
           "Achat ciblé Q1 ?", "CA HT Q1 (tous produits)",
           "Achat ciblé 6 mois ?", "CA HT total 6 mois", "Catégorie 20/80"]
for col_idx, h in enumerate(HEADERS, start=1):
    c = ws.cell(row=3, column=col_idx, value=h)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = HEADER_ALIGN
    c.border = BORDER
ws.row_dimensions[3].height = 36

# Data rows
for i, key in enumerate(zero_achat_clients, start=1):
    r = 3 + i
    ref, name = key
    banding = (i % 2 == 0)

    # N°
    c = ws.cell(row=r, column=1, value=i)
    c.font = BODY_FONT; c.alignment = BODY_ALIGN_CENTER; c.border = BORDER
    if banding: c.fill = BAND_FILL

    # Réf
    c = ws.cell(row=r, column=2, value=ref)
    c.font = BODY_FONT; c.alignment = BODY_ALIGN_CENTER; c.border = BORDER
    if banding: c.fill = BAND_FILL

    # Nom
    c = ws.cell(row=r, column=3, value=name)
    c.font = BODY_FONT; c.alignment = BODY_ALIGN_LEFT; c.border = BORDER
    if banding: c.fill = BAND_FILL

    # Achat ciblé Q1 ?
    c = ws.cell(row=r, column=4, value="Non")
    c.font = Font(name="Calibri", size=10, color="C00000", bold=True)
    c.alignment = BODY_ALIGN_CENTER; c.border = BORDER
    if banding: c.fill = BAND_FILL

    # CA HT Q1 (tous produits)
    c = ws.cell(row=r, column=5, value=client_ca_q1.get(key, 0.0))
    c.font = BODY_FONT; c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER
    c.number_format = CA_NUM_FMT
    if banding: c.fill = BAND_FILL

    # Achat ciblé 6 mois ?
    targeted_any = client_targeted_any.get(key, False)
    c = ws.cell(row=r, column=6, value=("Oui" if targeted_any else "Non"))
    if targeted_any:
        c.font = Font(name="Calibri", size=10, color="006100", bold=True)
    else:
        c.font = Font(name="Calibri", size=10, color="C00000", bold=True)
    c.alignment = BODY_ALIGN_CENTER; c.border = BORDER
    if banding: c.fill = BAND_FILL

    # CA HT total 6 mois
    c = ws.cell(row=r, column=7, value=client_ca_total.get(key, 0.0))
    c.font = BODY_FONT; c.alignment = BODY_ALIGN_RIGHT; c.border = BORDER
    c.number_format = CA_NUM_FMT
    if banding: c.fill = BAND_FILL

    # 20/80 marker
    c = ws.cell(row=r, column=8)
    c.border = BORDER
    c.alignment = BODY_ALIGN_CENTER
    if key in pareto_clients:
        c.value = PARETO_SYMBOL
        c.fill = PARETO_FILL
        c.font = PARETO_FONT
    else:
        if banding: c.fill = BAND_FILL
        c.font = BODY_FONT

# Column widths
ws.column_dimensions['A'].width = 6
ws.column_dimensions['B'].width = 16
ws.column_dimensions['C'].width = 48
ws.column_dimensions['D'].width = 16
ws.column_dimensions['E'].width = 22
ws.column_dimensions['F'].width = 18
ws.column_dimensions['G'].width = 22
ws.column_dimensions['H'].width = 14

# Freeze panes
ws.freeze_panes = "D4"

# Auto filter
ws.auto_filter.ref = f"A3:H{3 + len(zero_achat_clients)}"

# Save
wb.save(ZERO_ACHAT_FILE)
print(f"  Saved: {ZERO_ACHAT_FILE}")
print(f"  File size: {os.path.getsize(ZERO_ACHAT_FILE):,} bytes")

# ---- Summary stats ------------------------------------------------------
print(f"\n{'='*70}")
print(f"ZERO-ACHAT BREAKDOWN")
print(f"{'='*70}")
targeted_any_count = sum(1 for k in zero_achat_clients if client_targeted_any[k])
never_targeted = sum(1 for k in zero_achat_clients if not client_targeted_any[k])
pareto_in_zero = sum(1 for k in zero_achat_clients if k in pareto_clients)

print(f"  Total zero-achat clients (Jan-Mar)        : {len(zero_achat_clients)}")
print(f"    - of which bought targeted AFTER March  : {targeted_any_count}  (i.e. acquired later in Q2)")
print(f"    - of which NEVER bought any targeted    : {never_targeted}")
print(f"    - of which are in 20/80 (top clients)   : {pareto_in_zero}")
print(f"  Clients WITH targeted purchase in Q1      : {len(clients_all) - len(zero_achat_clients)}")
