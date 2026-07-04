"""Verify the cleaned file: confirm all data rows have État='Livrée'."""
from openpyxl import load_workbook
from collections import Counter

DST = "/home/z/my-project/download/ventes janv a juin 2026 - Livrées uniquement.xlsx"

wb = load_workbook(DST, read_only=True, data_only=True)

print("="*70)
print("VERIFICATION OF CLEANED FILE")
print("="*70)

total_data = 0
all_good = True
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    counter = Counter()
    # Check row 1 (title) and row 2 (headers) are preserved
    row1 = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    row2 = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    print(f"\n--- Sheet: {sheet_name} ---")
    print(f"  Row 1 (title)  : {str(row1[0])[:80]}")
    print(f"  Row 2 (header A): {row2[0]} | (header P): {row2[15]}")

    for row in ws.iter_rows(min_row=3, values_only=True):
        etat = row[15] if len(row) > 15 else None
        counter[etat] += 1
        total_data += 1

    print(f"  État distribution in data rows:")
    for val, count in counter.most_common():
        marker = "OK" if val == "Livrée" else "NOT OK"
        print(f"    {marker}  {repr(val)}: {count}")
    # Check that only 'Livrée' remains
    non_livree = {v: c for v, c in counter.items() if v != "Livrée"}
    if non_livree:
        all_good = False
        print(f"  !! Non-Livrée rows still present: {non_livree}")
    else:
        print(f"  -> All {counter['Livrée']} data rows are 'Livrée'.")

wb.close()

print(f"\n{'='*70}")
print(f"TOTAL data rows in cleaned file: {total_data}")
print(f"ALL ROWS 'Livrée'? {'YES' if all_good else 'NO'}")
print("="*70)
