"""
Écrit les valeurs calculées dans le fichier ADMINISTRATEUR DE VENTE.xlsx.

- Feuille 1 (CAHT par agence A): B117-B130 (CA 2026), D117-D130 (CA 2025)
- Feuille 2 (Ventes par produits A): L46-L54 (CA 2026), N46-N54 (CA 2025), K53 = "ALVEOLES"
- Feuille 3 (Vente en volume A): D24-D29 (Vol 2026), E24-E29 (Vol 2025)
"""
import openpyxl
import json
from copy import copy

# Charger les données calculées
with open('/home/z/my-project/scripts/admin_vente_data.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

# Charger le fichier (pas en read_only pour pouvoir écrire)
SRC = '/home/z/my-project/upload/ADMINISTRATEUR DE VENTE.xlsx'
OUT = '/home/z/my-project/download/ADMINISTRATEUR_DE_VENTE_rempli.xlsx'

wb = openpyxl.load_workbook(SRC)

# ===== FEUILLE 1 : CAHT par agence A =====
print("=== Feuille 1: CAHT par agence A ===")
ws1 = wb['CAHT par agence A']

agences_ordre = data['agences_ordre']
agence_ca_2025 = {k: float(v) for k, v in data['agence_ca_2025'].items()}
agence_ca_2026 = {k: float(v) for k, v in data['agence_ca_2026'].items()}

# B117-B130 = CA 2026 (14 agences), D117-D130 = CA 2025
for i, item in enumerate(agences_ordre):
    row = 117 + i
    agence = item['agence']
    ca_2026 = agence_ca_2026.get(agence, 0)
    ca_2025 = agence_ca_2025.get(agence, 0)
    
    ws1.cell(row=row, column=2, value=round(ca_2026, 2))  # B = 2026
    ws1.cell(row=row, column=4, value=round(ca_2025, 2))  # D = 2025
    
    print(f"  R{row} {item['label']:<20} B={ca_2026/1e6:>8.1f}M  D={ca_2025/1e6:>8.1f}M")

# ===== FEUILLE 2 : Ventes par produits A =====
print("\n=== Feuille 2: Ventes par produits A ===")
ws2 = wb['Ventes par produits A']

# K53 doit être "ALVEOLES" (au lieu de MAIS)
ws2.cell(row=53, column=11, value="ALVEOLES")
print(f"  K53 = 'ALVEOLES' (corrigé)")

# Mapping des rubriques aux lignes 46-54
rubriques_ordre = [
    (46, 'VENTES TOURTEAU SOJA', 'TOURTEAU SOJA'),
    (47, 'VENTES INGREDIENTS', 'INGREDIENTS'),
    (48, 'VENTES MATERIEL ELEVAGE', 'MATERIEL ELEVAGE'),
    (49, 'ALIMENTS POISSON & CHICK & PIGLET BOOSTER', 'ALIMENTS COMPLETS'),
    (50, 'VENTES CONCENTRES', 'CONCENTRES'),
    (51, 'VENTES DIVERSES', 'DIVERSES'),
    (52, 'PRODUITS ACCESSOIRES', 'PRODUITS ACCESSOIRES'),
    (53, 'ALVEOLES', 'ALVEOLES'),
    (54, 'MAIS', 'MAIS'),
]

rubriques_2025 = {k: {'ca': float(v['ca']), 'vol': float(v['vol'])} for k, v in data['rubriques_2025'].items()}
rubriques_2026 = {k: {'ca': float(v['ca']), 'vol': float(v['vol'])} for k, v in data['rubriques_2026'].items()}

for row, label, key in rubriques_ordre:
    ca_2026 = rubriques_2026.get(key, {}).get('ca', 0)
    ca_2025 = rubriques_2025.get(key, {}).get('ca', 0)
    
    ws2.cell(row=row, column=12, value=round(ca_2026, 2))  # L = 2026
    ws2.cell(row=row, column=14, value=round(ca_2025, 2))  # N = 2025
    
    print(f"  R{row} {label:<45} L={ca_2026/1e6:>9.2f}M  N={ca_2025/1e6:>9.2f}M")

# ===== FEUILLE 3 : Vente en volume A =====
print("\n=== Feuille 3: Vente en volume A ===")
ws3 = wb['Vente en volume A']

# Lignes 24-29, D = 2026, E = 2025
rubriques_vol = [
    (24, 'VENTES TOURTEAU SOJA', 'TOURTEAU SOJA'),
    (25, 'VENTES INGREDIENTS', 'INGREDIENTS'),
    (26, 'VENTES MAIS', 'MAIS'),
    (27, 'VENTES ALIMENTS COMPLETS', 'ALIMENTS COMPLETS'),
    (28, 'VENTES CONCENTRES', 'CONCENTRES'),
    (29, 'VENTES DIVERSES', 'DIVERSES'),
]

for row, label, key in rubriques_vol:
    vol_2026 = rubriques_2026.get(key, {}).get('vol', 0)
    vol_2025 = rubriques_2025.get(key, {}).get('vol', 0)
    
    ws3.cell(row=row, column=4, value=round(vol_2026, 2))  # D = 2026
    ws3.cell(row=row, column=5, value=round(vol_2025, 2))  # E = 2025
    
    print(f"  R{row} {label:<35} D={vol_2026:>10.2f}t  E={vol_2025:>10.2f}t")

# Sauvegarder
wb.save(OUT)
print(f"\n✓ Fichier sauvegardé: {OUT}")
import os
print(f"  Taille: {os.path.getsize(OUT):,} octets")
