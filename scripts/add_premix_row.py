"""
Ajoute une ligne VENTES PREMIX dans la feuille 2 (Ventes par produits A)
et met à jour toutes les formules.

Nouvelle structure:
- R46: VENTES TOURTEAU SOJA
- R47: VENTES INGREDIENTS
- R48: VENTES MATERIEL ELEVAGE
- R49: ALIMENTS POISSON & CHICK & PIGLET BOOSTER
- R50: VENTES CONCENTRES
- R51: VENTES PREMIX  (← NOUVELLE LIGNE)
- R52: VENTES DIVERSES  (anciennement R51)
- R53: PRODUITS ACCESSOIRES  (anciennement R52)
- R54: ALVEOLES  (anciennement R53)
- R55: MAIS  (anciennement R54)
- R56: Total  (anciennement R55)
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from copy import copy

# Charger valeurs PREMIX
import json
with open('/home/z/my-project/scripts/premix_ca.json') as f:
    premix = json.load(f)

SRC = '/home/z/my-project/download/ADMINISTRATEUR_DE_VENTE_rempli.xlsx'
OUT = '/home/z/my-project/download/ADMINISTRATEUR_DE_VENTE_rempli.xlsx'  # écrase

wb = openpyxl.load_workbook(SRC)
ws = wb['Ventes par produits A']

# Étape 1: Insérer une ligne à R51 (tout décale vers le bas)
ws.insert_rows(51)

# Étape 2: Remplir la nouvelle ligne R51 "VENTES PREMIX"
# K = label, L = CA 2026, M = %, N = CA 2025, O = %, P = Var Valeur, Q = Var %
ws['K51'] = 'VENTES PREMIX'
ws['L51'] = round(premix['ca_2026'], 2)
ws['M51'] = '=+L51/L$56'
ws['N51'] = round(premix['ca_2025'], 2)
ws['O51'] = '=+N51/N$56'
ws['P51'] = '=+L51-N51'
ws['Q51'] = '=+P51/N51'

# Étape 3: Mettre à jour les formules % pour TOUTES les lignes (puisque Total est passé de R55 à R56)
# R46-R50 (avant PREMIX) et R52-R55 (après PREMIX) → tous utilisent L$56 et N$56 maintenant
for row in range(46, 56):
    if row == 51: continue  # déjà fait
    ws.cell(row=row, column=13, value=f'=+L{row}/L$56')  # M = % 2026
    ws.cell(row=row, column=15, value=f'=+N{row}/N$56')  # O = % 2025
    ws.cell(row=row, column=16, value=f'=+L{row}-N{row}')  # P = Var Valeur
    ws.cell(row=row, column=17, value=f'=+P{row}/N{row}')  # Q = Var %

# Étape 4: Mettre à jour la ligne Total (R56)
ws['K56'] = 'Total'
ws['L56'] = '=SUM(L46:L55)'
ws['M56'] = '=SUM(M46:M55)'
ws['N56'] = '=SUM(N46:N55)'
ws['O56'] = '=SUM(O46:O55)'
ws['P56'] = '=SUM(P46:P55)'
ws['Q56'] = '=+P56/N56'

# Vérifier le contenu
print("=== Feuille 2 après ajout PREMIX ===")
for row in range(44, 57):
    vals = []
    for col in range(11, 19):
        v = ws.cell(row, col).value
        vals.append(str(v) if v is not None else "")
    if any(vals):
        print(f"R{row}: {vals}")

wb.save(OUT)
print(f"\n✓ Fichier sauvegardé: {OUT}")
