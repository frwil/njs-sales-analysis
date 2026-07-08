"""
Audit des lignes en état "En cours" et "Validée" (hors PONT_BASCULE et CONTRIBUTION_CARBURANT).
Objectif: avoir une idée du CA/volume à venir quand elles seront "Livrées".
"""
import openpyxl
import json
import re
from collections import defaultdict

# Charger mapping
with open('/home/z/my-project/scripts/product_category_map.json') as f:
    cat_map = json.load(f)

# Poids manuels
MANUAL_WEIGHTS = {"M1051": 50.0, "CF101": 1.0, "S101": 1.0}
SERVICES_VALIDEE = {"PONT_BASCULE", "CONTRIBUTION_CARBURANT"}

WEIGHT_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(KG|Kg|kg|KG|GRAMMES?|G|L)\b", re.IGNORECASE)
def parse_weight_kg(ref, desc):
    ref_str = str(ref).strip() if ref else ""
    if ref_str in MANUAL_WEIGHTS:
        return MANUAL_WEIGHTS[ref_str]
    if not desc: return 0.0
    matches = WEIGHT_RE.findall(str(desc))
    if not matches: return 0.0
    val_str, unit = matches[-1]
    val = float(val_str)
    u = unit.upper()
    if u == "KG": return val
    if u in ("G", "GRAMME", "GRAMMES"): return val/1000.0
    if u == "L": return val
    return 0.0

wb = openpyxl.load_workbook('/home/z/my-project/upload/ventes janv a juin 2026.xlsx', read_only=True, data_only=True)

# Stats par état et catégorie
etat_cat_stats = defaultdict(lambda: defaultdict(lambda: {"vol_kg": 0, "ca": 0, "lignes": 0}))

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    if ws.cell(2, 1).value != 'Réf. produit':
        continue
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or len(row) < 16: continue
        ref = row[0]
        desc = row[1]
        qte = row[2]
        ca = row[8]
        etat = row[15]
        if not ref: continue
        ref_str = str(ref).strip()
        etat_str = str(etat) if etat else ""
        
        # Exclure PONT_BASCULE et CONTRIBUTION_CARBURANT de l'audit (déjà inclus via Validée)
        if ref_str in SERVICES_VALIDEE:
            continue
        
        # Exclure lignes M1051 à CA=0 (régularisation stock)
        if ref_str == "M1051" and (ca is None or float(ca) == 0):
            continue
        
        try:
            q = float(qte) if qte else 0
            c = float(ca) if ca else 0
        except:
            continue
        
        cat = cat_map.get(ref_str, "DIVERS")
        weight = parse_weight_kg(ref_str, desc)
        vol_kg = q * weight
        
        etat_cat_stats[etat_str][cat]["vol_kg"] += vol_kg
        etat_cat_stats[etat_str][cat]["ca"] += c
        etat_cat_stats[etat_str][cat]["lignes"] += 1

wb.close()

# === AFFICHAGE ===
print("=" * 110)
print("AUDIT DES LIGNES PAR ÉTAT (hors PONT_BASCULE et CONTRIBUTION_CARBURANT)")
print("=" * 110)
print()
print(f"{'État':<15} {'Catégorie':<25} {'Lignes':>8} {'Volume (t)':>12} {'CA (M FCFA)':>14}")
print("-" * 80)

etat_summary = defaultdict(lambda: {"vol_kg": 0, "ca": 0, "lignes": 0})
for etat in sorted(etat_cat_stats.keys()):
    for cat in sorted(etat_cat_stats[etat].keys()):
        d = etat_cat_stats[etat][cat]
        vol_t = d["vol_kg"] / 1000.0
        ca_m = d["ca"] / 1e6
        print(f"{etat:<15} {cat:<25} {d['lignes']:>8} {vol_t:>12.1f} {ca_m:>14.2f}")
        etat_summary[etat]["vol_kg"] += d["vol_kg"]
        etat_summary[etat]["ca"] += d["ca"]
        etat_summary[etat]["lignes"] += d["lignes"]
    s = etat_summary[etat]
    print(f"{etat:<15} {'SOUS-TOTAL':<25} {s['lignes']:>8} {s['vol_kg']/1000:>12.1f} {s['ca']/1e6:>14.2f}")
    print()

# === FOCUS SUR "EN COURS" ET "VALIDÉE" ===
print("=" * 110)
print("FOCUS: LIGNES 'EN COURS' ET 'VALIDÉE' (À VENIR LORSQUE LIVRÉES)")
print("=" * 110)
print()
print(f"{'État':<12} {'Catégorie':<25} {'Lignes':>8} {'Volume (t)':>12} {'CA (M FCFA)':>14}")
print("-" * 80)
total_a_venir = {"vol_kg": 0, "ca": 0, "lignes": 0}
for etat in ["En cours", "Validée"]:
    if etat not in etat_cat_stats: continue
    for cat in sorted(etat_cat_stats[etat].keys()):
        d = etat_cat_stats[etat][cat]
        vol_t = d["vol_kg"] / 1000.0
        ca_m = d["ca"] / 1e6
        print(f"{etat:<12} {cat:<25} {d['lignes']:>8} {vol_t:>12.1f} {ca_m:>14.2f}")
        total_a_venir["vol_kg"] += d["vol_kg"]
        total_a_venir["ca"] += d["ca"]
        total_a_venir["lignes"] += d["lignes"]
print("-" * 80)
print(f"{'TOTAL À VENIR':<12} {'':<25} {total_a_venir['lignes']:>8} {total_a_venir['vol_kg']/1000:>12.1f} {total_a_venir['ca']/1e6:>14.2f}")

# === IMPACT SUR LE TOTAL S1 ===
print()
print("=" * 110)
print("IMPACT SUR LE TOTAL S1 2026")
print("=" * 110)
livree_total = etat_summary.get("Livrée", {"vol_kg": 0, "ca": 0, "lignes": 0})
print(f"Livré actuellement (S1 2026):")
print(f"  Volume: {livree_total['vol_kg']/1000:.1f} t")
print(f"  CA HT: {livree_total['ca']/1e6:.1f} M FCFA")
print(f"  Lignes: {livree_total['lignes']}")
print()
print(f"À venir (En cours + Validée hors services):")
print(f"  Volume: {total_a_venir['vol_kg']/1000:.1f} t (+{total_a_venir['vol_kg']/livree_total['vol_kg']*100:.1f}%)")
print(f"  CA HT: {total_a_venir['ca']/1e6:.1f} M FCFA (+{total_a_venir['ca']/livree_total['ca']*100:.1f}%)")
print()
print(f"TOTAL PROJETÉ (Livré + À venir):")
print(f"  Volume: {(livree_total['vol_kg']+total_a_venir['vol_kg'])/1000:.1f} t")
print(f"  CA HT: {(livree_total['ca']+total_a_venir['ca'])/1e6:.1f} M FCFA")

# Sauvegarder
output = {
    "audit": {
        etat: {
            cat: {
                "lignes": etat_cat_stats[etat][cat]["lignes"],
                "vol_t": etat_cat_stats[etat][cat]["vol_kg"]/1000,
                "ca_m": etat_cat_stats[etat][cat]["ca"]/1e6
            } for cat in etat_cat_stats[etat]
        } for etat in etat_cat_stats
    },
    "a_venir": {
        "lignes": total_a_venir["lignes"],
        "vol_t": total_a_venir["vol_kg"]/1000,
        "ca_m": total_a_venir["ca"]/1e6
    },
    "livree": {
        "lignes": livree_total["lignes"],
        "vol_t": livree_total["vol_kg"]/1000,
        "ca_m": livree_total["ca"]/1e6
    }
}
with open('/home/z/my-project/scripts/audit_pipeline.json', 'w', encoding='utf-8') as f:
    json.dump(output, f, ensure_ascii=False, indent=2)
print(f"\n✓ Sauvegardé dans audit_pipeline.json")
