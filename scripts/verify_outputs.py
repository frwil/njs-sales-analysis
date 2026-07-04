"""Verify both output files."""
from openpyxl import load_workbook

CHECK = "/home/z/my-project/download/Clients uniques - check mensuel.xlsx"
ZERO = "/home/z/my-project/download/Clients zero achat - Janv-Mars.xlsx"

print("="*80)
print("FILE 1: Clients uniques - check mensuel.xlsx (UPDATED WITH 20/80)")
print("="*80)
wb = load_workbook(CHECK, data_only=True)
ws = wb.active
print(f"Sheet: {ws.title}  (max_row={ws.max_row}, max_col={ws.max_column})")
print(f"Row 1 title: {ws['A1'].value}")
print(f"\nHeaders (row 3):")
for c in range(1, ws.max_column + 1):
    print(f"  Col {chr(64+c)}: {ws.cell(row=3, column=c).value}")

print(f"\nFirst 10 rows (showing only key columns):")
print(f"  {'N°':<4} {'Réf':<14} {'Nom':<35} {'Jan':<4} {'Fev':<4} {'Mar':<4} {'CA HT':>15} {'20/80':<6}")
for r in range(4, 14):
    n = ws.cell(row=r, column=1).value
    ref = ws.cell(row=r, column=2).value or ""
    name = (ws.cell(row=r, column=3).value or "")[:33]
    jan = "✓" if ws.cell(row=r, column=4).value else "-"
    fev = "✓" if ws.cell(row=r, column=5).value else "-"
    mar = "✓" if ws.cell(row=r, column=6).value else "-"
    ca = ws.cell(row=r, column=11).value or 0
    p = ws.cell(row=r, column=12).value or ""
    print(f"  {n:<4} {ref:<14} {name:<35} {jan:<4} {fev:<4} {mar:<4} {ca:>15,.0f} {p:<6}")

# Count Pareto clients in the file
pareto_count = 0
for r in range(4, ws.max_row + 1):
    if ws.cell(row=r, column=12).value:
        pareto_count += 1
print(f"\nTotal Pareto (★) clients: {pareto_count}")
wb.close()

print("\n" + "="*80)
print("FILE 2: Clients zero achat - Janv-Mars.xlsx")
print("="*80)
wb = load_workbook(ZERO, data_only=True)
ws = wb.active
print(f"Sheet: {ws.title}  (max_row={ws.max_row}, max_col={ws.max_column})")
print(f"Row 1 title: {ws['A1'].value}")
print(f"Row 2 subtitle: {ws['A2'].value}")
print(f"\nHeaders (row 3):")
for c in range(1, ws.max_column + 1):
    print(f"  Col {chr(64+c)}: {ws.cell(row=3, column=c).value}")

print(f"\nFirst 15 rows:")
for r in range(4, 19):
    n = ws.cell(row=r, column=1).value
    ref = ws.cell(row=r, column=2).value or ""
    name = (ws.cell(row=r, column=3).value or "")[:35]
    q1 = ws.cell(row=r, column=4).value or ""
    ca_q1 = ws.cell(row=r, column=5).value or 0
    any6 = ws.cell(row=r, column=6).value or ""
    ca_6 = ws.cell(row=r, column=7).value or 0
    p = ws.cell(row=r, column=8).value or ""
    print(f"  {n:<4} {ref:<14} {name:<35} {q1:<4} {ca_q1:>12,.0f}  {any6:<4} {ca_6:>15,.0f}  {p}")

print(f"\nLast 5 rows:")
for r in range(ws.max_row - 4, ws.max_row + 1):
    n = ws.cell(row=r, column=1).value
    ref = ws.cell(row=r, column=2).value or ""
    name = (ws.cell(row=r, column=3).value or "")[:35]
    q1 = ws.cell(row=r, column=4).value or ""
    ca_q1 = ws.cell(row=r, column=5).value or 0
    any6 = ws.cell(row=r, column=6).value or ""
    ca_6 = ws.cell(row=r, column=7).value or 0
    p = ws.cell(row=r, column=8).value or ""
    print(f"  {n:<4} {ref:<14} {name:<35} {q1:<4} {ca_q1:>12,.0f}  {any6:<4} {ca_6:>15,.0f}  {p}")

print(f"\nTotal zero-achat clients: {ws.max_row - 3}")
wb.close()
