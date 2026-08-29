"""
Excel - Commandes SPC par agence × produit × format × scénario.
"""
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import os

orders = pd.read_csv("/home/z/my-project/scripts/spc_orders.csv")
print(f"Loaded {len(orders)} order lines")

# Styles
HEAD_FILL = PatternFill('solid', fgColor='1F4E78')
HEAD_FONT = Font(bold=True, color='FFFFFF', size=11)
SUBHEAD_FILL = PatternFill('solid', fgColor='D9E1F2')
SUBHEAD_FONT = Font(bold=True, color='1F4E78', size=11)
TOTAL_FILL = PatternFill('solid', fgColor='FFF2CC')
TOTAL_FONT = Font(bold=True, size=11)
SCEN_COLORS = {'prudent': 'FCE4D6', 'realiste': 'C6EFCE', 'optimiste': 'BDD7EE'}
THIN = Side(border_style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def style_header_row(ws, row, n_cols):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER

def style_total_row(ws, row, n_cols):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = TOTAL_FILL
        cell.font = TOTAL_FONT
        cell.border = BORDER

def style_data_row(ws, row, n_cols, color=None):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.border = BORDER
        if color:
            cell.fill = PatternFill('solid', fgColor=color)

wb = openpyxl.Workbook()
wb.remove(wb.active)

# === Sheet 1: Synthèse ===
ws = wb.create_sheet("1. Synthèse")
ws['A1'] = 'BELGOCAM SA - Commandes Aliment Complet SPC'
ws['A1'].font = Font(bold=True, size=16, color='1F4E78')
ws['A2'] = 'Substitution soja → aliment complet SPC | 3 scénarios | Toutes agences sauf BUEA | Ratio 1:3'
ws['A2'].font = Font(italic=True, size=10, color='595959')

ws['A4'] = 'SYNTHÈSE PAR SCÉNARIO'
ws['A4'].font = SUBHEAD_FONT

headers = ['Scénario', 'Taux bascule', 'Total unités', 'Total sacs éq. 50kg', 'Description']
for i, h in enumerate(headers, 1):
    ws.cell(row=5, column=i, value=h)
style_header_row(ws, 5, 5)

scn_info = [
    ('prudent', '10%', 'Prudent - 10% des clients soja basculent vers SPC'),
    ('realiste', '30%', 'Réaliste - 30% des clients soja basculent (recommandé)'),
    ('optimiste', '50%', 'Optimiste - 50% des clients soja basculent'),
]

row = 6
for scn, taux, desc in scn_info:
    sub = orders[orders['scenario'] == scn]
    total_units = sub['unites_commander'].sum()
    total_sacs = sub['sacs_50_eq'].sum()
    ws.cell(row=row, column=1, value=scn.upper())
    ws.cell(row=row, column=2, value=taux)
    ws.cell(row=row, column=3, value=f"{total_units:,.0f}")
    ws.cell(row=row, column=4, value=f"{total_sacs:,.0f}")
    ws.cell(row=row, column=5, value=desc)
    style_data_row(ws, row, 5, color=SCEN_COLORS[scn])
    row += 1

row += 1
ws.cell(row=row, column=1, value='SYNTHÈSE PAR CATÉGORIE (Scénario réaliste 30%)')
ws.cell(row=row, column=1).font = SUBHEAD_FONT
row += 1

headers = ['Catégorie', 'Unités', 'Part (%)', 'Sacs éq. 50kg']
for i, h in enumerate(headers, 1):
    ws.cell(row=row, column=i, value=h)
style_header_row(ws, row, 4)
row += 1

realiste = orders[orders['scenario'] == 'realiste']
cat_sum = realiste.groupby('categorie').agg(
    unites=('unites_commander', 'sum'),
    sacs=('sacs_50_eq', 'sum'),
).sort_values('unites', ascending=False)

total_u = cat_sum['unites'].sum()
for cat, r in cat_sum.iterrows():
    ws.cell(row=row, column=1, value=cat)
    ws.cell(row=row, column=2, value=f"{r['unites']:,.0f}")
    ws.cell(row=row, column=3, value=f"{r['unites']/total_u*100:.1f}%")
    ws.cell(row=row, column=4, value=f"{r['sacs']:,.0f}")
    style_data_row(ws, row, 4, color='C6EFCE')
    row += 1

ws.cell(row=row, column=1, value='TOTAL')
ws.cell(row=row, column=2, value=f"{total_u:,.0f}")
ws.cell(row=row, column=3, value='100.0%')
ws.cell(row=row, column=4, value=f"{cat_sum['sacs'].sum():,.0f}")
style_total_row(ws, row, 4)

ws.column_dimensions['A'].width = 18
ws.column_dimensions['B'].width = 15
ws.column_dimensions['C'].width = 18
ws.column_dimensions['D'].width = 20
ws.column_dimensions['E'].width = 55

# === Sheet 2: Par Agence ===
ws = wb.create_sheet("2. Par Agence")
ws['A1'] = 'COMMANDES SPC PAR AGENCE (3 scénarios)'
ws['A1'].font = Font(bold=True, size=14, color='1F4E78')

row = 3
headers = ['Scénario', 'Agence', 'Région', 'Unités 5kg', 'Unités 10kg', 'Total unités', 'Sacs éq. 50kg']
for i, h in enumerate(headers, 1):
    ws.cell(row=row, column=i, value=h)
style_header_row(ws, row, 7)
row += 1

for scn, _, _ in scn_info:
    sub = orders[orders['scenario'] == scn]
    ag_sum = sub.groupby(['agence', 'region', 'format'])['unites_commander'].sum().unstack(fill_value=0)
    if '5kg' not in ag_sum.columns: ag_sum['5kg'] = 0
    if '10kg' not in ag_sum.columns: ag_sum['10kg'] = 0
    ag_sum['total'] = ag_sum['5kg'] + ag_sum['10kg']
    sacs = sub.groupby(['agence', 'region'])['sacs_50_eq'].sum()
    ag_sum = ag_sum.merge(sacs, left_index=True, right_index=True)
    ag_sum = ag_sum.sort_values('total', ascending=False)
    
    for idx, r in ag_sum.iterrows():
        ag, reg = idx
        ws.cell(row=row, column=1, value=scn.upper())
        ws.cell(row=row, column=2, value=ag)
        ws.cell(row=row, column=3, value=reg)
        ws.cell(row=row, column=4, value=f"{r['5kg']:,.0f}")
        ws.cell(row=row, column=5, value=f"{r['10kg']:,.0f}")
        ws.cell(row=row, column=6, value=f"{r['total']:,.0f}")
        ws.cell(row=row, column=7, value=f"{r['sacs_50_eq']:,.0f}")
        style_data_row(ws, row, 7, color=SCEN_COLORS[scn])
        row += 1
    # Total
    ws.cell(row=row, column=1, value=scn.upper())
    ws.cell(row=row, column=2, value='TOTAL')
    ws.cell(row=row, column=3, value='—')
    ws.cell(row=row, column=4, value=f"{ag_sum['5kg'].sum():,.0f}")
    ws.cell(row=row, column=5, value=f"{ag_sum['10kg'].sum():,.0f}")
    ws.cell(row=row, column=6, value=f"{ag_sum['total'].sum():,.0f}")
    ws.cell(row=row, column=7, value=f"{ag_sum['sacs_50_eq'].sum():,.0f}")
    style_total_row(ws, row, 7)
    row += 2

ws.column_dimensions['A'].width = 14
ws.column_dimensions['B'].width = 18
ws.column_dimensions['C'].width = 12
for col in 'DEFG':
    ws.column_dimensions[col].width = 16

# === Sheet 3: Par Produit ===
ws = wb.create_sheet("3. Par Produit SPC")
ws['A1'] = 'COMMANDES SPC PAR PRODUIT (3 scénarios)'
ws['A1'].font = Font(bold=True, size=14, color='1F4E78')

row = 3
headers = ['Scénario', 'Réf', 'Description', 'Catégorie', 'Unités 5kg', 'Unités 10kg', 'Total unités', 'Sacs éq. 50kg']
for i, h in enumerate(headers, 1):
    ws.cell(row=row, column=i, value=h)
style_header_row(ws, row, 8)
row += 1

for scn, _, _ in scn_info:
    sub = orders[orders['scenario'] == scn]
    prod_sum = sub.groupby(['produit_spc', 'description', 'categorie', 'format'])['unites_commander'].sum().unstack(fill_value=0)
    if '5kg' not in prod_sum.columns: prod_sum['5kg'] = 0
    if '10kg' not in prod_sum.columns: prod_sum['10kg'] = 0
    prod_sum['total'] = prod_sum['5kg'] + prod_sum['10kg']
    sacs = sub.groupby(['produit_spc', 'description', 'categorie'])['sacs_50_eq'].sum()
    prod_sum = prod_sum.merge(sacs, left_index=True, right_index=True)
    prod_sum = prod_sum.sort_values('total', ascending=False)
    
    for idx, r in prod_sum.iterrows():
        ref, desc, cat = idx
        ws.cell(row=row, column=1, value=scn.upper())
        ws.cell(row=row, column=2, value=ref)
        ws.cell(row=row, column=3, value=desc)
        ws.cell(row=row, column=4, value=cat)
        ws.cell(row=row, column=5, value=f"{r['5kg']:,.0f}")
        ws.cell(row=row, column=6, value=f"{r['10kg']:,.0f}")
        ws.cell(row=row, column=7, value=f"{r['total']:,.0f}")
        ws.cell(row=row, column=8, value=f"{r['sacs_50_eq']:,.0f}")
        style_data_row(ws, row, 8, color=SCEN_COLORS[scn])
        row += 1
    row += 1

ws.column_dimensions['A'].width = 14
ws.column_dimensions['B'].width = 12
ws.column_dimensions['C'].width = 28
ws.column_dimensions['D'].width = 12
for col in 'EFGH':
    ws.column_dimensions[col].width = 16

# === Sheet 4: Détail complet ===
ws = wb.create_sheet("4. Détail complet")
ws['A1'] = 'DÉTAIL COMPLET - COMMANDES SPC (Produit × Agence × Format × Scénario)'
ws['A1'].font = Font(bold=True, size=14, color='1F4E78')

row = 3
headers = ['Scénario', 'Agence', 'Région', 'Réf SPC', 'Description', 'Catégorie', 'Format', 'Unités à commander', 'Sacs éq. 50kg']
for i, h in enumerate(headers, 1):
    ws.cell(row=row, column=i, value=h)
style_header_row(ws, row, 9)
row += 1

for scn in ['prudent', 'realiste', 'optimiste']:
    sub = orders[orders['scenario'] == scn].sort_values(['agence', 'categorie', 'produit_spc', 'format'])
    for _, r in sub.iterrows():
        ws.cell(row=row, column=1, value=scn.upper())
        ws.cell(row=row, column=2, value=r['agence'])
        ws.cell(row=row, column=3, value=r['region'])
        ws.cell(row=row, column=4, value=r['produit_spc'])
        ws.cell(row=row, column=5, value=r['description'])
        ws.cell(row=row, column=6, value=r['categorie'])
        ws.cell(row=row, column=7, value=r['format'])
        ws.cell(row=row, column=8, value=f"{r['unites_commander']:,.0f}")
        ws.cell(row=row, column=9, value=f"{r['sacs_50_eq']:,.1f}")
        style_data_row(ws, row, 9, color=SCEN_COLORS[scn])
        row += 1

ws.column_dimensions['A'].width = 14
ws.column_dimensions['B'].width = 15
ws.column_dimensions['C'].width = 12
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 28
ws.column_dimensions['F'].width = 12
ws.column_dimensions['G'].width = 8
for col in 'HI':
    ws.column_dimensions[col].width = 18
ws.freeze_panes = 'A4'

# === Sheet 5: Hypothèses ===
ws = wb.create_sheet("5. Hypothèses")
ws['A1'] = 'HYPOTHÈSES ET MÉTHODOLOGIE'
ws['A1'].font = Font(bold=True, size=14, color='1F4E78')

hyp = [
    ('OBJECTIF', 'Définir les quantités d\'aliment complet SPC à commander pour pallier le manque de soja'),
    ('', ''),
    ('PARAMÈTRES', ''),
    ('Ratio de substitution', '1 sac soja (50kg) = 3 sacs aliment complet (le soja représente ~30% de la formulation)'),
    ('Répartition formats', '30% en 5kg / 70% en 10kg (le 10kg plus économique pour l\'éleveur)'),
    ('Agences concernées', '13 agences (toutes sauf BUEA)'),
    ('Régions', 'Littoral (pilote) + Ouest + Centre (calque proportionnel au soja)'),
    ('', ''),
    ('SCÉNARIOS DE BASCULE', ''),
    ('Prudent (10%)', '10% des clients soja basculent vers l\'aliment complet SPC'),
    ('Réaliste (30%)', '30% des clients soja basculent (scénario recommandé)'),
    ('Optimiste (50%)', '50% des clients soja basculent'),
    ('', ''),
    ('MIX PRODUITS SPC AJUSTÉ', ''),
    ('Source SPC', 'Fichier "Proportion ou contribution type d\'aliment région Littoral.xlsx" (1 mois, 9 708 sacs)'),
    ('Ajustement BELGOCAM', 'Le mix SPC (77% chair, 4% ponte, 13% porc) est ajusté selon le mix BELGOCAM par agence (via les ventes concentrés)'),
    ('Note SPC', 'Le responsable SPC indique: "Chair contribue plus, puis porc, puis ponte. Cette tendance ne sera pas pareille chez BELGOCAM"'),
    ('Résultat', 'BELGOCAM a proportionnellement plus de PONTE et moins de CHAIR que SPC → le mix est ajusté en conséquence'),
    ('', ''),
    ('PRODUITS SPC (22 références)', ''),
    ('VOLAILLE CHAIR', 'BM1, BM2, BM3, BM1E (4 produits)'),
    ('VOLAILLE PONTE', 'AM1, AM2, AM3 (3 produits)'),
    ('PORC', 'P1, P2, P3, P4, P5, P1E (6 produits)'),
    ('LAPIN', 'L3G (1 produit)'),
    ('POISSON', 'CL2, CL3 (2 produits)'),
    ('CHEVAL', 'CH3 (1 produit)'),
    ('CHIEN', 'CN1 (1 produit)'),
    ('BOVIN/CHEVRE', 'ALT BOVIN EN/EMB, VACHE LAITIERE, CHEVRES (4 produits, volume 0)'),
    ('', ''),
    ('CONVERSION FORMAT', ''),
    ('Format 5kg', '1 sac éq. 50kg = 10 sacs de 5kg'),
    ('Format 10kg', '1 sac éq. 50kg = 5 sacs de 10kg'),
]

row = 3
for label, value in hyp:
    if label and not value:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=1).font = SUBHEAD_FONT
        ws.cell(row=row, column=1).fill = SUBHEAD_FILL
    elif label:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        ws.cell(row=row, column=1).border = BORDER
        ws.cell(row=row, column=2).border = BORDER
    row += 1

ws.column_dimensions['A'].width = 30
ws.column_dimensions['B'].width = 85

# === Save ===
output_path = "/home/z/my-project/download/commandes_spc_aliment_complet.xlsx"
wb.save(output_path)
print(f"\n=== EXCEL SPC SAVED ===")
print(f"Path: {output_path}")
print(f"Size: {os.path.getsize(output_path) / 1024:.0f} KB")
print(f"Sheets: {wb.sheetnames}")
