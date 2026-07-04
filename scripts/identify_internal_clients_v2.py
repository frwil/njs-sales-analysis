"""
Refined identification of internal/filiale clients to exclude.
Focus on:
1. Filiales du Groupe NJS (explicit "FILIALE GROUPE NJS")
2. SPC (Société de Provenderie du Cameroun - internal shops)
3. SOLDE COMPTA (accounting balances, not real sales)
4. BELGOCAM itself
"""
from openpyxl import load_workbook
import re

SRC = "/home/z/my-project/download/ventes_livrees.xlsx"

wb = load_workbook(SRC, read_only=True, data_only=True)

# Refined patterns - only TRUE internal clients
INTERNAL_PATTERNS = [
    r"FILIALE\s+GROUPE\s+NJS",
    r"\bSPC\b",
    r"SOLDE\s+COMPTA",
    r"\bBELGOCAM\b",
    r"\bNJS\s+GROUP\b",
]
INTERNAL_RE = re.compile("|".join(INTERNAL_PATTERNS), re.IGNORECASE)

clients = {}
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for row in ws.iter_rows(min_row=3, values_only=True):
        tiers = row[5] if len(row) > 5 else None
        if tiers is None or str(tiers).strip() == "":
            continue
        if tiers not in clients:
            clients[tiers] = 0
        ca = row[8] if len(row) > 8 else 0
        try:
            clients[tiers] += float(ca) if ca is not None else 0
        except:
            pass

wb.close()

print(f"Total unique clients: {len(clients)}")
print(f"\nClients to EXCLUDE (true internal/filiale):")
print(f"{'Tiers':<80} {'CA HT 6 mois':>15}")
print("-" * 100)
internal_clients = []
for tiers, ca in sorted(clients.items(), key=lambda x: -x[1]):
    if INTERNAL_RE.search(str(tiers)):
        internal_clients.append((tiers, ca))
        print(f"{str(tiers)[:78]:<80} {ca:>15,.0f}")

print(f"\nTotal clients to exclude: {len(internal_clients)}")
total_ca_internal = sum(c for _, c in internal_clients)
print(f"Total CA excluded: {total_ca_internal:,.0f} FCFA")

# Save the list of tiers to exclude for use in the consolidate script
import json
excluded_tiers = [t for t, _ in internal_clients]
with open("/home/z/my-project/scripts/excluded_clients.json", "w", encoding="utf-8") as f:
    json.dump(excluded_tiers, f, ensure_ascii=False, indent=2)
print(f"\nExcluded list saved to: /home/z/my-project/scripts/excluded_clients.json")
