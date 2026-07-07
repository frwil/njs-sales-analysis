"""
Compute 2025 sales by category x month, then compare with 2026 S1.
Build refined forecast S2 2026 using 2025 seasonality.
"""
import os
import json
import re
from openpyxl import load_workbook
from collections import defaultdict

SRC_2025 = "/home/z/my-project/upload/86d96135-9db7-45bc-bba6-a69efa2c5ee3.xlsx"

# Load excluded clients
with open("/home/z/my-project/scripts/excluded_clients.json", "r", encoding="utf-8") as f:
    excluded_tiers = set(json.load(f))

# Patterns to also exclude (in case 2025 has CLIENT(S) COMPTOIR variants we missed)
COMPTOIR_RE = re.compile(r"CLIENTS?\s+COMPTOIR", re.IGNORECASE)
INTERNAL_RE = re.compile(r"FILIALE\s+GROUPE\s+NJS|SOLDE\s+COMPTA|BELGOCAM\b|NJS\s+GROUP\b", re.IGNORECASE)

# ===== READ 2025 SALES =====
print("Reading 2025 sales file...")
wb = load_workbook(SRC_2025, read_only=True, data_only=True)
ws = wb.active

# category_2025[category][month] = {"vol_t": float, "ca": float}
category_2025 = defaultdict(lambda: defaultdict(lambda: {"vol_t": 0.0, "ca": 0.0}))
total_excluded = 0
total_rows = 0

for row in ws.iter_rows(min_row=2, values_only=True):
    tiers = row[4] if len(row) > 4 else None  # E
    if tiers is None or str(tiers).strip() == "": continue
    tiers_str = str(tiers)
    # Apply exclusion
    if tiers_str in excluded_tiers or COMPTOIR_RE.search(tiers_str) or INTERNAL_RE.search(tiers_str):
        total_excluded += 1
        continue
    total_rows += 1

    state = row[9] if len(row) > 9 else None  # J
    if state and str(state).strip() != "Livrée":
        continue

    cat = row[10] if len(row) > 10 else None  # K
    vol_t = row[11] if len(row) > 11 else 0  # L (already in tons)
    ca_ht = row[6] if len(row) > 6 else 0  # G
    date_cmd = row[5] if len(row) > 5 else None  # F

    if cat is None: continue
    cat_str = str(cat).strip()

    # Extract month
    month = None
    if date_cmd:
        if hasattr(date_cmd, 'month'):
            month = date_cmd.month
        else:
            s = str(date_cmd)
            parts = s.split("/")
            if len(parts) == 3:
                try: month = int(parts[1])
                except: pass
            elif len(s) >= 7:
                try: month = int(s[5:7])
                except: pass
    if month is None: continue

    try:
        vol_t_f = float(vol_t) if vol_t is not None else 0.0
    except: vol_t_f = 0.0
    try:
        ca_f = float(ca_ht) if ca_ht is not None else 0.0
    except: ca_f = 0.0

    category_2025[cat_str][month]["vol_t"] += vol_t_f
    category_2025[cat_str][month]["ca"] += ca_f

wb.close()
print(f"Total rows (after exclusion): {total_rows}")
print(f"Total excluded rows: {total_excluded}")

# Print summary 2025
print("\n=== 2025 SALES BY CATEGORY (tons) ===")
print(f"{'Category':<28}", end='')
for m in range(1, 13):
    print(f"  M{m:>2}", end='')
print(f"{'TOTAL':>10}{'CA (M FCFA)':>14}")
for cat in sorted(category_2025.keys()):
    print(f"{cat[:28]:<28}", end='')
    total_vol = 0
    total_ca = 0
    for m in range(1, 13):
        v = category_2025[cat][m]["vol_t"]
        total_vol += v
        total_ca += category_2025[cat][m]["ca"]
        print(f"{v:>6.1f}", end='')
    print(f"{total_vol:>10.1f}{total_ca/1e6:>14.1f}")

# ===== COMPUTE YoY GROWTH (S1 2025 vs S1 2026) =====
print("\n=== YoY GROWTH (S1 2025 vs S1 2026) ===")
# Load 2026 S1 data
with open("/home/z/my-project/scripts/objectives_comparison.json", "r") as f:
    data_2026 = json.load(f)

# Map 2025 categories to 2026 categories (they should be the same, except DIVERS2)
def map_cat(c):
    if c == "DIVERS2": return "DIVERS"
    return c

# Compute S1 2025 (months 1-6) per category
s1_2025 = {}
for cat in category_2025:
    mapped = map_cat(cat)
    if mapped not in s1_2025:
        s1_2025[mapped] = {"vol_t": 0, "ca": 0}
    for m in range(1, 7):
        s1_2025[mapped]["vol_t"] += category_2025[cat][m]["vol_t"]
        s1_2025[mapped]["ca"] += category_2025[cat][m]["ca"]

# Compute S2 2025 (months 7-12) per category
s2_2025 = {}
for cat in category_2025:
    mapped = map_cat(cat)
    if mapped not in s2_2025:
        s2_2025[mapped] = {"vol_t": 0, "ca": 0}
    for m in range(7, 13):
        s2_2025[mapped]["vol_t"] += category_2025[cat][m]["vol_t"]
        s2_2025[mapped]["ca"] += category_2025[cat][m]["ca"]

# Compare
print(f"\n{'Category':<28}{'S1 2025 (t)':>14}{'S1 2026 (t)':>14}{'YoY %':>10}{'S2 2025 (t)':>14}")
yoy_data = {}
for cat in sorted(set(list(s1_2025.keys()) + [d['category'] for d in data_2026['comparison']])):
    s1_25 = s1_2025.get(cat, {"vol_t": 0})["vol_t"]
    # Find S1 2026 real
    d = next((x for x in data_2026['comparison'] if x['category'] == cat), None)
    s1_26 = d['s1_real'] if d else 0
    s2_25 = s2_2025.get(cat, {"vol_t": 0})["vol_t"]
    yoy = ((s1_26 - s1_25) / s1_25 * 100) if s1_25 > 0 else 0
    print(f"{cat[:28]:<28}{s1_25:>14.1f}{s1_26:>14.1f}{yoy:>9.1f}%{s2_25:>14.1f}")
    yoy_data[cat] = {"s1_2025": s1_25, "s1_2026": s1_26, "yoy_pct": yoy, "s2_2025": s2_25}

# ===== REFINED FORECAST S2 2026 =====
# Method: S2 2026 forecast = S2 2025 * (1 + YoY growth S1 2026 vs S1 2025)
# Adjusted by scenario:
#   Pessimist: YoY - 10 points
#   Realist: YoY
#   Optimist: YoY + 10 points (plan d'action + rupture concurrente)
print("\n=== REFINED FORECAST S2 2026 (using 2025 seasonality + YoY growth) ===")
print(f"\n{'Category':<28}{'S2 2025 (t)':>14}{'YoY S1':>10}{'🔴 Pess.':>12}{'🟡 Real.':>12}{'🟢 Opt.':>12}{'S2 obj.':>12}")
forecast_refined = {}
for cat in sorted(yoy_data.keys()):
    s2_25 = yoy_data[cat]["s2_2025"]
    yoy = yoy_data[cat]["yoy_pct"] / 100
    # S2 objective
    s2_obj = sum(data_2026['global_objectives'].get(cat, {}).get(str(m), 0) for m in range(7, 13))

    # 3 scenarios using YoY as baseline growth
    pess_mult = max(0, 1 + yoy - 0.10)
    real_mult = 1 + yoy
    opt_mult = 1 + yoy + 0.10

    pess = s2_25 * pess_mult
    real = s2_25 * real_mult
    opt = s2_25 * opt_mult

    print(f"{cat[:28]:<28}{s2_25:>14.1f}{yoy*100:>9.1f}%{pess:>12.1f}{real:>12.1f}{opt:>12.1f}{s2_obj:>12.1f}")
    forecast_refined[cat] = {"s2_2025": s2_25, "yoy_pct": yoy*100, "pess": pess, "real": real, "opt": opt, "s2_obj": s2_obj}

# Totals
total_s2_25 = sum(forecast_refined[c]["s2_2025"] for c in forecast_refined)
total_pess = sum(forecast_refined[c]["pess"] for c in forecast_refined)
total_real = sum(forecast_refined[c]["real"] for c in forecast_refined)
total_opt = sum(forecast_refined[c]["opt"] for c in forecast_refined)
total_s2_obj = sum(forecast_refined[c]["s2_obj"] for c in forecast_refined)
print(f"\n{'TOTAL':<28}{total_s2_25:>14.1f}{'':>10}{total_pess:>12.1f}{total_real:>12.1f}{total_opt:>12.1f}{total_s2_obj:>12.1f}")
print(f"\nvs objectif S2: Pess={total_pess/total_s2_obj*100:.1f}%  Real={total_real/total_s2_obj*100:.1f}%  Opt={total_opt/total_s2_obj*100:.1f}%")

# Save data
output = {
    "category_2025_monthly": {cat: {str(m): category_2025[cat][m] for m in range(1, 13)} for cat in category_2025},
    "yoy_data": yoy_data,
    "forecast_refined": forecast_refined,
    "totals": {
        "s2_2025": total_s2_25,
        "pess": total_pess,
        "real": total_real,
        "opt": total_opt,
        "s2_obj": total_s2_obj,
    }
}
with open("/home/z/my-project/scripts/yoy_forecast.json", "w", encoding="utf-8") as f:
    json.dump(output, f, ensure_ascii=False, indent=2, default=str)
print(f"\nData saved to /home/z/my-project/scripts/yoy_forecast.json")
