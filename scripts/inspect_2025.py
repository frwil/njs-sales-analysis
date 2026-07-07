"""Inspect 2025 sales file in detail."""
from openpyxl import load_workbook
from collections import Counter

SRC = "/home/z/my-project/upload/86d96135-9db7-45bc-bba6-a69efa2c5ee3.xlsx"

wb = load_workbook(SRC, read_only=True, data_only=True)
ws = wb.active

print(f"Sheet: {ws.title}")
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")

# Print first 5 rows
print("\n=== First 5 rows ===")
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True), start=1):
    print(f"\n--- Row {i} ---")
    headers = ["A:Réf.prod", "B:Desc", "C:Qté", "D:Réf.cmd", "E:Tiers", "F:Date",
               "G:MontantHT", "H:MontantTTC", "I:Facturé", "J:État",
               "K:Catégorie", "L:Qté tonnes", "M:Agence", "N:Région", "O:typeClient"]
    for j, v in enumerate(row):
        if v is not None:
            print(f"  {headers[j]}: {repr(v)[:80]}")

# Distinct categories
print("\n=== Distinct categories (col K) ===")
cats = Counter()
states = Counter()
agences = Counter()
months = Counter()
for row in ws.iter_rows(min_row=2, values_only=True):
    cat = row[10] if len(row) > 10 else None  # K
    state = row[9] if len(row) > 9 else None  # J
    agence = row[12] if len(row) > 12 else None  # M
    date_cmd = row[5] if len(row) > 5 else None  # F
    if cat: cats[str(cat).strip()] += 1
    if state: states[str(state).strip()] += 1
    if agence: agences[str(agence).strip()] += 1
    if date_cmd:
        s = str(date_cmd)
        parts = s.split("/")
        if len(parts) == 3:
            months[f"2025-{parts[1]}"] += 1
        elif len(s) >= 7:
            months[s[:7]] += 1

print(f"Categories: {len(cats)}")
for c, n in cats.most_common():
    print(f"  {c}: {n} lignes")

print(f"\nÉtats: {len(states)}")
for s, n in states.most_common():
    print(f"  {s}: {n} lignes")

print(f"\nMois: {len(months)}")
for m, n in sorted(months.items()):
    print(f"  {m}: {n} lignes")

print(f"\nAgences: {len(agences)}")
for a, n in agences.most_common():
    print(f"  {a}: {n} lignes")

wb.close()
