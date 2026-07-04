"""
Extract unique clients from the cleaned (Livrée only) file and produce
a per-month purchase check matrix.

Source: /home/z/my-project/download/ventes janv a juin 2026 - Livrées uniquement.xlsx
Output: /home/z/my-project/download/Clients uniques - check mensuel.xlsx

For each unique client (column F "Tiers"):
- Split into Réf. client (CU code) + Nom du client
- For each of the 6 months (Sheet 1=Jan, Feuil1=Fev, ..., Feuil5=Juin),
  check whether the client appears in that sheet
- Mark with green ✓ if purchases were made, empty otherwise
"""
import re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict
import os

SRC = "/home/z/my-project/download/ventes janv a juin 2026 - Livrées uniquement.xlsx"
DST = "/home/z/my-project/download/Clients uniques - check mensuel.xlsx"

MONTHS = [
    ("Sheet 1", "Janvier"),
    ("Feuil1",  "Février"),
    ("Feuil2",  "Mars"),
    ("Feuil3",  "Avril"),
    ("Feuil4",  "Mai"),
    ("Feuil5",  "Juin"),
]

# Pattern to split "CU2601-14325 -  TCHONBEUA AUGUSTIN" -> ("CU2601-14325", "TCHONBEUA AUGUSTIN")
# Also handles ZPCL..., ZXCL..., or any alnum code prefix.
# The separator is " - " (with possible extra spaces)
TIERS_RE = re.compile(r"^\s*([A-Z]{2,}[\w-]*?)\s*-\s+(.+?)\s*$")

def split_tiers(tiers):
    """Split a Tiers value into (client_ref, client_name)."""
    if tiers is None:
        return ("", "")
    s = str(tiers).strip()
    m = TIERS_RE.match(s)
    if m:
        return (m.group(1).strip(), m.group(2).strip())
    # Fallback: no clear ref - put everything in name
    return ("", s)

# --- 1. Read source file and collect per-sheet client sets -----------------
print(f"Reading source: {SRC}")
wb_src = load_workbook(SRC, read_only=True, data_only=True)

# Map: client_key -> {month_idx: True/False}
# client_key = (ref, name)
client_months = defaultdict(set)  # client_key -> set of month indices (0..5)

for month_idx, (sheet_name, month_label) in enumerate(MONTHS):
    if sheet_name not in wb_src.sheetnames:
        print(f"  WARNING: sheet '{sheet_name}' not found, skipping")
        continue
    ws = wb_src[sheet_name]
    n = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        tiers = row[5] if len(row) > 5 else None  # column F = index 5 (0-based)
        if tiers is None or str(tiers).strip() == "":
            continue
        ref, name = split_tiers(tiers)
        client_months[(ref, name)].add(month_idx)
        n += 1
    print(f"  {month_label:<10} ({sheet_name}): {n} client occurrences")

wb_src.close()

# --- 2. Build sorted client list -------------------------------------------
# Sort by name (case-insensitive), then ref
clients = sorted(client_months.keys(), key=lambda k: (k[1].upper(), k[0]))
print(f"\nTotal unique clients: {len(clients)}")

# --- 3. Build output workbook ----------------------------------------------
print(f"\nWriting output: {DST}")
from openpyxl import Workbook
wb = Workbook()
wb.properties.creator = "Z.ai"
ws = wb.active
ws.title = "Clients uniques"

# --- Styles ---
# Header style
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Title style
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F4E78")
TITLE_ALIGN = Alignment(horizontal="left", vertical="center")

# Body styles
BODY_FONT = Font(name="Calibri", size=10)
BODY_ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)
BODY_ALIGN_CENTER = Alignment(horizontal="center", vertical="center")

# Green check fill for "purchased" cells
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
GREEN_FONT = Font(name="Calibri", size=12, bold=True, color="006100")
CHECK_SYMBOL = "✓"

# Alternating row banding (very light gray for even rows)
BAND_FILL = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

# Border
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# --- Layout ---
# Row 1: Title
ws.cell(row=1, column=1, value="BELGOCAM SA - Clients uniques (Janvier - Juin 2026) — Check mensuel d'achats").font = TITLE_FONT
ws.cell(row=1, column=1).alignment = TITLE_ALIGN
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
ws.row_dimensions[1].height = 26

# Row 2: empty spacer
ws.row_dimensions[2].height = 6

# Row 3: Headers
HEADERS = ["N°", "Réf. client", "Nom du client", "Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Nb mois d'achat"]
for col_idx, h in enumerate(HEADERS, start=1):
    c = ws.cell(row=3, column=col_idx, value=h)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = HEADER_ALIGN
    c.border = BORDER
ws.row_dimensions[3].height = 32

# Data rows starting at row 4
START_ROW = 4
for i, (ref, name) in enumerate(clients, start=1):
    r = START_ROW + i - 1
    # Banded rows for readability (use light fill for even rows)
    banding = (i % 2 == 0)

    # N°
    c = ws.cell(row=r, column=1, value=i)
    c.font = BODY_FONT
    c.alignment = BODY_ALIGN_CENTER
    c.border = BORDER
    if banding: c.fill = BAND_FILL

    # Réf. client
    c = ws.cell(row=r, column=2, value=ref)
    c.font = BODY_FONT
    c.alignment = BODY_ALIGN_CENTER
    c.border = BORDER
    if banding: c.fill = BAND_FILL

    # Nom du client
    c = ws.cell(row=r, column=3, value=name)
    c.font = BODY_FONT
    c.alignment = BODY_ALIGN_LEFT
    c.border = BORDER
    if banding: c.fill = BAND_FILL

    # Month checks (columns 4..9)
    months_active = client_months.get((ref, name), set())
    for m_idx in range(6):
        col = 4 + m_idx
        c = ws.cell(row=r, column=col)
        c.border = BORDER
        c.alignment = BODY_ALIGN_CENTER
        if m_idx in months_active:
            c.value = CHECK_SYMBOL
            c.fill = GREEN_FILL
            c.font = GREEN_FONT
        else:
            # Empty cell - keep banding if applicable
            if banding:
                c.fill = BAND_FILL
            c.font = BODY_FONT

    # Nb mois d'achat
    c = ws.cell(row=r, column=10, value=len(months_active))
    c.font = Font(name="Calibri", size=10, bold=True)
    c.alignment = BODY_ALIGN_CENTER
    c.border = BORDER
    if banding: c.fill = BAND_FILL

# --- Column widths ---
ws.column_dimensions['A'].width = 6      # N°
ws.column_dimensions['B'].width = 16     # Réf. client
ws.column_dimensions['C'].width = 42     # Nom du client
for col_letter in ['D', 'E', 'F', 'G', 'H', 'I']:
    ws.column_dimensions[col_letter].width = 11
ws.column_dimensions['J'].width = 16     # Nb mois d'achat

# Freeze panes: keep header + N°/ref/name visible
ws.freeze_panes = "D4"

# Auto filter on header row
ws.auto_filter.ref = f"A3:J{START_ROW + len(clients) - 1}"

# Save
wb.save(DST)
print(f"Saved: {DST}")
print(f"File size: {os.path.getsize(DST):,} bytes")

# --- Summary printout ---
print("\nSUMMARY:")
print(f"  Total unique clients: {len(clients)}")
# Distribution of "Nb mois d'achat"
from collections import Counter
dist = Counter(len(client_months.get(k, set())) for k in clients)
print(f"  Distribution of months active per client:")
for n_months in sorted(dist.keys()):
    print(f"    {n_months} mois : {dist[n_months]} clients")
