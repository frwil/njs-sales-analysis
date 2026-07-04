"""
Generate charts (PNG) for the consolidated analysis file.

Charts:
1. Evolution mensuelle du nombre de clients actifs (ciblé vs concentrés)
2. Histogramme : segments de transition Q1->Q2 (ciblé)
3. Histogramme : segments de transition Q1->Q2 (concentrés)
4. Top 10 pertes Q1 (global) en CA
5. Top 10 pertes Q1 (concentrés) en CA
6. Répartition des clients Zero achat global Q1 (20/80 vs non-20/80)
7. Comparatif bilan net Ciblé vs Concentrés

All charts saved as PNG to /home/z/my-project/scripts/charts/
"""
import re
import os
import json
import matplotlib
matplotlib.use("Agg")
import matplotlib.font_manager as fm
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from collections import defaultdict

# Setup fonts - use Liberation Sans (a clean Latin font) as primary
plt.rcParams['font.sans-serif'] = ['Liberation Sans', 'DejaVu Sans']
plt.rcParams['axes.unicode_minus'] = False

SRC = "/home/z/my-project/download/ventes_livrees.xlsx"
CHARTS_DIR = "/home/z/my-project/scripts/charts"
os.makedirs(CHARTS_DIR, exist_ok=True)

# Load excluded clients
with open("/home/z/my-project/scripts/excluded_clients.json", "r", encoding="utf-8") as f:
    excluded_tiers = set(json.load(f))

# Product categories
SOJA_REFS = {"T102", "T1021", "T1023", "T1024"}
CONCENTRE_REFS = {
    "C102", "C1022", "C104", "C1042", "C1043", "C1044",
    "C105", "C1053", "C1054", "C1055", "C101", "C103",
}
TARGET_REFS = SOJA_REFS | CONCENTRE_REFS

WEIGHT_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(KG|Kg|kg|KG|GRAMMES?|G|L)\b", re.IGNORECASE)
def parse_weight_kg(desc):
    if not desc: return 0.0
    matches = WEIGHT_RE.findall(str(desc))
    if not matches: return 0.0
    val_str, unit = matches[-1]
    val = float(val_str)
    u = unit.upper()
    if u == "KG": return val
    if u in ("G","GRAMME","GRAMMES"): return val/1000.0
    if u == "L": return val
    return 0.0

sheet_to_month = {"Sheet 1":1,"Feuil1":2,"Feuil2":3,"Feuil3":4,"Feuil4":5,"Feuil5":6}
Q1_SHEETS = {"Sheet 1","Feuil1","Feuil2"}
Q2_SHEETS = {"Feuil3","Feuil4","Feuil5"}
MONTH_NAMES = ["", "Jan", "Fev", "Mar", "Avr", "Mai", "Juin"]

TIERS_RE = re.compile(r"^\s*([A-Z]{2,}[\w-]*?)\s*-\s+(.+?)\s*$")
def split_tiers(tiers):
    if tiers is None: return ("","")
    s = str(tiers).strip()
    m = TIERS_RE.match(s)
    if m: return (m.group(1).strip(), m.group(2).strip())
    return ("", s)

# Read source
print("Reading source...")
wb = load_workbook(SRC, read_only=True, data_only=True)

product_weight = {}
client_ca_total = defaultdict(float)
client_months_any = defaultdict(set)
# Per-month active clients count (for chart 1)
month_active_clients_target = defaultdict(set)  # month -> set of client keys
month_active_clients_concentre = defaultdict(set)
month_active_clients_any = defaultdict(set)

# Per client aggregates (targeted + concentré)
client_t_q1_kg = defaultdict(float)
client_t_q2_kg = defaultdict(float)
client_t_q1_ca = defaultdict(float)
client_t_q2_ca = defaultdict(float)
client_t_q1_months = defaultdict(set)
client_t_q2_months = defaultdict(set)
client_c_q1_kg = defaultdict(float)
client_c_q2_kg = defaultdict(float)
client_c_q1_ca = defaultdict(float)
client_c_q2_ca = defaultdict(float)
client_c_q1_months = defaultdict(set)
client_c_q2_months = defaultdict(set)

client_q2_by_product_target = defaultdict(lambda: defaultdict(lambda: {"vol_kg":0.0, "ca":0.0, "months":set()}))
client_q2_by_product_concentre = defaultdict(lambda: defaultdict(lambda: {"vol_kg":0.0, "ca":0.0, "months":set()}))
client_q2_active_months_target = defaultdict(set)
client_q2_active_months_concentre = defaultdict(set)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    is_q1 = sheet_name in Q1_SHEETS
    is_q2 = sheet_name in Q2_SHEETS
    if not (is_q1 or is_q2): continue
    month_num = sheet_to_month.get(sheet_name)
    if month_num is None: continue
    for row in ws.iter_rows(min_row=3, values_only=True):
        tiers = row[5] if len(row)>5 else None
        if tiers is None or str(tiers).strip()=="": continue
        if tiers in excluded_tiers: continue  # EXCLUDE internal clients
        ref_prod = row[0] if len(row)>0 else None
        desc = row[1] if len(row)>1 else None
        qte = row[2] if len(row)>2 else 0
        ca_ht = row[8] if len(row)>8 else 0
        ref_c, name_c = split_tiers(tiers)
        key = (ref_c, name_c)
        ref_prod_str = str(ref_prod).strip() if ref_prod is not None else ""
        if ref_prod_str and ref_prod_str not in product_weight:
            product_weight[ref_prod_str] = parse_weight_kg(desc)
        weight_kg = product_weight.get(ref_prod_str, 0.0)
        try: qte_f = float(qte) if qte is not None else 0.0
        except: qte_f = 0.0
        try: ca_f = float(ca_ht) if ca_ht is not None else 0.0
        except: ca_f = 0.0
        client_ca_total[key] += ca_f
        client_months_any[key].add(month_num)
        month_active_clients_any[month_num].add(key)
        is_target = ref_prod_str in TARGET_REFS
        is_concentre = ref_prod_str in CONCENTRE_REFS
        vol_kg = qte_f * weight_kg
        if is_target:
            month_active_clients_target[month_num].add(key)
            if is_q1:
                client_t_q1_kg[key] += vol_kg
                client_t_q1_ca[key] += ca_f
                client_t_q1_months[key].add(month_num)
            else:
                client_t_q2_kg[key] += vol_kg
                client_t_q2_ca[key] += ca_f
                client_t_q2_months[key].add(month_num)
                rec = client_q2_by_product_target[key][ref_prod_str]
                rec["vol_kg"] += vol_kg
                rec["ca"] += ca_f
                if month_num: rec["months"].add(month_num)
                if month_num: client_q2_active_months_target[key].add(month_num)
        if is_concentre:
            month_active_clients_concentre[month_num].add(key)
            if is_q1:
                client_c_q1_kg[key] += vol_kg
                client_c_q1_ca[key] += ca_f
                client_c_q1_months[key].add(month_num)
            else:
                client_c_q2_kg[key] += vol_kg
                client_c_q2_ca[key] += ca_f
                client_c_q2_months[key].add(month_num)
                rec = client_q2_by_product_concentre[key][ref_prod_str]
                rec["vol_kg"] += vol_kg
                rec["ca"] += ca_f
                if month_num: rec["months"].add(month_num)
                if month_num: client_q2_active_months_concentre[key].add(month_num)

wb.close()

clients_all = sorted(client_ca_total.keys(), key=lambda k:(k[1].upper(), k[0]))
print(f"Total unique clients (after exclusion): {len(clients_all)}")

# Pareto
total_ca = sum(client_ca_total.values())
sorted_by_ca = sorted(clients_all, key=lambda k:-client_ca_total[k])
pareto_threshold = 0.80 * total_ca
pareto_clients = set()
cumulative = 0.0
for k in sorted_by_ca:
    cumulative += client_ca_total[k]
    pareto_clients.add(k)
    if cumulative >= pareto_threshold: break

# Zero achat
zero_global = [k for k in clients_all if not client_t_q1_months[k]]
zero_concentre = [k for k in clients_all if not client_c_q1_months[k]]

# Compute losses
def compute_loss(client_key, category_refs, products_dict, active_months_dict):
    products_q2 = products_dict.get(client_key, {})
    active_months_set = active_months_dict.get(client_key, set())
    n_active = len(active_months_set)
    frequency = n_active / 3.0
    sum_avg = 0.0; loss_t = 0.0; loss_c = 0.0
    for ref, rec in products_q2.items():
        if ref not in category_refs: continue
        total_t = rec["vol_kg"]/1000.0
        n_m = len(rec["months"])
        if n_m == 0: continue
        m_avg = total_t / n_m
        sum_avg += m_avg
        loss_t += m_avg * frequency * 3
        price = (rec["ca"]/total_t) if total_t > 0 else 0
        loss_c += m_avg * frequency * 3 * price
    return {"loss_tons": loss_t, "loss_fcfa": loss_c, "sum_avg": sum_avg,
            "frequency": frequency, "n_active": n_active, "n_products": len([1 for r in products_q2.values() if True])}

losses_global = {k: compute_loss(k, TARGET_REFS, client_q2_by_product_target, client_q2_active_months_target) for k in zero_global}
losses_concentre = {k: compute_loss(k, CONCENTRE_REFS, client_q2_by_product_concentre, client_q2_active_months_concentre) for k in zero_concentre}

# Transition segments
def get_seg(q1m, q2m):
    return ("Active Q1" if len(q1m)>0 else "Zero Q1", "Active Q2" if len(q2m)>0 else "Zero Q2")
client_seg_target = {k: get_seg(client_t_q1_months[k], client_t_q2_months[k]) for k in clients_all}
client_seg_concentre = {k: get_seg(client_c_q1_months[k], client_c_q2_months[k]) for k in clients_all}

# ===== CHARTS =====
print("\nGenerating charts...")

# Color palette
COLORS = {
    "primary": "#1F4E78",
    "secondary": "#2E75B6",
    "accent": "#FFC000",
    "success": "#70AD47",
    "danger": "#C00000",
    "warning": "#ED7D31",
    "gray": "#A6A6A6",
    "light_blue": "#DDEBF7",
    "light_green": "#C6EFCE",
    "light_red": "#FCE4E4",
    "light_gray": "#D9D9D9",
}

# ---- Chart 1: Evolution mensuelle du nombre de clients actifs ----
fig, ax = plt.subplots(figsize=(10, 5.5), constrained_layout=True)
months = [1,2,3,4,5,6]
counts_any = [len(month_active_clients_any[m]) for m in months]
counts_target = [len(month_active_clients_target[m]) for m in months]
counts_concentre = [len(month_active_clients_concentre[m]) for m in months]
x = np.arange(len(months))
width = 0.27
b1 = ax.bar(x - width, counts_any, width, label="Tous produits", color=COLORS["primary"])
b2 = ax.bar(x, counts_target, width, label="Ciblé (16 produits)", color=COLORS["accent"])
b3 = ax.bar(x + width, counts_concentre, width, label="Concentrés (12 produits)", color=COLORS["warning"])
ax.set_xticks(x)
ax.set_xticklabels([MONTH_NAMES[m] for m in months])
ax.set_ylabel("Nombre de clients actifs")
ax.set_title("Évolution mensuelle du nombre de clients actifs", fontsize=13, fontweight="bold", color=COLORS["primary"])
ax.legend(loc="lower right", framealpha=0.95)
ax.grid(axis="y", alpha=0.3, linestyle="--")
for bars in [b1, b2, b3]:
    for bar in bars:
        h = bar.get_height()
        ax.text(bar.get_x() + bar.get_width()/2., h + 20, f'{int(h)}', ha='center', va='bottom', fontsize=8)
ax.set_ylim(0, max(counts_any) * 1.15)
plt.savefig(f"{CHARTS_DIR}/chart1_evolution.png", dpi=120)
plt.close()
print("  ✓ chart1_evolution.png")

# ---- Chart 2: Segments de transition Q1->Q2 (ciblé) ----
seg_counts_target = {
    "Q1 Zero → Q2 Zero\n(persistant)": sum(1 for k in clients_all if client_seg_target[k]==("Zero Q1","Zero Q2")),
    "Q1 Zero → Q2 Active\n(réactivé)": sum(1 for k in clients_all if client_seg_target[k]==("Zero Q1","Active Q2")),
    "Q1 Active → Q2 Active\n(retenu)": sum(1 for k in clients_all if client_seg_target[k]==("Active Q1","Active Q2")),
    "Q1 Active → Q2 Zero\n(churned)": sum(1 for k in clients_all if client_seg_target[k]==("Active Q1","Zero Q2")),
}
fig, ax = plt.subplots(figsize=(10, 5.5), constrained_layout=True)
labels = list(seg_counts_target.keys())
values = list(seg_counts_target.values())
colors_seg = [COLORS["light_gray"], COLORS["success"], COLORS["secondary"], COLORS["danger"]]
bars = ax.bar(labels, values, color=colors_seg, edgecolor="white", linewidth=1.5)
ax.set_ylabel("Nombre de clients")
ax.set_title("Segments de transition Q1 → Q2 (16 produits ciblés)", fontsize=13, fontweight="bold", color=COLORS["primary"])
ax.grid(axis="y", alpha=0.3, linestyle="--")
for bar, val in zip(bars, values):
    h = bar.get_height()
    ax.text(bar.get_x() + bar.get_width()/2., h + 10, f'{val}', ha='center', va='bottom', fontsize=11, fontweight="bold")
ax.set_ylim(0, max(values) * 1.15)
plt.savefig(f"{CHARTS_DIR}/chart2_segments_cible.png", dpi=120)
plt.close()
print("  ✓ chart2_segments_cible.png")

# ---- Chart 3: Segments de transition Q1->Q2 (concentrés) ----
seg_counts_concentre = {
    "Q1 Zero → Q2 Zero\n(persistant)": sum(1 for k in clients_all if client_seg_concentre[k]==("Zero Q1","Zero Q2")),
    "Q1 Zero → Q2 Active\n(réactivé)": sum(1 for k in clients_all if client_seg_concentre[k]==("Zero Q1","Active Q2")),
    "Q1 Active → Q2 Active\n(retenu)": sum(1 for k in clients_all if client_seg_concentre[k]==("Active Q1","Active Q2")),
    "Q1 Active → Q2 Zero\n(churned)": sum(1 for k in clients_all if client_seg_concentre[k]==("Active Q1","Zero Q2")),
}
fig, ax = plt.subplots(figsize=(10, 5.5), constrained_layout=True)
labels = list(seg_counts_concentre.keys())
values = list(seg_counts_concentre.values())
bars = ax.bar(labels, values, color=colors_seg, edgecolor="white", linewidth=1.5)
ax.set_ylabel("Nombre de clients")
ax.set_title("Segments de transition Q1 → Q2 (12 concentrés)", fontsize=13, fontweight="bold", color=COLORS["primary"])
ax.grid(axis="y", alpha=0.3, linestyle="--")
for bar, val in zip(bars, values):
    h = bar.get_height()
    ax.text(bar.get_x() + bar.get_width()/2., h + 10, f'{val}', ha='center', va='bottom', fontsize=11, fontweight="bold")
ax.set_ylim(0, max(values) * 1.15)
plt.savefig(f"{CHARTS_DIR}/chart3_segments_concentre.png", dpi=120)
plt.close()
print("  ✓ chart3_segments_concentre.png")

# ---- Chart 4: Top 10 pertes Q1 (ciblé) en CA ----
top10_global = sorted(zero_global, key=lambda k: -losses_global[k]["loss_fcfa"])[:10]
names_g = [k[1][:25] for k in top10_global]
values_g = [losses_global[k]["loss_fcfa"]/1e6 for k in top10_global]  # in millions
fig, ax = plt.subplots(figsize=(11, 6), constrained_layout=True)
y_pos = np.arange(len(names_g))
bars = ax.barh(y_pos, values_g, color=COLORS["danger"], edgecolor="white", linewidth=1)
ax.set_yticks(y_pos)
ax.set_yticklabels(names_g, fontsize=9)
ax.invert_yaxis()
ax.set_xlabel("Perte Q1 estimée (millions FCFA)")
ax.set_title("Top 10 pertes Q1 — Clients zéro achat global (16 produits)", fontsize=12, fontweight="bold", color=COLORS["primary"])
ax.grid(axis="x", alpha=0.3, linestyle="--")
for bar, val in zip(bars, values_g):
    ax.text(val + max(values_g)*0.01, bar.get_y() + bar.get_height()/2., f'{val:.1f} M', va='center', fontsize=9)
ax.set_xlim(0, max(values_g) * 1.15)
plt.savefig(f"{CHARTS_DIR}/chart4_top10_pertes_global.png", dpi=120)
plt.close()
print("  ✓ chart4_top10_pertes_global.png")

# ---- Chart 5: Top 10 pertes Q1 (concentrés) en CA ----
top10_conc = sorted(zero_concentre, key=lambda k: -losses_concentre[k]["loss_fcfa"])[:10]
names_c = [k[1][:25] for k in top10_conc]
values_c = [losses_concentre[k]["loss_fcfa"]/1e6 for k in top10_conc]
fig, ax = plt.subplots(figsize=(11, 6), constrained_layout=True)
y_pos = np.arange(len(names_c))
bars = ax.barh(y_pos, values_c, color=COLORS["warning"], edgecolor="white", linewidth=1)
ax.set_yticks(y_pos)
ax.set_yticklabels(names_c, fontsize=9)
ax.invert_yaxis()
ax.set_xlabel("Perte Q1 estimée (millions FCFA)")
ax.set_title("Top 10 pertes Q1 — Clients zéro achat concentrés (12 produits)", fontsize=12, fontweight="bold", color=COLORS["primary"])
ax.grid(axis="x", alpha=0.3, linestyle="--")
for bar, val in zip(bars, values_c):
    ax.text(val + max(values_c)*0.01, bar.get_y() + bar.get_height()/2., f'{val:.1f} M', va='center', fontsize=9)
ax.set_xlim(0, max(values_c) * 1.15)
plt.savefig(f"{CHARTS_DIR}/chart5_top10_pertes_concentre.png", dpi=120)
plt.close()
print("  ✓ chart5_top10_pertes_concentre.png")

# ---- Chart 6: Répartition Zero achat Q1 - 20/80 vs non ----
fig, axes = plt.subplots(1, 2, figsize=(11, 5.5), constrained_layout=True)
# Subplot 1: Zero achat global
za_pareto_g = sum(1 for k in zero_global if k in pareto_clients)
za_nonpareto_g = len(zero_global) - za_pareto_g
axes[0].pie([za_pareto_g, za_nonpareto_g],
            labels=[f"20/80\n({za_pareto_g} clients)", f"Autres\n({za_nonpareto_g} clients)"],
            colors=[COLORS["accent"], COLORS["light_gray"]],
            autopct='%1.1f%%', startangle=90, textprops={'fontsize': 10})
axes[0].set_title(f"Zero achat Q1 (ciblé)\n{len(zero_global)} clients", fontsize=11, fontweight="bold", color=COLORS["primary"])

# Subplot 2: Zero achat concentrés
za_pareto_c = sum(1 for k in zero_concentre if k in pareto_clients)
za_nonpareto_c = len(zero_concentre) - za_pareto_c
axes[1].pie([za_pareto_c, za_nonpareto_c],
            labels=[f"20/80\n({za_pareto_c} clients)", f"Autres\n({za_nonpareto_c} clients)"],
            colors=[COLORS["accent"], COLORS["light_gray"]],
            autopct='%1.1f%%', startangle=90, textprops={'fontsize': 10})
axes[1].set_title(f"Zero achat Q1 (concentrés)\n{len(zero_concentre)} clients", fontsize=11, fontweight="bold", color=COLORS["primary"])

plt.savefig(f"{CHARTS_DIR}/chart6_zero_achat_2080.png", dpi=120)
plt.close()
print("  ✓ chart6_zero_achat_2080.png")

# ---- Chart 7: Bilan net Ciblé vs Concentrés ----
gain_cible = sum(client_t_q2_ca[k] for k in clients_all if client_seg_target[k]==("Zero Q1","Active Q2"))
loss_cible = sum(client_t_q1_ca[k] for k in clients_all if client_seg_target[k]==("Active Q1","Zero Q2"))
gain_conc = sum(client_c_q2_ca[k] for k in clients_all if client_seg_concentre[k]==("Zero Q1","Active Q2"))
loss_conc = sum(client_c_q1_ca[k] for k in clients_all if client_seg_concentre[k]==("Active Q1","Zero Q2"))

fig, ax = plt.subplots(figsize=(10, 5.5), constrained_layout=True)
categories = ["Ciblé (16 produits)", "Concentrés (12 produits)"]
gains = [gain_cible/1e6, gain_conc/1e6]
losses = [loss_cible/1e6, loss_conc/1e6]
nets = [(gain_cible-loss_cible)/1e6, (gain_conc-loss_conc)/1e6]
x = np.arange(len(categories))
width = 0.27
b1 = ax.bar(x - width, gains, width, label="Gain (réactivés)", color=COLORS["success"])
b2 = ax.bar(x, losses, width, label="Perte (churned)", color=COLORS["danger"])
b3 = ax.bar(x + width, nets, width, label="Bilan net", color=COLORS["primary"])
ax.set_xticks(x)
ax.set_xticklabels(categories)
ax.set_ylabel("Montant (millions FCFA)")
ax.set_title("Bilan net Q1 → Q2 : Ciblé vs Concentrés", fontsize=13, fontweight="bold", color=COLORS["primary"])
ax.legend(loc="upper right")
ax.grid(axis="y", alpha=0.3, linestyle="--")
ax.axhline(y=0, color="black", linewidth=0.8)
for bars in [b1, b2, b3]:
    for bar in bars:
        h = bar.get_height()
        offset = 10 if h >= 0 else -25
        ax.text(bar.get_x() + bar.get_width()/2., h + offset, f'{h:.0f}', ha='center', va='bottom' if h>=0 else 'top', fontsize=9)
plt.savefig(f"{CHARTS_DIR}/chart7_bilan_net.png", dpi=120)
plt.close()
print("  ✓ chart7_bilan_net.png")

print(f"\nAll charts saved to: {CHARTS_DIR}")
print("Done.")
