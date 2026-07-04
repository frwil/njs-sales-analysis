"""Quick verification: read back the output file and show first few rows + stats."""
from openpyxl import load_workbook
from collections import Counter

DST = "/home/z/my-project/download/Clients uniques - check mensuel.xlsx"
wb = load_workbook(DST, data_only=True)
ws = wb.active

print(f"Sheet: {ws.title}  (max_row={ws.max_row}, max_col={ws.max_column})")
print(f"\n--- Row 1 (title) ---")
print(f"  A1: {ws['A1'].value}")
print(f"\n--- Row 3 (headers) ---")
headers = [ws.cell(row=3, column=c).value for c in range(1, 11)]
print(f"  {headers}")

print(f"\n--- First 10 clients ---")
for r in range(4, 14):
    row = [ws.cell(row=r, column=c).value for c in range(1, 11)]
    print(f"  {row}")

print(f"\n--- Last 5 clients ---")
for r in range(ws.max_row - 4, ws.max_row + 1):
    row = [ws.cell(row=r, column=c).value for c in range(1, 11)]
    print(f"  {row}")

# Verify counts: total data rows = max_row - 3
total_clients = ws.max_row - 3
print(f"\nTotal client rows: {total_clients}")

# Distribution of "Nb mois d'achat"
dist = Counter()
for r in range(4, ws.max_row + 1):
    nb = ws.cell(row=r, column=10).value
    dist[nb] += 1
print("Distribution 'Nb mois d'achat':")
for k in sorted(dist.keys()):
    print(f"  {k} mois : {dist[k]} clients")

wb.close()
