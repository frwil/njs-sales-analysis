"""
Inspect product list: find all unique (Réf. produit, Description) pairs
that contain 'SOJA', 'BELGO 10%', or 'BELGO 5%' in their description.
"""
import re
from openpyxl import load_workbook
from collections import Counter

SRC = "/home/z/my-project/download/ventes janv a juin 2026 - Livrées uniquement.xlsx"

wb = load_workbook(SRC, read_only=True, data_only=True)

all_products = Counter()  # (ref, desc) -> total rows
soja_products = Counter()
belgo10_products = Counter()
belgo5_products = Counter()
other_belgo = Counter()

# Patterns
SOJA_RE = re.compile(r"SOJA", re.IGNORECASE)
BELGO10_RE = re.compile(r"BELGO\s*10\s*%", re.IGNORECASE)
BELGO5_RE = re.compile(r"BELGO\s*5\s*%", re.IGNORECASE)
BELGO_RE = re.compile(r"BELGO", re.IGNORECASE)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for row in ws.iter_rows(min_row=3, values_only=True):
        ref = row[0] if len(row) > 0 else None       # Réf. produit
        desc = row[1] if len(row) > 1 else None      # Description
        if ref is None and desc is None:
            continue
        key = (ref, desc)
        all_products[key] += 1
        desc_str = str(desc) if desc else ""
        if SOJA_RE.search(desc_str):
            soja_products[key] += 1
        if BELGO10_RE.search(desc_str):
            belgo10_products[key] += 1
        if BELGO5_RE.search(desc_str):
            belgo5_products[key] += 1
        elif BELGO_RE.search(desc_str) and not BELGO10_RE.search(desc_str) and not BELGO5_RE.search(desc_str):
            other_belgo[key] += 1

wb.close()

print("="*80)
print(f"TOUS LES PRODUITS UNIQUES ({len(all_products)} produits):")
print("="*80)
for (ref, desc), n in sorted(all_products.items(), key=lambda x: -x[1]):
    print(f"  {str(ref):<8} | {str(desc):<55} | {n} lignes")

print(f"\n{'='*80}")
print(f"PRODUITS 'SOJA' ({len(soja_products)} produits):")
print("="*80)
for (ref, desc), n in sorted(soja_products.items(), key=lambda x: -x[1]):
    print(f"  {str(ref):<8} | {str(desc):<55} | {n} lignes")

print(f"\n{'='*80}")
print(f"PRODUITS 'BELGO 10%' ({len(belgo10_products)} produits):")
print("="*80)
for (ref, desc), n in sorted(belgo10_products.items(), key=lambda x: -x[1]):
    print(f"  {str(ref):<8} | {str(desc):<55} | {n} lignes")

print(f"\n{'='*80}")
print(f"PRODUITS 'BELGO 5%' ({len(belgo5_products)} produits):")
print("="*80)
for (ref, desc), n in sorted(belgo5_products.items(), key=lambda x: -x[1]):
    print(f"  {str(ref):<8} | {str(desc):<55} | {n} lignes")

print(f"\n{'='*80}")
print(f"AUTRES PRODUITS 'BELGO' (non 10%, non 5%) ({len(other_belgo)} produits):")
print("="*80)
for (ref, desc), n in sorted(other_belgo.items(), key=lambda x: -x[1]):
    print(f"  {str(ref):<8} | {str(desc):<55} | {n} lignes")
