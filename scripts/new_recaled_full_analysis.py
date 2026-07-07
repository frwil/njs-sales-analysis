"""
Complete analysis: new recalibrated S2 objectives vs old vs forecast, CA impact.
"""
import json
from openpyxl import load_workbook
from collections import defaultdict

# ===== LOAD NEW RECALIBRATED OBJECTIVES =====
SRC_NEW = "/home/z/my-project/upload/DOC-20260706-WA0017.xlsx"
wb = load_workbook(SRC_NEW, data_only=True)
ws = wb.active

CATEGORIES = ['ALIMENT COMPLET', 'INGREDIENTS', 'COMPLEMENT ALIMENTAIRE', 'MATERIEL ELEVAGE',
              'PREMIX', 'CONCENTRES', 'TOURTEAUX', 'ALVEOLE', 'DIVERS']
MONTHS_S2 = [7, 8, 9, 10, 11, 12]

# Extract new S2 objectives
new_s2 = {cat: {m: 0 for m in MONTHS_S2} for cat in CATEGORIES}
current_agence = None
for r in range(4, ws.max_row + 1):
    a = ws.cell(row=r, column=1).value
    b = ws.cell(row=r, column=2).value
    if a:
        current_agence = str(a).strip()
    if b and current_agence:
        cat = str(b).strip()
        if cat in CATEGORIES:
            for j, m in enumerate(MONTHS_S2):
                v = ws.cell(row=r, column=5 + j).value
                if v:
                    new_s2[cat][m] += float(v)
wb.close()

# ===== LOAD OLD DATA =====
with open('/home/z/my-project/scripts/objectives_comparison.json', 'r') as f:
    old_data = json.load(f)
with open('/home/z/my-project/scripts/s2_recaled_objectives.json', 'r') as f:
    old_recaled = json.load(f)
with open('/home/z/my-project/scripts/yoy_forecast.json', 'r') as f:
    yoy_data = json.load(f)
with open('/home/z/my-project/scripts/ventes_objectifs_data.json', 'r') as f:
    sales_data = json.load(f)

# Prices per ton (from S1 2026 sales)
prices = sales_data['price_per_ton']

# ===== COMPUTE EVERYTHING =====
print("="*90)
print("ANALYSE COMPLÈTE : NOUVEAU RECALIBRAGE S2 vs ANCIEN vs FORECAST")
print("="*90)

# 1. Volume comparison
print("\n1. VOLUME S2 (tonnes)")
print(f"{'Catégorie':<28}{'Avant recaled':>15}{'Ancien recaled':>16}{'Nouveau recaled':>17}{'Forecast YoY':>14}")
print("-"*90)
old_totals = {}
prev_totals = {}
new_totals = {}
for cat in CATEGORIES:
    avant = sum(old_data['global_objectives'].get(cat, {}).get(str(m), 0) for m in MONTHS_S2)
    ancien = sum(old_recaled['global_s2_recaled'].get(cat, {}).get(str(m), 0) for m in MONTHS_S2)
    nouveau = sum(new_s2[cat][m] for m in MONTHS_S2)
    forecast = yoy_data.get('forecast_refined', {}).get(cat, {}).get('real', 0)
    old_totals[cat] = avant
    prev_totals[cat] = ancien
    new_totals[cat] = nouveau
    print(f"{cat:<28}{avant:>15.1f}{ancien:>16.1f}{nouveau:>17.1f}{forecast:>14.1f}")

tot_avant = sum(old_totals.values())
tot_ancien = sum(prev_totals.values())
tot_nouveau = sum(new_totals.values())
tot_forecast = sum(yoy_data.get('forecast_refined', {}).get(c, {}).get('real', 0) for c in CATEGORIES)
print("-"*90)
print(f"{'TOTAL':<28}{tot_avant:>15.1f}{tot_ancien:>16.1f}{tot_nouveau:>17.1f}{tot_forecast:>14.1f}")

# 2. CA comparison
print("\n2. CA HT S2 (millions FCFA) — basé sur prix moyen S1 2026")
print(f"{'Catégorie':<28}{'Prix/t':>10}{'CA avant':>12}{'CA anc.rec':>12}{'CA nouv.rec':>12}{'CA forecast':>13}")
print("-"*87)
ca_avant = {}
ca_ancien = {}
ca_nouveau = {}
ca_forecast = {}
for cat in CATEGORIES:
    price = prices.get(cat, 0)
    ca_av = old_totals[cat] * price
    ca_an = prev_totals[cat] * price
    ca_no = new_totals[cat] * price
    ca_fc = yoy_data.get('forecast_refined', {}).get(cat, {}).get('real', 0) * price
    ca_avant[cat] = ca_av
    ca_ancien[cat] = ca_an
    ca_nouveau[cat] = ca_no
    ca_forecast[cat] = ca_fc
    print(f"{cat:<28}{price:>10,.0f}{ca_av/1e6:>12.1f}{ca_an/1e6:>12.1f}{ca_no/1e6:>12.1f}{ca_fc/1e6:>13.1f}")

tot_ca_avant = sum(ca_avant.values())
tot_ca_ancien = sum(ca_ancien.values())
tot_ca_nouveau = sum(ca_nouveau.values())
tot_ca_forecast = sum(ca_forecast.values())
print("-"*87)
print(f"{'TOTAL':<28}{'':>10}{tot_ca_avant/1e6:>12.1f}{tot_ca_ancien/1e6:>12.1f}{tot_ca_nouveau/1e6:>12.1f}{tot_ca_forecast/1e6:>13.1f}")

# 3. Impact du nouveau recalibrage
print("\n3. INCIDENCE DU NOUVEAU RECALIBRAGE SUR LE CA")
print(f"{'Catégorie':<28}{'Écart vol (t)':>15}{'Écart CA (M)':>14}{'% écart CA':>12}")
print("-"*69)
for cat in CATEGORIES:
    ecart_vol = new_totals[cat] - old_totals[cat]
    ecart_ca = (ca_nouveau[cat] - ca_avant[cat]) / 1e6
    pct = ((ca_nouveau[cat] / ca_avant[cat]) - 1) * 100 if ca_avant[cat] > 0 else 0
    if abs(ecart_ca) > 0.1:
        print(f"{cat:<28}{ecart_vol:>+15.1f}{ecart_ca:>+14.1f}{pct:>+11.1f}%")

ecart_total_vol = tot_nouveau - tot_avant
ecart_total_ca = (tot_ca_nouveau - tot_ca_avant) / 1e6
pct_total = ((tot_ca_nouveau / tot_ca_avant) - 1) * 100 if tot_ca_avant > 0 else 0
print("-"*69)
print(f"{'TOTAL':<28}{ecart_total_vol:>+15.1f}{ecart_total_ca:>+14.1f}{pct_total:>+11.1f}%")

# 4. Focus CONCENTRÉS + TOURTEAUX
print("\n4. FOCUS CONCENTRÉS + TOURTEAUX (impact CA du recalibrage)")
print(f"{'':>28}{'Avant':>12}{'Nouveau':>12}{'Écart CA':>12}{'Écart %':>10}")
print("-"*74)
for cat in ['CONCENTRES', 'TOURTEAUX']:
    print(f"{cat:<28}{ca_avant[cat]/1e6:>12.1f}{ca_nouveau[cat]/1e6:>12.1f}{(ca_nouveau[cat]-ca_avant[cat])/1e6:>+12.1f}{(ca_nouveau[cat]/ca_avant[cat]-1)*100:>+9.1f}%")
bilan_conc_tour = ((ca_nouveau['CONCENTRES'] - ca_avant['CONCENTRES']) + (ca_nouveau['TOURTEAUX'] - ca_avant['TOURTEAUX'])) / 1e6
print(f"{'Bilan net Conc+Tour':<28}{'':>12}{'':>12}{bilan_conc_tour:>+12.1f}{'':>10}")

# 5. Forecast vs nouveau recalibrage
print("\n5. FORECAST YOY (RÉALISTE) vs NOUVEAU RECALIBRÉ")
print(f"{'Catégorie':<28}{'Forecast (t)':>14}{'Nouv.obj (t)':>14}{'Atteinte %':>12}{'CA forecast':>13}{'CA nouv.obj':>13}")
print("-"*94)
for cat in CATEGORIES:
    fc = yoy_data.get('forecast_refined', {}).get(cat, {}).get('real', 0)
    obj = new_totals[cat]
    pct = (fc / obj * 100) if obj > 0 else 0
    print(f"{cat:<28}{fc:>14.1f}{obj:>14.1f}{pct:>11.1f}%{ca_forecast[cat]/1e6:>13.1f}{ca_nouveau[cat]/1e6:>13.1f}")
print("-"*94)
print(f"{'TOTAL':<28}{tot_forecast:>14.1f}{tot_nouveau:>14.1f}{tot_forecast/tot_nouveau*100:>11.1f}%{tot_ca_forecast/1e6:>13.1f}{tot_ca_nouveau/1e6:>13.1f}")

# 6. CONCENTRÉS - boucler +5% à +10% vs 2025
print("\n6. CONCENTRÉS — PEUT-ON BOUCLER +5% À +10% VS 2025 ?")
with open('/home/z/my-project/scripts/concentres_deepdive.json', 'r') as f:
    conc_data = json.load(f)
total_2025_conc = sum(conc_data['subcat_2025'].get(sc, {}).get(str(m), {}).get('vol_t', 0)
                      for sc in conc_data['subcat_2025'] for m in range(1, 13))
s1_2026_conc = sum(conc_data['subcat_2026'].get(sc, {}).get(str(m), {}).get('vol_t', 0)
                   for sc in conc_data['subcat_2026'] for m in range(1, 7))
s2_2025_conc = sum(conc_data['subcat_2025'].get(sc, {}).get(str(m), {}).get('vol_t', 0)
                   for sc in conc_data['subcat_2025'] for m in range(7, 13))
forecast_conc = yoy_data['forecast_refined']['CONCENTRES']['real']
new_obj_conc = new_totals['CONCENTRES']
price_conc = prices['CONCENTRES']

print(f"  2025 total CONCENTRÉS      : {total_2025_conc:>10,.1f} t  (CA = {total_2025_conc*price_conc/1e6:>8,.1f} M)")
print(f"  S1 2026 réel               : {s1_2026_conc:>10,.1f} t  ({(s1_2026_conc/sum(conc_data['subcat_2025'].get(sc,{}).get(str(m),{}).get('vol_t',0) for sc in conc_data['subcat_2025'] for m in range(1,7))-1)*100:+.1f}% YoY)")
print(f"  S2 2025                    : {s2_2025_conc:>10,.1f} t")
print(f"  Forecast S2 2026 réaliste  : {forecast_conc:>10,.1f} t  ({(forecast_conc/s2_2025_conc-1)*100:+.1f}% vs S2 2025)")
print(f"  Nouvel objectif S2 2026    : {new_obj_conc:>10,.1f} t  ({(new_obj_conc/s2_2025_conc-1)*100:+.1f}% vs S2 2025)")
print()
target_5 = total_2025_conc * 1.05
target_10 = total_2025_conc * 1.10
s2_needed_5 = target_5 - s1_2026_conc
s2_needed_10 = target_10 - s1_2026_conc
print(f"  Cible +5%  → total 2026 = {target_5:>10,.1f} t → S2 nécessaire = {s2_needed_5:>10,.1f} t")
print(f"  Cible +10% → total 2026 = {target_10:>10,.1f} t → S2 nécessaire = {s2_needed_10:>10,.1f} t")
print()
print(f"  vs forecast réaliste ({forecast_conc:,.1f} t) :")
print(f"    Pour +5%  : manque de {s2_needed_5 - forecast_conc:>+10,.1f} t ({(s2_needed_5 - forecast_conc)*price_conc/1e6:>+8,.1f} M FCFA)")
print(f"    Pour +10% : manque de {s2_needed_10 - forecast_conc:>+10,.1f} t ({(s2_needed_10 - forecast_conc)*price_conc/1e6:>+8,.1f} M FCFA)")
print()
print(f"  vs nouvel objectif recalibré ({new_obj_conc:,.1f} t) :")
print(f"    Pour +5%  : il faut {s2_needed_5 - new_obj_conc:>+10,.1f} t au-delà du nouvel objectif")
print(f"    Pour +10% : il faut {s2_needed_10 - new_obj_conc:>+10,.1f} t au-delà du nouvel objectif")
print()

# Scenarios YoY
total_2026_pess = s1_2026_conc + yoy_data['forecast_refined']['CONCENTRES']['pess']
total_2026_real = s1_2026_conc + forecast_conc
total_2026_opt = s1_2026_conc + yoy_data['forecast_refined']['CONCENTRES']['opt']
print(f"  Scénarios YoY full year 2026 vs 2025 ({total_2025_conc:,.1f} t) :")
print(f"    🔴 Pessimiste : {total_2026_pess:>10,.1f} t → YoY = {(total_2026_pess/total_2025_conc-1)*100:+.1f}%")
print(f"    🟡 Réaliste   : {total_2026_real:>10,.1f} t → YoY = {(total_2026_real/total_2025_conc-1)*100:+.1f}%")
print(f"    🟢 Optimiste  : {total_2026_opt:>10,.1f} t → YoY = {(total_2026_opt/total_2025_conc-1)*100:+.1f}%")

# Save
output = {
    'new_s2_objectives': {cat: {str(m): new_s2[cat][m] for m in MONTHS_S2} for cat in CATEGORIES},
    'ca_comparison': {
        'avant': {cat: ca_avant[cat] for cat in CATEGORIES},
        'nouveau': {cat: ca_nouveau[cat] for cat in CATEGORIES},
        'forecast': {cat: ca_forecast[cat] for cat in CATEGORIES},
    },
    'totals': {
        'vol_avant': tot_avant,
        'vol_nouveau': tot_nouveau,
        'vol_forecast': tot_forecast,
        'ca_avant': tot_ca_avant,
        'ca_nouveau': tot_ca_nouveau,
        'ca_forecast': tot_ca_forecast,
    }
}
with open('/home/z/my-project/scripts/new_recaled_analysis.json', 'w', encoding='utf-8') as f:
    json.dump(output, f, ensure_ascii=False, indent=2, default=str)
