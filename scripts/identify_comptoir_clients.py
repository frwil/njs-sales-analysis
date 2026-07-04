"""
Identify all CLIENT COMPTOIR entries to exclude from analysis.
"""
from openpyxl import load_workbook
import re
import json

SRC = "/home/z/my-project/download/ventes_livrees.xlsx"

wb = load_workbook(SRC, read_only=True, data_only=True)

# Pattern: CLIENT COMPTOIR (case insensitive)
COMPTOIR_RE = re.compile(r"CLIENT\s+COMPTOIR", re.IGNORECASE)

clients = {}
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for row in ws.iter_rows(min_row=3, values_only=True):
        tiers = row[5] if len(row) > 5 else None
        if tiers is None or str(tiers).strip() == "":
            continue
        if tiers not in clients:
            clients[tiers] = 0
        ca = row[8] if len(row) > 8 else 0
        try:
            clients[tiers] += float(ca) if ca is not None else 0
        except:
            pass

wb.close()

# Filter for CLIENT COMPTOIR
comptoir_clients = []
print(f"Total unique clients: {len(clients)}")
print(f"\nClients 'CLIENT COMPTOIR' to EXCLUDE:")
print(f"{'Tiers':<80} {'CA HT 6 mois':>15}")
print("-" * 100)
for tiers, ca in sorted(clients.items(), key=lambda x: -x[1]):
    if COMPTOIR_RE.search(str(tiers)):
        comptoir_clients.append(tiers)
        print(f"{str(tiers)[:78]:<80} {ca:>15,.0f}")

print(f"\nTotal CLIENT COMPTOIR clients to exclude: {len(comptoir_clients)}")
total_ca_comptoir = sum(clients[t] for t in comptoir_clients)
print(f"Total CA of CLIENT COMPTOIR clients: {total_ca_comptoir:,.0f} FCFA")

# Load existing excluded list and merge
with open("/home/z/my-project/scripts/excluded_clients.json", "r", encoding="utf-8") as f:
    existing_excluded = set(json.load(f))

# Add the new comptoir clients
new_excluded = sorted(set(existing_excluded) | set(comptoir_clients))
print(f"\nExisting excluded (internal/filiale): {len(existing_excluded)}")
print(f"New CLIENT COMPTOIR to add: {len(comptoir_clients)}")
print(f"Total excluded after merge: {len(new_excluded)}")

# Save merged list
with open("/home/z/my-project/scripts/excluded_clients.json", "w", encoding="utf-8") as f:
    json.dump(new_excluded, f, ensure_ascii=False, indent=2)
print(f"\nMerged excluded list saved to: /home/z/my-project/scripts/excluded_clients.json")
