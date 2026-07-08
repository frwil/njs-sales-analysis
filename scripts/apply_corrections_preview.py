"""
Applique les corrections de catégorisation et calcule l'impact sur les volumes S1 2026.
- BELGOFOS (P105, P1051, P1053) → INGREDIENTS (était COMPLEMENT ALIMENTAIRE)
- BELGOTOX (E101, E1011, E1014) → INGREDIENTS (était COMPLEMENT ALIMENTAIRE)
- PIERRE A LECHER (PL102) → DIVERS (était COMPLEMENT ALIMENTAIRE)
- Liquides 1L : conversion 1L ≈ 1 kg (déjà gérée par la regex)
"""
import json
import re
import os
from collections import defaultdict
from openpyxl import load_workbook

SRC_VENTES = "/home/z/my-project/download/ventes_livrees.xlsx"

# Charger l'ancien mapping
with open("/home/z/my-project/scripts/product_category_map.json") as f:
    OLD_MAP = json.load(f)

# Construire le nouveau mapping
NEW_MAP = dict(OLD_MAP)
# BELGOFOS → INGREDIENTS
for p in ["P105", "P1051", "P1053"]:
    NEW_MAP[p] = "INGREDIENTS"
# BELGOTOX → INGREDIENTS
for p in ["E101", "E1011", "E1014"]:
    NEW_MAP[p] = "INGREDIENTS"
# PIERRE A LECHER → DIVERS
NEW_MAP["PL102"] = "DIVERS"

# Sauvegarder le nouveau mapping
with open("/home/z/my-project/scripts/product_category_map.json", "w", encoding="utf-8") as f:
    json.dump(NEW_MAP, f, ensure_ascii=False, indent=2)
print("✓ product_category_map.json mis à jour")

# Vérifier la nouvelle composition
print("\n" + "=" * 80)
print("NOUVELLE COMPOSITION DES CATÉGORIES IMPACTÉES")
print("=" * 80)
cats = defaultdict(list)
for p, c in NEW_MAP.items():
    cats[c].append(p)
for c in ["COMPLEMENT ALIMENTAIRE", "INGREDIENTS", "DIVERS"]:
    print(f"\n{c} ({len(cats[c])} produits):")
    for p in sorted(cats[c]):
        old = OLD_MAP.get(p, "?")
        marker = " ← RECLASSÉ" if old != c else ""
        print(f"  {p} (était: {old}){marker}")

# Parser de poids (identique à compute_sales_by_category.py)
WEIGHT_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(KG|Kg|kg|KG|GRAMMES?|G|L)\b", re.IGNORECASE)
def parse_weight_kg(desc):
    if not desc: return 0.0
    matches = WEIGHT_RE.findall(str(desc))
    if not matches: return 0.0
    val_str, unit = matches[-1]
    val = float(val_str)
    u = unit.upper()
    if u == "KG": return val
    if u in ("G", "GRAMME", "GRAMMES"): return val/1000.0
    if u == "L": return val  # 1L ≈ 1 kg
    return 0.0

# Calculer les volumes S1 avec nouveau mapping
print("\n" + "=" * 80)
print("IMPACT SUR LES VOLUMES S1 2026 (en tonnes)")
print("=" * 80)

# Charger excluded clients
with open("/home/z/my-project/scripts/excluded_clients.json") as f:
    excluded_tiers = set(json.load(f))

# Volumes avec EXCLUSION (pour ventes vs objectifs)
new_cat_sales = defaultdict(lambda: {"vol_kg": 0.0, "ca": 0.0})
old_cat_sales = defaultdict(lambda: {"vol_kg": 0.0, "ca": 0.0})

# Aussi sans exclusion (pour ventes totales)
new_cat_sales_no_excl = defaultdict(lambda: {"vol_kg": 0.0, "ca": 0.0})
old_cat_sales_no_excl = defaultdict(lambda: {"vol_kg": 0.0, "ca": 0.0})

wb = load_workbook(SRC_VENTES, read_only=True, data_only=True)
sheet_to_month = {"Sheet 1": 1, "Feuil1": 2, "Feuil2": 3, "Feuil3": 4, "Feuil4": 5, "Feuil5": 6}

for sheet_name in wb.sheetnames:
    if sheet_name not in sheet_to_month: continue
    ws = wb[sheet_name]
    for row in ws.iter_rows(min_row=3, values_only=True):
        tiers = row[5] if len(row) > 5 else None
        ref_prod = row[0] if len(row) > 0 else None
        desc = row[1] if len(row) > 1 else None
        qte = row[2] if len(row) > 2 else 0
        ca_ht = row[8] if len(row) > 8 else 0

        ref_str = str(ref_prod).strip() if ref_prod else ""
        if not ref_str: continue
        try: qte_f = float(qte) if qte else 0.0
        except: qte_f = 0.0
        try: ca_f = float(ca_ht) if ca_ht else 0.0
        except: ca_f = 0.0
        weight_kg = parse_weight_kg(desc)
        vol_kg = qte_f * weight_kg

        new_cat = NEW_MAP.get(ref_str, "DIVERS")
        old_cat = OLD_MAP.get(ref_str, "DIVERS")

        new_cat_sales_no_excl[new_cat]["vol_kg"] += vol_kg
        new_cat_sales_no_excl[new_cat]["ca"] += ca_f
        old_cat_sales_no_excl[old_cat]["vol_kg"] += vol_kg
        old_cat_sales_no_excl[old_cat]["ca"] += ca_f

        if tiers and tiers not in excluded_tiers:
            new_cat_sales[new_cat]["vol_kg"] += vol_kg
            new_cat_sales[new_cat]["ca"] += ca_f
            old_cat_sales[old_cat]["vol_kg"] += vol_kg
            old_cat_sales[old_cat]["ca"] += ca_f

wb.close()

# Afficher comparaison (sans exclusion — ventes totales)
print("\nVOLUMES S1 2026 (VENTES TOTALES — sans exclusion comptoirs)")
print(f"{'Catégorie':<28}{'Ancien (t)':>12}{'Nouveau (t)':>14}{'Écart (t)':>12}{'Nouveau CA (M)':>16}")
print("-" * 82)
for cat in sorted(set(list(new_cat_sales_no_excl.keys()) + list(old_cat_sales_no_excl.keys()))):
    old_vol = old_cat_sales_no_excl[cat]["vol_kg"] / 1000.0
    new_vol = new_cat_sales_no_excl[cat]["vol_kg"] / 1000.0
    new_ca = new_cat_sales_no_excl[cat]["ca"] / 1e6
    diff = new_vol - old_vol
    print(f"{cat:<28}{old_vol:>12.1f}{new_vol:>14.1f}{diff:>+12.1f}{new_ca:>16.1f}")

print("\nVOLUMES S1 2026 (VENTES EXCLUANT COMPTOIRS — pour ventes vs objectifs)")
print(f"{'Catégorie':<28}{'Ancien (t)':>12}{'Nouveau (t)':>14}{'Écart (t)':>12}{'Nouveau CA (M)':>16}")
print("-" * 82)
for cat in sorted(set(list(new_cat_sales.keys()) + list(old_cat_sales.keys()))):
    old_vol = old_cat_sales[cat]["vol_kg"] / 1000.0
    new_vol = new_cat_sales[cat]["vol_kg"] / 1000.0
    new_ca = new_cat_sales[cat]["ca"] / 1e6
    diff = new_vol - old_vol
    print(f"{cat:<28}{old_vol:>12.1f}{new_vol:>14.1f}{diff:>+12.1f}{new_ca:>16.1f}")

# Vérifier les objectifs vs nouveau réel
print("\n" + "=" * 80)
print("RECONFRONTATION AVEC LES OBJECTIFS (corrigés)")
print("=" * 80)

# Charger les objectifs
with open("/home/z/my-project/scripts/objectives_comparison.json") as f:
    obj_data = json.load(f)

print(f"\n{'Catégorie':<28}{'Nouveau S1 (t)':>16}{'Obj. S2 (t)':>14}{'Atteinte prévisible':>22}")
print("-" * 80)
for cat, info in obj_data.get("objectives_per_category", {}).items():
    new_vol = new_cat_sales.get(cat, {"vol_kg": 0})["vol_kg"] / 1000.0
    obj_s2 = info.get("s2_volume_tons", 0)
    if obj_s2 > 0:
        # Approx: si S2 ≈ S1 (sans saisonnalité), atteinte ≈ S1/S2
        atteinte = new_vol / obj_s2 * 100 if obj_s2 > 0 else 0
        print(f"{cat:<28}{new_vol:>16.1f}{obj_s2:>14.1f}{atteinte:>22.1f}%")
