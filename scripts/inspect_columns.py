"""Inspect the structure of the Excel file - show first 5 rows of each sheet."""
from openpyxl import load_workbook

SRC = "/home/z/my-project/upload/ventes janv a juin 2026.xlsx"

wb = load_workbook(SRC, read_only=True, data_only=True)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n{'='*80}")
    print(f"SHEET: {sheet_name}  (max_row={ws.max_row}, max_col={ws.max_column})")
    print('='*80)
    # Print first 5 rows
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True), start=1):
        print(f"\n--- Row {i} ---")
        for j, val in enumerate(row, start=1):
            col_letter = chr(ord('A') + j - 1) if j <= 26 else f"A{j}"
            if val is not None:
                print(f"  {col_letter}: {repr(val)[:100]}")

wb.close()
