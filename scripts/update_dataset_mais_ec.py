"""
Mise à jour du dataset 2023-2026 — Inclusion MAIS + En cours/Validées tous mois.

CHANGEMENTS:
1. MAIS (M1051, M1052) inclus comme famille 'MAIS' (produit opportuniste mais vendu)
2. En cours/Validées de S1 2026 + Juil 2026 + Août 2026 inclus (pas seulement août)
3. Objectif: matcher les chiffres utilisateur (YTD 2026 ~57697 t, LY 2025 ~44588 t)
"""
import openpyxl
import pandas as pd
import os
import json

# === Load existing dataset ===
print("Loading existing dataset...")
df = pd.read_csv("/home/z/my-project/scripts/dataset_2023_2026.csv", parse_dates=['date'], low_memory=False)
print(f"  Current records: {len(df)}")
print(f"  Current YTD 2026 total: {df[(df['date'].dt.year==2026)&(df['date'].dt.month<=8)]['tonnes'].sum():.0f} t")

# === 1. Add MAIS records from all years ===
print("\n=== Adding MAIS records ===")

MAIS_WEIGHT = 50  # 50 kg/sac (confirmed from 2025 ERP: 4154 t / 83083 sacs = 50.0 kg)
AGENCE_MAP_SHORT = {
    'AGENCE FAMLA': ('Famla', 'Ouest'), 'Famla': ('Famla', 'Ouest'), 'FAMLA': ('Famla', 'Ouest'),
    'AGENCE MESSASSI': ('Messassi', 'Centre'), 'Messassi': ('Messassi', 'Centre'), 'MESSASSI': ('Messassi', 'Centre'),
    'AGENCE BERTOUA': ('Bertoua', 'Centre'), 'Bertoua': ('Bertoua', 'Centre'), 'BERTOUA': ('Bertoua', 'Centre'),
    'AGENCE NDOBO': ('Ndobo', 'Littoral'), 'Ndobo': ('Ndobo', 'Littoral'), 'NDOBO': ('Ndobo', 'Littoral'),
    'AGENCE DJELENG': ('Djeleng', 'Ouest'), 'Djeleng': ('Djeleng', 'Ouest'), 'DJELENG': ('Djeleng', 'Ouest'),
    'AGENCE VILLAGE': ('Village', 'Littoral'), 'Village': ('Village', 'Littoral'), 'VILLAGE': ('Village', 'Littoral'),
    'AGENCE PK11': ('Pk11', 'Littoral'), 'Pk11': ('Pk11', 'Littoral'), 'PK11': ('Pk11', 'Littoral'),
    'AGENCE NGAOUNDERE': ('Ngaoundere', 'Centre'), 'Ngaoundere': ('Ngaoundere', 'Centre'), 'NGAOUNDERE': ('Ngaoundere', 'Centre'),
    'AGENCE AHALA': ('Ahala', 'Centre'), 'Ahala': ('Ahala', 'Centre'), 'AHALA': ('Ahala', 'Centre'),
    'AGENCE NKONGSAMBA': ('Nkongsamba', 'Littoral'), 'Nkongsamba': ('Nkongsamba', 'Littoral'), 'NKONGSAMBA': ('Nkongsamba', 'Littoral'),
    'AGENCE NKOLBISSON': ('Nkolbisson', 'Centre'), 'Nkolbisson': ('Nkolbisson', 'Centre'), 'NKOLBISSON': ('Nkolbisson', 'Centre'),
    'AGENCE NKOABANG': ('Nkoabang', 'Centre'), 'Nkoabang': ('Nkoabang', 'Centre'), 'NKOABANG': ('Nkoabang', 'Centre'),
    'AGENCE DE BAMENDA - DEPOT MBOUDA': ('Mbouda', 'Ouest'), 'Mbouda': ('Mbouda', 'Ouest'), 'MBOUDA': ('Mbouda', 'Ouest'),
    'AGENCE BUEA': ('Buea', 'Littoral'), 'Buea': ('Buea', 'Littoral'), 'BUEA': ('Buea', 'Littoral'),
    'AGENCE MAROUA': ('Maroua', 'Centre'),
}

mais_records = []

# 2025 file (with tonnes column)
print("  Loading MAIS 2025...")
wb2025 = openpyxl.load_workbook("/home/z/my-project/upload/86d96135-9db7-45bc-bba6-a69efa2c5ee3.xlsx", read_only=True, data_only=True)
ws2025 = wb2025['Feuil1']
rows2025 = list(ws2025.iter_rows(values_only=True))
header2025 = rows2025[0]
col_idx_2025 = {}
for i, h in enumerate(header2025):
    if h: col_idx_2025[str(h)] = i

for r in rows2025[1:]:
    if not r or not r[0]: continue
    ref = str(r[col_idx_2025['Réf. produit']])
    if ref not in ('M1051', 'M1052'): continue
    etat = str(r[col_idx_2025['État']]).strip() if r[col_idx_2025['État']] else ''
    if etat != 'Livrée': continue
    agence_raw = r[col_idx_2025['tableauProprieteAgences.Agence']]
    if agence_raw not in AGENCE_MAP_SHORT: continue
    agence, region = AGENCE_MAP_SHORT[agence_raw]
    qte = r[col_idx_2025['Qté commandée']] if r[col_idx_2025['Qté commandée']] else 0
    tonnes = r[col_idx_2025['Qté commandée (en tonnes)']] if r[col_idx_2025['Qté commandée (en tonnes)']] else qte * MAIS_WEIGHT / 1000
    ht = r[col_idx_2025['Montant HT']] if r[col_idx_2025['Montant HT']] else 0
    ttc = r[col_idx_2025['Montant TTC']] if r[col_idx_2025['Montant TTC']] else 0
    date = r[col_idx_2025['Date de commande']]
    if pd.isna(date): continue
    date = pd.to_datetime(date)
    mais_records.append({
        'date': date, 'year': date.year, 'month': date.month,
        'ref': ref, 'family': 'MAIS', 'description': r[col_idx_2025['Description du produit']],
        'client': r[col_idx_2025['Tiers']] if r[col_idx_2025['Tiers']] else '',
        'agence': agence, 'region': region,
        'qte': qte, 'weight_kg': MAIS_WEIGHT, 'kg': qte * MAIS_WEIGHT,
        'tonnes': tonnes, 'sacs_50': qte,
        'montant_ttc': ttc, 'montant_ht': ht,
        'source': 'MAIS_2025'
    })
print(f"    MAIS 2025: {len(mais_records)} records")

# 2026 S1 file
print("  Loading MAIS S1 2026...")
wb_s1 = openpyxl.load_workbook("/home/z/my-project/upload/ventes janv a juin 2026.xlsx", read_only=True, data_only=True)
for sheet_name in wb_s1.sheetnames:
    ws = wb_s1[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 3: continue
    header = rows[1]
    cols = {}
    for i, h in enumerate(header):
        if h: cols[str(h)] = i
    if 'Réf. produit' not in cols or 'État' not in cols: continue
    for r in rows[2:]:
        if not r or not r[0]: continue
        ref = str(r[cols['Réf. produit']])
        if ref not in ('M1051', 'M1052'): continue
        etat = str(r[cols['État']]).strip() if r[cols['État']] else ''
        if etat != 'Livrée': continue
        agence_raw = r[cols['agence']]
        if agence_raw not in AGENCE_MAP_SHORT: continue
        agence, region = AGENCE_MAP_SHORT[agence_raw]
        qte = r[cols['Qté commandée']] if r[cols['Qté commandée']] else 0
        tonnes = qte * MAIS_WEIGHT / 1000
        ht = r[cols.get('Montant HT', 0)] if r[cols.get('Montant HT', 0)] else 0
        ttc = r[cols.get('Montant TTC', 0)] if r[cols.get('Montant TTC', 0)] else 0
        date = pd.to_datetime(str(r[cols['Date de commande']])[:10], format='%d/%m/%Y')
        mais_records.append({
            'date': date, 'year': date.year, 'month': date.month,
            'ref': ref, 'family': 'MAIS', 'description': r[cols['Description du produit']],
            'client': r[cols.get('Tiers', '')] if r[cols.get('Tiers', '')] else '',
            'agence': agence, 'region': region,
            'qte': qte, 'weight_kg': MAIS_WEIGHT, 'kg': qte * MAIS_WEIGHT,
            'tonnes': tonnes, 'sacs_50': qte,
            'montant_ttc': ttc, 'montant_ht': ht,
            'source': 'MAIS_S1_2026'
        })
print(f"    MAIS S1 2026: {len([m for m in mais_records if m['source']=='MAIS_S1_2026'])} records")

# 2026 Juil + Août
for fpath, label, date_filter in [
    ("/home/z/my-project/upload/NJS GROUP ERP - Lignes de commandes + multicompany (9).xlsx", "MAIS_Juil_2026", "/07/2026"),
    ("/home/z/my-project/upload/NJS GROUP ERP - Lignes de commandes + multicompany (27).xlsx", "MAIS_Aout_2026", "/08/2026"),
]:
    wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
    ws = wb['Sheet 1']
    rows = list(ws.iter_rows(values_only=True))
    header = rows[1]
    cols = {}
    for i, h in enumerate(header):
        if h: cols[str(h)] = i
    for r in rows[2:]:
        if not r or not r[0]: continue
        ref = str(r[cols['Réf. produit']])
        if ref not in ('M1051', 'M1052'): continue
        etat = str(r[cols['État']]).strip() if r[cols['État']] else ''
        if etat != 'Livrée': continue
        date_str = str(r[cols['Date de commande']])[:10]
        if date_filter not in date_str: continue
        agence_raw = r[cols['agence']]
        if agence_raw not in AGENCE_MAP_SHORT: continue
        agence, region = AGENCE_MAP_SHORT[agence_raw]
        qte = r[cols['Qté commandée']] if r[cols['Qté commandée']] else 0
        tonnes = qte * MAIS_WEIGHT / 1000
        ht = r[cols['Montant HT']] if r[cols['Montant HT']] else 0
        ttc = r[cols['Montant TTC']] if r[cols['Montant TTC']] else 0
        date = pd.to_datetime(date_str, format='%d/%m/%Y')
        mais_records.append({
            'date': date, 'year': date.year, 'month': date.month,
            'ref': ref, 'family': 'MAIS', 'description': r[cols['Description du produit']],
            'client': r[cols.get('Tiers', '')] if r[cols.get('Tiers', '')] else '',
            'agence': agence, 'region': region,
            'qte': qte, 'weight_kg': MAIS_WEIGHT, 'kg': qte * MAIS_WEIGHT,
            'tonnes': tonnes, 'sacs_50': qte,
            'montant_ttc': ttc, 'montant_ht': ht,
            'source': label
        })

# LY_24 (Jul-Dec 2024)
print("  Loading MAIS LY_24 (Jul-Dec 2024)...")
wb24 = openpyxl.load_workbook("/home/z/my-project/upload/21_24.xlsx", read_only=True, data_only=True)
ws24 = wb24['LY_24']
rows24 = list(ws24.iter_rows(values_only=True))
for r in rows24[1:]:
    if not r or not r[0]: continue
    ref = str(r[0])
    if ref not in ('M1051', 'M1052'): continue
    etat = str(r[10]).strip() if r[10] else ''
    if etat != 'Livrée': continue
    agence = str(r[14]).strip() if r[14] else ''
    agence_title = agence.upper()
    if agence_title not in AGENCE_MAP_SHORT: continue
    ag_short, region = AGENCE_MAP_SHORT[agence_title]
    tonnes = r[12] if r[12] else 0
    ht = r[7] if r[7] else 0
    date = r[5]
    if isinstance(date, str):
        try: date = pd.to_datetime(date, format='%d/%m/%Y')
        except: continue
    if pd.isna(date): continue
    qte = r[2] if r[2] else 0
    mais_records.append({
        'date': date, 'year': date.year, 'month': date.month,
        'ref': ref, 'family': 'MAIS', 'description': r[1],
        'client': '', 'agence': ag_short, 'region': region,
        'qte': qte, 'weight_kg': MAIS_WEIGHT, 'kg': qte * MAIS_WEIGHT,
        'tonnes': tonnes, 'sacs_50': qte,
        'montant_ttc': 0, 'montant_ht': ht,
        'source': 'MAIS_LY_24'
    })

print(f"  Total MAIS records: {len(mais_records)}")

# === 2. Add En cours/Validées from S1 2026 + Juil 2026 ===
print("\n=== Adding En cours/Validées from S1 + Juil 2026 ===")

# Product refs and weights
ALL_WEIGHTS = {
    'T102': 50, 'T1021': 1, 'T1023': 5, 'T1024': 25,
    'C101': 50, 'C102': 50, 'C103': 50, 'C104': 50, 'C1042': 1, 'C1043': 5, 'C1044': 25,
    'C105': 50, 'C1053': 1, 'C1054': 5, 'C1055': 25, 'C108': 50, 'C1022': 5,
    'CB100': 25, 'CB200': 25, 'CB101': 5, 'CB201': 5, 'PB100': 25, 'PB200': 25, 'DB100': 25, 'DB200': 25, 'ALAP25': 25,
    'B100': 25, 'E101': 25, 'I105': 25, 'P105': 25, 'I106': 25, 'I107': 25, 'F114': 50, 'F1145': 50, 'F1146': 25, 'F1147': 1,
    'B1001': 1, 'B1003': 5, 'B1004': 25, 'E1011': 1, 'E1013': 5, 'E1014': 0.2,
    'I1051': 1, 'I1053': 5, 'I1054': 25, 'I1061': 1, 'I1071': 1,
    'P1051': 1, 'P1053': 5, 'P102N2': 25, 'P104N2': 25, 'P109': 25,
    'V300': 1, 'CA003.1': 1, 'CA004.1': 1, 'CA006.1': 1, 'CA001.1': 1, 'CA002.1': 1, 'CA005.1': 1, 'CA007.1': 1, 'CA008.1': 1,
    'MAT011-80010002': 1, 'MAT014-80010003': 1, 'MAT015': 1, 'MAT017': 1,
    'M1051': 50, 'M1052': 50,  # MAIS
}

def get_family_ext(ref):
    SOJA = {'T102', 'T1021', 'T1023', 'T1024'}
    CONC = {'C101', 'C102', 'C103', 'C104', 'C1042', 'C1043', 'C1044', 'C105', 'C1053', 'C1054', 'C1055', 'C108', 'C1022'}
    ALIM = {'CB100', 'CB200', 'CB101', 'CB201', 'PB100', 'PB200', 'DB100', 'DB200', 'ALAP25'}
    ING = {'B100', 'E101', 'I105', 'B1001', 'B1003', 'B1004', 'E1011', 'E1013', 'E1014', 'I1051', 'I1053', 'I1054', 'I106', 'I1061', 'I107', 'I1071', 'P105', 'P1051', 'P1053', 'F114', 'F1145', 'F1146', 'F1147'}
    PREMIX = {'P102N2', 'P104N2', 'P109', 'PX101', 'PX102', 'PX103', 'PX104', 'PX105'}
    COMP = {'V300', 'CA003.1', 'CA004.1', 'CA006.1', 'CA001.1', 'CA002.1', 'CA005.1', 'CA007.1', 'CA008.1'}
    ALV = {'MAT011-80010002', 'MAT014-80010003', 'MAT015', 'MAT017'}
    MAIS = {'M1051', 'M1052'}
    if ref in SOJA: return 'TOURTEAUX'
    if ref in CONC: return 'CONCENTRES'
    if ref in ALIM: return 'ALIMENT_COMPLET'
    if ref in ING: return 'INGREDIENTS'
    if ref in PREMIX: return 'PREMIX'
    if ref in COMP: return 'COMPLEMENT_ALIMENTAIRE'
    if ref in ALV: return 'ALVEOLES'
    if ref in MAIS: return 'MAIS'
    if ref.startswith('MAT') or ref.startswith('ME'): return 'MATERIEL_ELEVAGE'
    return None

ec_records = []

# S1 2026 En cours/Validées
for sheet_name in wb_s1.sheetnames:
    ws = wb_s1[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 3: continue
    header = rows[1]
    cols = {}
    for i, h in enumerate(header):
        if h: cols[str(h)] = i
    if 'État' not in cols or 'Réf. produit' not in cols: continue
    for r in rows[2:]:
        if not r or not r[0]: continue
        etat = str(r[cols['État']]).strip() if r[cols['État']] else ''
        if etat not in ('En cours', 'Validée'): continue
        ref = str(r[cols['Réf. produit']])
        family = get_family_ext(ref)
        if family is None: continue
        weight = ALL_WEIGHTS.get(ref, 0)
        if weight == 0: continue
        agence_raw = r[cols['agence']]
        if agence_raw not in AGENCE_MAP_SHORT: continue
        agence, region = AGENCE_MAP_SHORT[agence_raw]
        qte = r[cols['Qté commandée']] if r[cols['Qté commandée']] else 0
        kg = qte * weight
        tonnes = kg / 1000
        sacs_50 = kg / 50 if family not in ('MATERIEL_ELEVAGE', 'COMPLEMENT_ALIMENTAIRE', 'ALVEOLES') else 0
        if family in ('MATERIEL_ELEVAGE', 'ALVEOLES'):
            tonnes = 0
        ht = r[cols.get('Montant HT', 0)] if r[cols.get('Montant HT', 0)] else 0
        ttc = r[cols.get('Montant TTC', 0)] if r[cols.get('Montant TTC', 0)] else 0
        date = pd.to_datetime(str(r[cols['Date de commande']])[:10], format='%d/%m/%Y')
        ec_records.append({
            'date': date, 'year': date.year, 'month': date.month,
            'ref': ref, 'family': family, 'description': r[cols['Description du produit']],
            'client': r[cols.get('Tiers', '')] if r[cols.get('Tiers', '')] else '',
            'agence': agence, 'region': region,
            'qte': qte, 'weight_kg': weight, 'kg': kg, 'tonnes': tonnes, 'sacs_50': sacs_50,
            'montant_ttc': ttc, 'montant_ht': ht,
            'source': 'EC_S1_2026'
        })

# Juil 2026 En cours/Validées
wb_juil = openpyxl.load_workbook("/home/z/my-project/upload/NJS GROUP ERP - Lignes de commandes + multicompany (9).xlsx", read_only=True, data_only=True)
ws_juil = wb_juil['Sheet 1']
rows_juil = list(ws_juil.iter_rows(values_only=True))
header_juil = rows_juil[1]
cols_juil = {}
for i, h in enumerate(header_juil):
    if h: cols_juil[str(h)] = i
for r in rows_juil[2:]:
    if not r or not r[0]: continue
    etat = str(r[cols_juil['État']]).strip() if r[cols_juil['État']] else ''
    if etat not in ('En cours', 'Validée'): continue
    date_str = str(r[cols_juil['Date de commande']])[:10]
    if '/07/2026' not in date_str: continue
    ref = str(r[cols_juil['Réf. produit']])
    family = get_family_ext(ref)
    if family is None: continue
    weight = ALL_WEIGHTS.get(ref, 0)
    if weight == 0: continue
    agence_raw = r[cols_juil['agence']]
    if agence_raw not in AGENCE_MAP_SHORT: continue
    agence, region = AGENCE_MAP_SHORT[agence_raw]
    qte = r[cols_juil['Qté commandée']] if r[cols_juil['Qté commandée']] else 0
    kg = qte * weight
    tonnes = kg / 1000
    sacs_50 = kg / 50 if family not in ('MATERIEL_ELEVAGE', 'COMPLEMENT_ALIMENTAIRE', 'ALVEOLES') else 0
    if family in ('MATERIEL_ELEVAGE', 'ALVEOLES'):
        tonnes = 0
    ht = r[cols_juil['Montant HT']] if r[cols_juil['Montant HT']] else 0
    ttc = r[cols_juil['Montant TTC']] if r[cols_juil['Montant TTC']] else 0
    date = pd.to_datetime(date_str, format='%d/%m/%Y')
    ec_records.append({
        'date': date, 'year': date.year, 'month': date.month,
        'ref': ref, 'family': family, 'description': r[cols_juil['Description du produit']],
        'client': r[cols_juil.get('Tiers', '')] if r[cols_juil.get('Tiers', '')] else '',
        'agence': agence, 'region': region,
        'qte': qte, 'weight_kg': weight, 'kg': kg, 'tonnes': tonnes, 'sacs_50': sacs_50,
        'montant_ttc': ttc, 'montant_ht': ht,
        'source': 'EC_Juil_2026'
    })

print(f"  En cours/Validées (S1+Juil): {len(ec_records)} records")

# === Merge all ===
print("\n=== Merging ===")
df_mais = pd.DataFrame(mais_records)
df_ec = pd.DataFrame(ec_records)
df_all = pd.concat([df, df_mais, df_ec], ignore_index=True)
print(f"  Total records: {len(df_all)}")

# Verify
ytd_2026 = df_all[(df_all['date'].dt.year == 2026) & (df_all['date'].dt.month <= 8)]
ytd_2025 = df_all[(df_all['date'].dt.year == 2025) & (df_all['date'].dt.month <= 8)]
print(f"\n=== VERIFICATION ===")
print(f"  YTD 2026 total: {ytd_2026['tonnes'].sum():.0f} t (user: 57697)")
print(f"  YTD 2026 TOURTEAUX: {ytd_2026[ytd_2026['family']=='TOURTEAUX']['tonnes'].sum():.0f} t (user: 41560)")
print(f"  YTD 2025 total: {ytd_2025['tonnes'].sum():.0f} t (user: 44588)")
print(f"  MAIS 2026 YTD: {ytd_2026[ytd_2026['family']=='MAIS']['tonnes'].sum():.0f} t")
print(f"  MAIS 2025 YTD: {ytd_2025[ytd_2025['family']=='MAIS']['tonnes'].sum():.0f} t")
print(f"  EC 2026 YTD: {ytd_2026[ytd_2026['source'].str.startswith('EC', na=False)]['tonnes'].sum():.0f} t")

# Save
output_path = "/home/z/my-project/scripts/dataset_2023_2026.csv"
df_all.to_csv(output_path, index=False)
print(f"\nDataset saved: {output_path}")
print(f"File size: {os.path.getsize(output_path) / 1024:.0f} KB")
print(f"\nFamilies: {sorted(df_all['family'].unique().tolist())}")
