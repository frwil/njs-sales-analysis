"""
Clean the sales file: keep only rows where 'État' column (P, index 16) == 'Livrée'.

Structure of each sheet:
- Row 1: Title (BELGOCAM SA - NJS GROUP ERP - ...)
- Row 2: Headers (Réf. produit, Description, ..., État, ...)
- Row 3+: Data

Strategy:
- Load workbook with openpyxl (preserves formatting)
- For each sheet, identify rows where column P != 'Livrée'
- Delete them from bottom to top (to keep row indices stable)
"""
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import shutil
import os
import time

SRC = "/home/z/my-project/upload/ventes janv a juin 2026.xlsx"
DST = "/home/z/my-project/download/ventes janv a juin 2026 - Livrées uniquement.xlsx"

# Make a copy first, then edit in place - safer and faster than rebuilding styles
os.makedirs(os.path.dirname(DST), exist_ok=True)
print(f"Copying source file to {DST} ...")
shutil.copy2(SRC, DST)

print("Opening workbook (this preserves formatting) ...")
t0 = time.time()
wb = load_workbook(DST)
print(f"  Loaded in {time.time()-t0:.1f}s")

ETAT_COL = 16  # Column P (1-based)

total_kept = 0
total_removed = 0

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    t1 = time.time()
    print(f"\nProcessing sheet '{sheet_name}' (max_row={ws.max_row}) ...")

    # Collect row indices to delete (data rows start at row 3)
    rows_to_delete = []
    for row_idx in range(3, ws.max_row + 1):
        cell_val = ws.cell(row=row_idx, column=ETAT_COL).value
        if cell_val != "Livrée":
            rows_to_delete.append(row_idx)

    print(f"  Rows to delete: {len(rows_to_delete)}")
    print(f"  Rows to keep (Livrée): {(ws.max_row - 2) - len(rows_to_delete)}")

    # Delete from bottom to top to keep indices stable
    # Group consecutive rows into ranges for batch deletion (faster)
    if rows_to_delete:
        # Build ranges of consecutive indices
        ranges = []
        start = rows_to_delete[0]
        prev = start
        for r in rows_to_delete[1:]:
            if r == prev + 1:
                prev = r
            else:
                ranges.append((start, prev))
                start = r
                prev = r
        ranges.append((start, prev))

        print(f"  Deleting in {len(ranges)} batch(es) ...")
        # Delete from bottom to top
        for start, end in reversed(ranges):
            count = end - start + 1
            ws.delete_rows(start, count)

    kept = (ws.max_row - 2) if ws.max_row >= 2 else 0
    total_kept += kept
    total_removed += len(rows_to_delete)
    print(f"  Sheet done in {time.time()-t1:.1f}s — final rows: {ws.max_row} (kept {kept} data rows)")

print(f"\nSaving to {DST} ...")
t2 = time.time()
wb.save(DST)
print(f"  Saved in {time.time()-t2:.1f}s")

wb.close()

print(f"\n{'='*70}")
print(f"SUMMARY")
print(f"{'='*70}")
print(f"Total 'Livrée' rows kept : {total_kept}")
print(f"Total rows removed       : {total_removed}")
print(f"Output file              : {DST}")
print(f"File size                : {os.path.getsize(DST):,} bytes")
